"""Ticket 47: the engine computes the QEP major power producers' gas price to Ofgem WMI NBP day-ahead ratio,
yearly 2010-25, with the mean, min and max, from two registered raw files (DESNZ QEP 3.2.1 (Annual) column F,
and the Ofgem WMI gas day-ahead monthly series). The independent source of truth is the Researcher's derivation
CSV, registered as a test reference (never read by the engine). CSV values are 4 dp, so tolerances are 5e-5."""
from __future__ import annotations

import csv
import hashlib
import statistics
from pathlib import Path

import openpyxl
import pytest

from gascf import prices as P

CSV = Path(__file__).resolve().parent.parent / "data/raw/s36/s36-qep-ratio-1065-derivation.csv"
QEP_SHA = "b58d7dc9dcb3313876b77f31af59a1e2b218b7872211550602fe024b5d57740c"
WMI_SHA = "a304e0726a41a838d71b31d913c3505fa611a4e79b4392ec97871463e527ff57"


def _derivation():
    by_year, mean = {}, None
    for row in csv.DictReader(CSV.open()):
        if row["year"] == "mean_of_16_annual_ratios":
            mean = float(row["ratio"])
        elif row["year"].isdigit():
            by_year[int(row["year"])] = {k: row[k] for k in row}
    return by_year, mean


def test_yearly_ratios_match_the_derivation_csv():
    d, _ = _derivation()
    got = P.qep_dayahead_ratio_2010_25()["by_year"]
    assert sorted(got) == list(range(2010, 2026)) == sorted(d)
    for y, exp in d.items():
        assert got[y]["ratio"] == pytest.approx(float(exp["ratio"]), abs=5e-5)
        assert got[y]["qep_p_per_kwh"] == float(exp["qep_p_per_kwh"])       # ==, the workbook's stored double
        assert got[y]["wmi_months"] == 12
        # CSV stores the mean to 3 dp, so the gap is at most half of that, 5e-4; 2023 sits exactly on the
        # boundary (99.2225 vs 99.222, a gap of 0.0005000000000024), so allow a floating-point margin of 1e-12.
        assert got[y]["wmi_mean_p_per_therm"] == pytest.approx(float(exp["wmi_mean_p_per_therm"]), abs=5e-4 + 1e-12)


def test_mean_matches_the_derivation_csv():
    d, csv_mean = _derivation()
    r = P.qep_dayahead_ratio_2010_25()
    assert r["mean"] == pytest.approx(csv_mean, abs=5e-5)
    # plain mean of the 16 yearly ratios, not a ratio of mean prices
    plain = statistics.mean(r["by_year"][y]["ratio"] for y in r["by_year"])
    assert r["mean"] == pytest.approx(plain, abs=1e-12)
    ratio_of_means = (statistics.mean(r["by_year"][y]["qep_gbp_mwh"] for y in r["by_year"])
                      / statistics.mean(r["by_year"][y]["wmi_gbp_mwh"] for y in r["by_year"]))
    assert abs(r["mean"] - ratio_of_means) > 5e-5          # the two are different numbers
    assert round(r["mean"], 3) == 1.065


def test_min_is_2021_and_max_is_2020():
    d, _ = _derivation()
    r = P.qep_dayahead_ratio_2010_25()
    assert r["min"]["year"] == 2021 and r["max"]["year"] == 2020
    assert r["min"]["value"] == pytest.approx(float(d[2021]["ratio"]), abs=5e-5)
    assert r["max"]["value"] == pytest.approx(float(d[2020]["ratio"]), abs=5e-5)
    ratios = {y: r["by_year"][y]["ratio"] for y in r["by_year"]}
    assert r["min"]["value"] == min(ratios.values()) and r["max"]["value"] == max(ratios.values())


def test_inputs_are_the_registered_files_with_matching_sha256():
    r = P.qep_dayahead_ratio_2010_25()
    shas = {Path(i["path"]).name: i["sha256"] for i in r["inputs"]}
    assert shas["s18a-QEP_table_321.xlsx"] == QEP_SHA
    assert shas["s18a-WMI_gas_dayahead_monthly_YFsS4g98P.js"] == WMI_SHA
    for i in r["inputs"]:
        assert hashlib.sha256(Path(i["path"]).read_bytes()).hexdigest() == i["sha256"]


def test_values_are_computed_sensitivity_wmi_and_qep(tmp_path):
    base = P.qep_dayahead_ratio_2010_25()
    # WMI: a temp everviz file with every 2020 month doubled -> 2020 ratio halves (path keyed, nothing in data/raw)
    import json as J
    series = P.read_wmi(P.WMI_GAS)
    lines = ["Category,NBP"]
    for (y, m), v in sorted(series.items()):
        lines.append(f'"{m:02d}/01/{y}",{v * 2 if y == 2020 else v}')
    wmi2 = tmp_path / "wmi2020x2.js"
    wmi2.write_text('{"csv":' + J.dumps("\n".join(lines)) + "}")
    r_wmi = P.qep_dayahead_ratio_2010_25(wmi_path=wmi2)
    assert r_wmi["by_year"][2020]["ratio"] == pytest.approx(base["by_year"][2020]["ratio"] / 2, abs=1e-12)
    assert r_wmi["max"]["year"] != 2020
    assert r_wmi["mean"] != base["mean"]
    # QEP: a temp workbook with the 2021 annual cell (F43) doubled -> 2021 ratio doubles, min moves off 2021
    wb = openpyxl.load_workbook(P.QEP)
    ws = wb["3.2.1 (Annual)"]
    ws["F43"] = ws["F43"].value * 2
    qep2 = tmp_path / "qep_f43x2.xlsx"
    wb.save(qep2)
    r_qep = P.qep_dayahead_ratio_2010_25(qep_path=qep2)
    assert r_qep["by_year"][2021]["ratio"] == pytest.approx(base["by_year"][2021]["ratio"] * 2, abs=1e-9)
    assert r_qep["min"]["year"] != 2021


def test_no_typed_ratio_literals_in_source():
    root = Path(__file__).resolve().parent.parent
    bad = ("1.065", "1.0646", "1.06465", "0.7384", "0.73844", "1.4056", "1.40559")
    for base in ("src", "tools"):
        for p in (root / base).rglob("*.py"):
            t = p.read_text()
            for lit in bad:
                assert lit not in t, f"{p} contains ratio literal {lit}"
