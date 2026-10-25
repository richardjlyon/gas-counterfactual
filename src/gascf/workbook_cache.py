"""Read each Excel workbook once per process (card [internal ticket], 20 Oct 2026).

One export used to open the EGC cost workbooks 598 times (`fossil_fleet.edition` has no cache, and
EGC 2020 KeyData re-parses a 20.7 MB external-link part on every open) and the ECR capacity workbook
1,351 times. This module keeps the rows of each sheet it has read, so the same sheet is parsed once.

What it caches is only what openpyxl reads, in read-only, values-only mode:
`list(load_workbook(path, read_only=True, data_only=True)[sheet].iter_rows(values_only=True))`, and
the sheet names. The readers that turn rows into plants and factors still run on every call, so
every layout guard in them still fires on every call (s36 B-merge cause 4: a cache once hid a guard
by caching the body that held it).

The key is the file's resolved path, the sha256 of its bytes and the read options. A workbook whose
contents change is read again, whatever its timestamp says; a copy at another path is read on its
own. The file is hashed on every call; that cost is inside the measured export time (163 s on the
Mac mini, against 1,028 s before this change).

Rows come back as a new list of tuples each call, so a caller that changes the list cannot change
what the next caller reads (the cells are numbers, strings, dates or None).
"""
from __future__ import annotations

import hashlib
from pathlib import Path

import openpyxl

_NAMES: dict[tuple, tuple[str, ...]] = {}
_ROWS: dict[tuple, tuple[tuple, ...]] = {}


def _key(path: Path | str, read_only: bool, data_only: bool) -> tuple:
    p = Path(path).resolve()
    return (str(p), hashlib.sha256(p.read_bytes()).hexdigest(), read_only, data_only)


def _read(path: Path | str, key: tuple, sheet: str | None, read_only: bool, data_only: bool) -> None:
    wb = openpyxl.load_workbook(path, read_only=read_only, data_only=data_only)
    try:
        _NAMES[key] = tuple(wb.sheetnames)
        if sheet is not None:
            _ROWS[key + (sheet,)] = tuple(wb[sheet].iter_rows(values_only=True))   # KeyError if absent
    finally:
        wb.close()


def sheetnames(path: Path | str, *, read_only: bool = True, data_only: bool = True) -> list[str]:
    """`load_workbook(path, ...).sheetnames`, read once per file contents."""
    key = _key(path, read_only, data_only)
    if key not in _NAMES:
        _read(path, key, None, read_only, data_only)
    return list(_NAMES[key])


def sheet_rows(path: Path | str, sheet: str, *, read_only: bool = True, data_only: bool = True) -> list[tuple]:
    """`list(load_workbook(path, ...)[sheet].iter_rows(values_only=True))`, read once per file
    contents and sheet. Raises KeyError for a sheet the workbook does not have, as openpyxl does."""
    key = _key(path, read_only, data_only)
    if key + (sheet,) not in _ROWS:
        _read(path, key, sheet, read_only, data_only)
    return list(_ROWS[key + (sheet,)])


def cache_clear() -> None:
    _NAMES.clear()
    _ROWS.clear()
