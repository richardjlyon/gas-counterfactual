"""Layer 2: policy costs moved from the bill to the taxpayer.

Spec §6, step 4(a). Two items:

1. **Renewables Obligation (RO) paid by the Treasury, April 2026 to March 2029.** The RO is
   the old renewables subsidy: suppliers must buy certificates (ROCs) for each MWh they sell.
   From 1 April 2026 the government pays suppliers 75% of the RO cost attributable to domestic
   supply (RO SI explanatory memorandum, para 4.2). The bill keeps 25%. Ofgem's cap carries the
   25% (Annex 4 tab 1a, RO £/MWh supplied); so the Treasury's 75% is exactly 3 x that bill share
   on the same basis. At the receipt's 2,500 kWh this is £61.37 a year (spec §6, spec.md:126;
   3 x the bill share rounded to £20.46 would give £61.38, which is why we never round first).
   Ofgem's Annex 4 £/MWh is an ex-ante allowance, so it is tagged ``allowance`` (review B2).

   Cross-checks, each from the government's own words:
     * DESNZ supplier guidance: discount = 0.75 x buy-out price x obligation level
       = 0.75 x £69.34/ROC x 0.472 ROC/MWh = 2.45 p/kWh for 2026/27;
     * Budget page (27 Nov 2025): "funding 75% of RO ... should save £88" per *average*
       household, and "£24.80/MWh off the unit rate for 75% of RO (excluding VAT)". The page
       counts the VAT saving separately ("These measures together should save £7 in VAT"), so
       its £88 is ex VAT;
     * RO SI explanatory memorandum (SI 2026/452, para 9.3): "£87 on average or £66 for the
       typical dual fuel household". £66 = 2.7 MWh x the 2.45 p/kWh discount. **The EM does not
       state whether £87 includes VAT** (the word does not appear in it; review N2). Ex VAT it
       implies 3.54 MWh a household; inc VAT, 3.38 MWh. DUKES 5.1 domestic / GB households
       gives 3.36 MWh (2024) and 3.40 MWh (2025), so the data cannot tell the two apart and we
       do not infer a basis. Both government figures are *average household*, not 2,500 kWh,
       which is why they exceed £61.

2. **Energy Price Guarantee (EPG), October 2022 to June 2023 (note only).** The Treasury paid
   suppliers a per-kWh discount off the cap's unit rates. DESNZ's policy paper states the
   discount for each quarter (GB, ex VAT). We report the Treasury cost per household per
   quarter at 2,500 kWh (electricity), at Ofgem's 3,100 kWh then in force, and the gas part
   for a dual-fuel note. The standing charge was not discounted. The NAO puts the scheme's
   total at £23,999m (GB and NI, Oct 2022 to Mar 2024).
"""
from __future__ import annotations

import html
import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from .system_costs import Line, gb_households, gb_households_fy, gb_households_series  # noqa: F401

# The engine's one household count (review S3) is defined in system_costs and re-exported here;
# goods.py imports it from this module. ONS LFS for England, Wales and Scotland (households.py).

ROOT = Path(__file__).parents[2]
GOVUK = ROOT / "data/raw/govuk"
ANNEX4 = ROOT / "data/raw/ofgem/Annex-4-policy-cost-allowance-methodology-v1.25.xlsx"
CAP_MODEL = ROOT / "data/raw/ofgem/Default-tariff-cap-level-v1.31.xlsx"
RO_GUIDANCE = GOVUK / "RO_to_Exchequer_supplier_guidance.html"
BUDGET_PAGE = GOVUK / "Autumn_Budget_energy_bills_archived_20260302.html"
RO_EM = GOVUK / "RO_to_Exchequer_EM_uksi_2026_452.pdf"
EPG_PAPER = GOVUK / "EPG_up_to_30_June_2023.html"
EPG_RATES = {"Oct-Dec 2022": GOVUK / "EPG_regional_rates_Oct-Dec_2022.html",
             "Jan-Mar 2023": GOVUK / "EPG_regional_rates_Jan-Mar_2023.html",
             "Apr-Jun 2023": GOVUK / "EPG_regional_rates_Apr-Jun_2023.html"}
NAO = GOVUK / "NAO_Energy_bills_support_an_update_HC232_2024.pdf"

RECEIPT_MWH = 2.5
TREASURY_SHARE = 0.75
RO_SCHEME_YEARS = 3          # April 2026 to March 2029 (EM para 5.4: "between 2026 and 2029")

SRC_A4 = "ofgem-annex4-v1.25"
VAT_DOMESTIC = 0.05            # reduced rate on domestic fuel to Sep 2026 (VATA 1994 Sch 7A)
BUDGET_VAT_QUOTE = "These measures together should save £7 in VAT"
SRC_GUIDE = "desnz_ro_to_exchequer_supplier_guidance"
SRC_BUDGET = "hmt_desnz_autumn_budget_energy_bills_2025"
SRC_EM = "uksi_2026_452_explanatory_memorandum"
SRC_EPG = "desnz_epg_up_to_30_june_2023"
SRC_EPG_RATES = "desnz_epg_regional_rates"
SRC_CAP = "ofgem-cap-model-v1.31"
SRC_NAO = "nao_energy_bills_support_update_2024"


def _page_text(path: Path) -> str:
    """Visible text of a stored gov.uk page (from <main>), whitespace-collapsed."""
    s = path.read_text(encoding="utf-8")
    i = s.find("<main")
    if i < 0:
        raise ValueError(f"{path.name}: no <main> element; page layout changed")
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", s[i:])))


def _pdf_text(path: Path) -> str:
    import pymupdf
    return re.sub(r"\s+", " ", " ".join(p.get_text() for p in pymupdf.open(path)))


def _grab(text: str, pattern: str, what: str) -> re.Match:
    m = re.search(pattern, text)
    if not m:
        raise ValueError(f"{what}: anchor {pattern!r} not found; source wording changed")
    return m


# --------------------------------------------------------------------------- RO

def read_annex4_ro() -> dict[str, float]:
    """{cap period label: RO £/MWh supplied}, GB average, Annex 4 tab 1a row 53."""
    rows = list(openpyxl.load_workbook(ANNEX4, read_only=True, data_only=True)
                ["1a Policy Cost Allowance"].iter_rows(values_only=True))
    if not str(rows[44][1]).startswith("2. Scheme by scheme estimates (GB average)"):
        raise ValueError("Annex 4 1a: B45 heading moved")
    if (rows[52][1], rows[52][2], rows[52][3]) != (
            "Electricity - Single-Rate Metering Arrangement", "RO", "£/MWh supplied"):
        raise ValueError("Annex 4 1a: RO row moved from row 53")
    out = {}
    for j, label in enumerate(rows[11]):
        if label and isinstance(rows[52][j], (int, float)) and re.match(r"\w+ \d{4} - \w+ \d{4}$", str(label)):
            out[str(label)] = float(rows[52][j])
    return out


def ro_treasury(period: str = "April 2026 - June 2026") -> dict[str, Line]:
    """RO paid by the Treasury per household per year at 2,500 kWh, for a cap period from April 2026."""
    ro = read_annex4_ro()
    if period not in ro:
        raise KeyError(f"UNSOURCED: Annex 4 has no RO value for {period!r}")
    bill = ro[period]                                    # 25% left on the bill, £/MWh
    treasury = bill * TREASURY_SHARE / (1 - TREASURY_SHARE)   # = 3 x bill share
    pre = ro["January 2026 - March 2026"]
    unit = "£/household/yr (nominal, ex VAT)"
    return {
        "ro_bill_share_mwh": Line(bill, "£/MWh", "allowance", SRC_A4,
                                  f"{period}, 25% left on bill (Ofgem ex-ante allowance)"),
        "ro_bill_share_per_household": Line(bill * RECEIPT_MWH, unit, "modelled", SRC_A4, "x 2.5 MWh"),
        "ro_treasury_mwh": Line(treasury, "£/MWh", "modelled", SRC_A4, "3 x bill share (75/25)"),
        "ro_treasury_per_household": Line(treasury * RECEIPT_MWH, unit, "modelled", SRC_A4,
                                          "3 x bill share, 2.5 MWh; per year, Apr 2026 to Mar 2029"),
        "ro_treasury_scheme_total_per_household": Line(treasury * RECEIPT_MWH * RO_SCHEME_YEARS,
                                                       "£/household over 3 years (nominal)", "modelled",
                                                       SRC_A4, "if the 2026/27 rate held for all 3 years"),
        "ro_full_jan_mar_2026_mwh": Line(pre, "£/MWh", "allowance", SRC_A4,
                                         "last period before the transfer (100% on bill)"),
    }


def ro_government_statements() -> dict[str, Line]:
    """The government's own RO-to-Exchequer figures, read from the stored documents."""
    g = _page_text(RO_GUIDANCE)
    m = _grab(g, r"\[£(\d+\.\d+) using CPI percentage of ([\d.]+)% for 2025\]", "RO guidance buy-out")
    buyout = float(m.group(1))
    obl = float(_grab(g, r"\[(\d\.\d+) ROCs ?/MWh in Great Britain\]", "RO guidance obligation").group(1))
    disc = float(_grab(g, r"volumetric discount for 2026/2027 of (\d\.\d+) p/ ?kWh", "RO guidance discount").group(1))
    b = _page_text(BUDGET_PAGE)
    b88 = float(_grab(b, r"funding 75% of RO through general taxation should save £(\d+)", "Budget £88").group(1))
    b_mwh = float(_grab(b, r"£(\d+\.\d+)/ ?MWh off the unit rate for 75% of RO", "Budget £/MWh").group(1))
    em = _pdf_text(RO_EM)
    m = _grab(em, r"by £(\d+) on average or £(\d+) for the typical dual fuel household", "EM para 9.3")
    em_avg, em_typ = float(m.group(1)), float(m.group(2))
    em_bn = float(_grab(em, r"£([\d.]+)bn \(nominal\) over a three-year period", "EM £bn").group(1))
    formula = TREASURY_SHARE * buyout * obl            # £/MWh
    return {
        "guidance_buyout_gbp_per_roc": Line(buyout, "£/ROC", "measured", SRC_GUIDE, "2026/27, CPI 3.4%"),
        "guidance_obligation_roc_per_mwh": Line(obl, "ROC/MWh", "measured", SRC_GUIDE, "GB 2026/27"),
        "guidance_discount_p_kwh": Line(disc, "p/kWh", "measured", SRC_GUIDE, "ex VAT, 2026/27"),
        "guidance_formula_mwh": Line(formula, "£/MWh", "modelled", SRC_GUIDE, "0.75 x a x b"),
        "guidance_per_2500kwh": Line(formula * RECEIPT_MWH, "£/household/yr", "modelled", SRC_GUIDE,
                                     "0.75 x a x b x 2.5 MWh"),
        "budget_average_household": Line(b88, "£/household/yr", "measured", SRC_BUDGET,
                                         "Budget forecast, average household (all households' "
                                         "average use), ex VAT; announced 27 Nov 2025"),
        "budget_mwh_off_unit_rate": Line(b_mwh, "£/MWh", "measured", SRC_BUDGET,
                                         "on the then-latest cap, ex VAT"),
        "budget_implied_mwh": Line(b88 / b_mwh, "MWh/household", "modelled", SRC_BUDGET,
                                   "£88 / £24.80 per MWh"),
        "em_average_household": Line(em_avg, "£/household/yr", "measured", SRC_EM,
                                     "para 9.3, average; VAT basis not stated by the EM"),
        "em_typical_dual_fuel": Line(em_typ, "£/household/yr", "measured", SRC_EM,
                                     "para 9.3, typical dual fuel (2.7 MWh)"),
        "em_implied_mwh": Line(em_avg / formula, "MWh/household", "modelled", SRC_EM,
                               "£87 / 0.75ab, IF ex VAT (the EM does not say)"),
        "em_implied_mwh_if_inc_vat": Line(em_avg / formula / (1 + VAT_DOMESTIC), "MWh/household",
                                          "modelled", SRC_EM, "£87 / 1.05 / 0.75ab, IF inc VAT"),
        "em_total_bn_3yr": Line(em_bn, "£bn (nominal)", "measured", SRC_EM, "para 9.2, three years"),
    }


def ro_em_total_per_gb_household(year: int = 2025) -> Line:
    """EM's £6.7bn over three years spread over all GB households (a cross-check only)."""
    s = ro_government_statements()
    hh, parts = gb_households(year)
    basis = "estimated" if any(p.basis == "estimated" for p in parts.values()) else "modelled"
    return Line(s["em_total_bn_3yr"].value * 1e9 / RO_SCHEME_YEARS / hh, "£/household/yr", basis, SRC_EM,
                f"£{s['em_total_bn_3yr'].value}bn / 3 / GB households {year} ({hh:,.0f})")


# --------------------------------------------------------------------------- EPG

@dataclass(frozen=True)
class EpgQuarter:
    label: str
    elec_p_kwh: float     # Treasury discount, ex VAT
    gas_p_kwh: float


_EPG_PATTERNS = (
    ("Apr-Jun 2023", r"from April to June 2023 reductions of ([\d.]+)p/ ?kWh for electricity and ([\d.]+)p/ ?kWh for gas"),
    ("Jan-Mar 2023", r"from January to March 2023 reductions of ([\d.]+)p/ ?kWh for electricity and ([\d.]+)p/ ?kWh for gas"),
    ("Oct-Dec 2022", r"from October to December 2022, unit rate discount of ([\d.]+)p/ ?kWh for electricity and ([\d.]+)p/ ?kWh for gas"),
)


def read_epg_discounts() -> dict[str, EpgQuarter]:
    """EPG unit-rate discounts per quarter (GB, ex VAT), DESNZ policy paper.

    April to June 2023 uses the paper's 3-decimal footnote (16.570p / 2.185p).
    """
    t = _page_text(EPG_PAPER)
    out = {}
    for label, pat in _EPG_PATTERNS:
        m = _grab(t, pat, f"EPG paper {label}")
        out[label] = EpgQuarter(label, float(m.group(1)), float(m.group(2)))
    m = _grab(t, r"discount rates to 3 decimal places in Great Britain were ([\d.]+)p/ ?kWh for electricity "
                 r"and ([\d.]+)p/ ?kWh for gas", "EPG footnote")
    e3, g3 = float(m.group(1)), float(m.group(2))
    q = out["Apr-Jun 2023"]
    if round(e3, 1) != q.elec_p_kwh or round(g3, 1) != q.gas_p_kwh:
        raise ValueError("EPG footnote does not match the April-June text")
    out["Apr-Jun 2023"] = EpgQuarter(q.label, e3, g3)
    return out


def read_epg_regional_dd(label: str) -> tuple[float, float]:
    """GB unweighted mean of the 14 regional direct-debit EPG unit rates (elec, gas), p/kWh ex VAT."""
    t = _page_text(EPG_RATES[label])
    head = "gas (p/kwh ex vat)"
    i = t.lower().find(head)
    if i < 0:
        raise ValueError(f"EPG regional rates {label}: table header missing")
    j = t.find("Standard credit", i)
    if j < 0:
        raise ValueError(f"EPG regional rates {label}: 'Standard credit' table missing")
    body = t[i + len(head):j]
    regions = [p for p in re.findall(r"([A-Z][A-Za-z ]+?) (\d+\.\d+) (\d+\.\d+)", body) if "GB" not in p[0]]
    if len(regions) != 14:
        raise ValueError(f"EPG regional rates {label}: expected 14 regions, found {len(regions)}")
    return (sum(float(r[1]) for r in regions) / 14, sum(float(r[2]) for r in regions) / 14)


def read_cap_unit_rates() -> dict[str, tuple[float, float]]:
    """Ofgem cap unit rates for the EPG quarters, GB average, Other Payment Method, p/kWh ex VAT.

    (typical - nil) / benchmark consumption, from cap model v1.31 tab '1b Historical level tables'
    (3.1 MWh electricity, 12 MWh gas then in force: Notes rows 163-165).
    """
    wb = openpyxl.load_workbook(CAP_MODEL, read_only=True, data_only=True)
    R = list(wb["1b Historical level tables"].iter_rows(values_only=True))
    if (R[7][1], R[9][1], R[11][1], R[29][1], R[26][1], R[44][1]) != (
            "Other Payment Method", "Electricity: Single-Rate Metering Arrangement", "Nil consumption",
            "Typical consumption", "Total_GB average", "Total_GB average"):
        raise ValueError("cap model 1b layout moved")
    gas_col = next(j for j, v in enumerate(R[9]) if v == "Gas")
    N = list(wb["Notes"].iter_rows(values_only=True))
    if not str(N[163][1]).startswith("Electricity - Single-Rate") or N[165][1] != "Gas":
        raise ValueError("cap model Notes benchmark table moved")
    e_mwh, g_mwh = float(N[163][2]), float(N[165][2])
    want = {"Oct-Dec 2022": "Oct 2022 - Dec 2022", "Jan-Mar 2023": "Jan 2023 - Mar 2023",
            "Apr-Jun 2023": "Apr 2023 - June 2023"}
    out = {}
    for k, lab in want.items():
        je = next(j for j, v in enumerate(R[11][:gas_col]) if v == lab)
        jg = next(j for j, v in enumerate(R[11]) if v == lab and j >= gas_col)
        e = (R[44][je] - R[26][je]) / e_mwh / 10
        g = (R[44][jg] - R[26][jg]) / g_mwh / 10
        out[k] = (e, g)
    return out


def epg_note(mwh: float = RECEIPT_MWH) -> dict[str, dict[str, Line]]:
    """Treasury cost per household per quarter under the EPG (note only, not in the receipt)."""
    disc = read_epg_discounts()
    out = {}
    for label, q in disc.items():
        quarter = {
            "elec_discount_p_kwh": Line(q.elec_p_kwh, "p/kWh ex VAT", "measured", SRC_EPG, label),
            "gas_discount_p_kwh": Line(q.gas_p_kwh, "p/kWh ex VAT", "measured", SRC_EPG, label),
            "elec_per_household_quarter": Line(q.elec_p_kwh * mwh * 10 / 4, "£/household/quarter",
                                               "modelled", SRC_EPG,
                                               f"{mwh} MWh/yr, a quarter = 1/4 of a year (no seasonality)"),
            "elec_per_household_annualised": Line(q.elec_p_kwh * mwh * 10, "£/household/yr", "modelled",
                                                  SRC_EPG, f"{mwh} MWh/yr at this quarter's rate"),
            "gas_per_household_quarter_12mwh": Line(q.gas_p_kwh * 12 * 10 / 4, "£/household/quarter",
                                                    "modelled", SRC_EPG, "Ofgem's 12 MWh gas then in force"),
        }
        out[label] = quarter
    return out


def read_nao_epg_total() -> Line:
    t = _pdf_text(NAO)
    m = _grab(t, r"Energy Price Guarantee \(EPG\) 1 Oct 2022–Mar 2024 in Great Britain \(GB\) "
                 r"Nov 2022–Mar 2024 in Northern Ireland \(NI\) £([\d,]+) million", "NAO EPG total")
    return Line(float(m.group(1).replace(",", "")), "£m", "measured", SRC_NAO,
                "EPG total, GB and NI, Oct 2022 to Mar 2024 (domestic)")
