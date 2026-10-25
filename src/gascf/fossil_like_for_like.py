"""Like-for-like fossil Britain: the actual system with net zero taken out (strand s22).

Plain words. The old step-3 method (`fleet="all_new"`, kept as a switch) priced fossil
Britain's whole coal-and-gas fleet as if built new, while net-zero Britain paid the market
price of a real fleet of old, largely paid-off plant. That charges one world for plant the
other never bought. Here fossil Britain is the actual system less net zero:

1. **The plant Britain had** (coal, gas and oil stations in GB, DUKES 5.7, closures on
   their actual dates) runs as it actually ran and is paid the actual wholesale price
   **less the carbon in it** (spec §6, the same price GO1k uses).
2. **Power renewables made** is replaced. Renewables replaced, all inside the volume V
   both bills are priced on: transmission wind and biomass (NESO historic generation mix,
   2009 on; NESO's "other" column counted as biomass in half-hours before NESO split biomass
   out in 2017; DUKES 5.6 wind spread flat before 2009), plus embedded wind and embedded
   solar (NESO Historic Demand Data, the same columns that make V). Natural-flow hydro is
   kept (mostly pre-1990 schemes Britain had anyway), priced like the existing plant. First the existing fossil plant runs harder where it had spare
   capacity that half-hour (at its own fuel and variable cost, no carbon: DUKES 5.10 fleet
   efficiencies, DESNZ QEP / Ofgem fuel prices, and the edition's variable O&M); where the
   existing plant runs out, **new** CCGT/OCGT is built and priced new (edition in force,
   `fossil_fleet.cost_fossil_year` on the shortfall).
3. **Firm capacity:** if existing plant (de-rated) plus the new plant falls short of the
   residual peak × (1 + margin), the shortfall is bought as new OCGT, fixed cost only.
4. **Merit-order uplift:** without renewables the marginal plant is dearer, so the price
   on every MWh sold at the market price (existing plant's actual and extra output, and the
   kept nuclear, imports and pumped storage) rises. Modelled from this module's own no-carbon
   supply stack (Δ = SRMC of the marginal plant at fossil Britain's load − at the actual
   fossil load, each half-hour), cross-checked against Shao et al. (2022), the only
   peer-reviewed GB year-by-year estimate found (£0.61-£1.83/MWh per GW of wind, Table 5;
   £2.98/MWh per GW of solar, Table 3). Central = the larger of the two each year (the
   conservative side: it raises fossil Britain's bill and shrinks the cost of net zero);
   the smaller is the low switch. New plant is paid its full cost, not the uplift.

**s23b, spec §15.3 (binding): fossil Britain contains no net-zero policy of any kind.**

5. **No merit-order uplift** in the central case (`moe="none"`); the old rules are kept as
   switches for the methods page only (never in the low-high range).
6. **Closures for legitimate reasons only.** `read_closures()` reads the closure register
   (`reference/fossil_closures.yaml`, strand s23a). A station whose `fossil_britain_close` is
   null (still running at end-2025) or later than its actual close stays in fossil Britain's
   existing fleet for the part of each year between the two dates, at its real capacity, the
   DUKES fleet efficiency for its fuel, coal at QEP 3.2.1, and the running costs of the edition
   in force (`fossil_fleet.keep_open_gbp_kw_yr`) plus the annualised cost of the pollution
   filters it needed (`fossil_fleet.abatement_gbp_kw_yr`). Switch `closures="actual"`: every
   closure on its actual date (the s22 rule).
7. **No ban on new coal.** Where existing plant runs out, the new plant is the cheapest of
   CCGT, OCGT and coal with FGD and SCR, strip by strip, at the cost edition in force; nuclear
   only where the edition prints it and not before its earliest commissioning year
   (`fossil_fleet.nuclear_available`; never the case 2002-2025). Switch `newbuild="gas_only"`
   is the s22 rule.

**s27a, spec §15.5 (binding, 12 Oct 2026; overrides §15.3 item 1):**

8. **Merit-order uplift back in the central case** (`moe="shao"`): Shao et al. (2022) on the
   market-priced MWh (actual fossil output, the extra running of the same plant from v4-s37 (0144 item 1A),
   and kept nuclear, imports, pumped storage). The engine's own stack estimate (`moe="model"`) is a
   methods-page sensitivity from v4-s37 (0144 item 4A); `"none"` (the §15.3 rule) is a methods-page sensitivity.
9. **Gas-price feedback** (`gas_feedback()`, rules central/krichene_lr/high/none): fossil Britain
   burns more gas, so gas is dearer by dQ / ((e_s + e_d) x Q) (0144 item 2A: central Albrizio's pair,
   0.06 + 0.169), applied to at-cost gas fuel and to market-priced MWh x gas-marginal share / CCGT efficiency.
10. **Coal capped at 65% load factor** (`coal_lf="dukes_max"`, DUKES 5.10.B best recorded
   coal-fleet year, read from the file): part of the coal block is moved to the top of the
   stack (it runs only at the peak) until
   its yearly output meets the cap; the rest goes to the next blocks and new build.
   `coal_lf="none"` is the s23b rule.
11. **No silent fallbacks:** every fallback is declared and recorded in `gascf.fallbacks`;
   coal efficiency 2025 (DUKES prints 0) = the 2015-2023 fleet mean.

Units: MW per half-hour, TWh, £m nominal (money of the day).
"""
from __future__ import annotations

import bisect
import csv
import json
import os
import re
from dataclasses import dataclass, replace
from datetime import date, datetime, timedelta
from functools import lru_cache
from pathlib import Path

import openpyxl

from gascf import capacity, prices
from gascf import fallbacks as fb
from gascf import fossil_fleet as ff
from gascf import residual_neso as rn

ROOT = Path(__file__).parents[2]
RAW = ROOT / "data/raw"
NESO_MIX = RAW / "neso_s22/s22-df_fuel_ckan.csv"        # NESO Historic generation mix (CKAN)
SHAO_PDF = RAW / "papers/s22-Shao_etal_2022_RO_wind_merit_order.pdf"
DUKES_57 = ff.DUKES_57
DUKES_511 = ff.DUKES_511

SRC_MIX = "neso-historic-generation-mix-df_fuel_ckan"
SRC_SHAO = "shao-chen-li-liu-2022-ro-wind-merit-order"
SRC_57 = "desnz-dukes-5.7"
SRC_511 = "desnz-dukes-5.11"

MOE_RULES = ("larger", "smaller", "model", "shao", "none")
DEFAULT_MOE = "shao"                 # spec §15.5 item 1 (Richard, 12 Oct 2026, binding; overrides §15.3)
MOE_HIGH = "model"                   # the engine's own supply-stack estimate: high-side sensitivity
MOE_METHODS = "none"                 # §15.3 rule, kept for the methods page

# s27a, spec §15.5 item 3: kept coal capped at the best recorded GB coal-fleet load factor.
# Card [internal ticket] (built OFF; decision 0166 makes 'dukes_max+ied_permits_1500+existing' central in v4-s39, the 'route'
# reading a check row only): "dukes_max+ied_<reading>_<hours>" also
# holds each kept coal unit in reference/ied_hours_caps.json to an IED hours cap from 1 July 2020
# (reading "permits" = the stations' real permits; "route" = also Drax 1-4 and Lynemouth; hours 1500 or 2250).
IED_HOURS_RULES = tuple(f"dukes_max+ied_{r}_{h}" for r in ("permits", "route") for h in (1500, 2250))
# Card [internal ticket] (built OFF; central in v4-s39 under 0166): "<kept rule>+existing" also holds existing coal
# to its real caps (the file's "existing" part): the LCPD opt-outs' 20,000 hours from 2008, and West Burton A and
# Drax 5-6 to their permits from 1 July 2020 while still open in reality.
IED_KEPT_RULES = IED_HOURS_RULES
IED_HOURS_RULES += tuple(f"{r}+existing" for r in IED_HOURS_RULES)
# Card [internal ticket] (measure only, OFF; not used under 0166, which rules SCR-instead out): a suffix "+ex_cap_<hours>" holds the
# EXISTING coal stations in the same file's "existing_jul2020" list (West Burton A, Drax 5-6) to the hours cap from
# 1 July 2020 while they are in the existing block; "+ex_scr" charges them SCR instead (no cap), costed as kept SCR.
# rb-v4-s39-rc merge: the suffix goes on dukes_max or a kept-coal rule, not on a '+existing' rule (the two cards'
# existing-coal treatments are alternatives, never stacked).
IED_EXISTING_SUFFIXES = ("+ex_cap_1500", "+ex_cap_2250", "+ex_scr")
# Card [internal ticket] (built OFF; ON by default in v4-s39 under 0166, see below): under a '+existing' rule, $GASCF_IED_LCPD_KEPT=on
# also holds the LCPD opt-outs that M2's closure test KEEPS past their real closes (Tilbury B, Ironbridge) to the
# 20,000 hours (`ied_existing_units`).
LCPD_KEPT_ENV = "GASCF_IED_LCPD_KEPT"
# v4-s39 central (card [internal ticket]; decision 0166, Richard Lyon 22 Oct 2026: "As long as it was lawful, then it's
# legitimate. Approved."): ON when the variable is unset; "off" turns it off (for reruns of the older line).


def lcpd_kept_on() -> bool:
    """Kept LCPD opt-outs held to their hours: True unless $GASCF_IED_LCPD_KEPT is "off" (0166)."""
    return os.environ.get(LCPD_KEPT_ENV, "on").strip() != "off"


# v4-s40 item 1 (step-3 sheet item 1 (a); decision 0175, Richard Lyon 23 Oct 2026, "Yes." at 06:13:51 UTC to "The coal
# fix caps coal at its legal energy limit, without building a station-by-station timetable."; verified on card
# [internal ticket]). To v4-s39 the s29 actual-output floor overrode the permit-derived energy limit (2012-13 breaches,
# review [internal ticket] item 1). "absolute" (central, unset): where a permit-derived limit is supplied it binds; the floor
# stays only under the DUKES load-factor cap, which is not a legal limit. "floored" = v4-s39 (explicit only). Set the
# variable before the process starts (lfl_year is cached).
COAL_CEILING_ENV = "GASCF_COAL_CEILING"
COAL_CEILING_READINGS = ("absolute", "floored")


def coal_ceiling_reading() -> str:
    v = (os.environ.get(COAL_CEILING_ENV) or COAL_CEILING_READINGS[0]).strip()
    if v not in COAL_CEILING_READINGS:
        raise ValueError(f"${COAL_CEILING_ENV}={v!r}: one of {COAL_CEILING_READINGS}")
    return v
COAL_LF_RULES = (("dukes_max", "none") + IED_HOURS_RULES
                 + tuple(b + s for b in ("dukes_max",) + IED_KEPT_RULES for s in IED_EXISTING_SUFFIXES))
DEFAULT_COAL_LF = "dukes_max"        # max of DUKES 5.10.B "of which coal-fired stations" (65.04%, 2003)
IED_HOURS_FILE = Path(__file__).parents[2] / "reference/ied_hours_caps.json"
SRC_510 = "desnz-dukes-5.10"

# s27a, spec §15.5 item 2: gas-price feedback. v4-s37, decision 0144 item 2A (Richard Lyon, 20 Oct 2026: "Yes,
# take all the recommendations"; supersedes 0033's central 0.16). When one buyer wants more gas, the price rises
# until suppliers sell more AND other buyers cut back, so dP/P = dQ / ((e_s + e_d) Q), Albrizio et al. (2022),
# IMF WP/22/143, equation 5. One source and one market for each pair (GAS_ELASTICITY_PAIRS, (supply, demand)):
# - central (0.06, 0.169) = 0.229, Albrizio's own pair, 12-month horizon, European market: supply 0.06 = Krichene
#   (2005), Table 5 (printed p.20, PDF p.21), natural gas supply, short-run price elasticity, SEM with the NEER,
#   1974-2004, as Albrizio use it (§2.4), the short-run worst case (12 months, no new investment); demand 0.169 =
#   Albrizio Table 2, sector demand elasticities weighted by EU consumption shares (industry 0.084 x 23%,
#   households 0.24 x 34%, power 0.136 x 32%, other 0.22 x 11% = 0.1686).
# - krichene_lr (0.16, 0.45) = 0.61, the sensitivity row (0144 item 2, option B): Krichene (2005) Table 2 (PDF
#   p.13), world natural gas, long-run, VECM, 1918-2004: supply 0.16, demand -0.45. A world pair applied to the
#   European Q (Neuhoff's 4,000 TWh plus GB), as in the sheet.
# Direction: a LARGER elasticity gives a smaller price rise, so a smaller Reliable Britain bill and a larger gap;
# a smaller elasticity gives a bigger feedback (a smaller gap). (v4-s36's comment here had that the wrong way round.)
# - high 0.06, supply only, NO demand side: kept only because case B (decision 0048, `case_b.CASES`) sets it; what
#   case B's "0.06" means under the paired formula is open (0144 note 2, put to Richard). Not in any range row.
# Engine names: "central", "krichene_lr", "high" (case B only), "none". v4-s36's "low" (0.34, supply only) is gone
# with the supply-only formula (0144 item 4A: feedback rows rebuilt with paired elasticities).
# Q = EU historic gas demand 4,000 TWh (Neuhoff 2022 Table 2) + GB gas demand (Energy Trends 4.1
# 'Total demand', by year). dQ = fossil Britain's gas burn for power less the actual (DUKES 5.6.A
# MPP gas fuel used), floored at 0.
# - krichene_lr_1974 (0.21, 0.06) = 0.27, RB v1.1 A5 (J3; decision 0160, card [internal ticket]), a held row for ruling R14,
#   never in the range: Krichene (2005) Table 2 (PDF p.13), VECM panel, 1974-2004: supply 0.21 (t 6.51), demand -0.06
#   (t -1.67, not significant at 5%). Confirmed by the analyst (card [internal ticket]) from the same table as krichene_lr.
GAS_FEEDBACK_RULES = ("central", "krichene_lr", "high", "none", "krichene_lr_1974")
DEFAULT_GAS_FEEDBACK = "central"
GAS_ELASTICITY_PAIRS = {"central": (0.06, 0.169), "krichene_lr": (0.16, 0.45), "high": (0.06, 0.0),
                        "krichene_lr_1974": (0.21, 0.06)}
GAS_ELASTICITY = {k: s + d for k, (s, d) in GAS_ELASTICITY_PAIRS.items()}   # the e in dQ / (e Q): supply + demand
KRICHENE_2005_PDF = RAW / "s36/s36-krichene2005_IMF_WP0532.pdf"
CASOLI_2023_PDF = RAW / "s36/s36-casoli_manera_valenti_FEEM_NDL2022-045.pdf"
ALBRIZIO_2022_TXT = RAW / "s36/s36-albrizio2022_imf_wp22-143.txt"
EU_GAS_DEMAND_TWH = 4000.0
NEUHOFF_PDF = RAW / "s26a/s26a-Neuhoff_2022_EPRG2212_gas_price_limits.pdf"
ET_41 = RAW / "s26a/s26a-ET_4.1_SEP_26.xlsx"
SRC_NEUHOFF = "neuhoff-2022-eprg2212-gas-price-limits"
SRC_ET41 = "desnz-energy-trends-4.1"
SRC_56 = "desnz-dukes-5.6"

CLOSURES_FILE = ROOT / "reference/fossil_closures.yaml"
CLOSURES_ENV = "GASCF_FOSSIL_CLOSURES"     # path override (tests and the fixture run)
CLOSURE_RULES = ("legitimate", "actual")
DEFAULT_CLOSURES = "legitimate"            # spec §15.3 items 2 and 5
NEWBUILD_RULES = ("cheapest", "gas_only")
DEFAULT_NEWBUILD = "cheapest"              # spec §15.3 item 4
CLOSURE_REASONS = ("age", "air_pollution", "net_zero_policy", "mixed", "other")
# RB v1.1 O1 (decision 0163 R17 (a)): closure classes, register data on the net_zero_policy and mixed rows (the
# net_zero_policy rows' classes are R34 (a), ruled in 0164); a class row closes that class on its real dates
CLOSURE_CLASSES = ("biomass_conversion", "coal_phase_out_deadline", "mixed")
CLOSURE_FUELS = ("coal", "gas", "oil")
SRC_CLOSURES = "s23a-fossil-closures-register"

# Shao, Chen, Li & Liu (2022), "An evaluation of the consumer-funded renewable obligation
# scheme in the UK for wind power generation", Swansea University (Cronfa). Table 5,
# Prais-Winsten by reporting year (April-March), daily average hourly GB price (£/MWh)
# on daily average hourly wind generation (GW). Key = first calendar year of the FY.
SHAO_WIND_FY = {2009: -1.39, 2010: -1.57, 2011: -1.83, 2012: -1.71, 2013: -1.57, 2014: -1.22,
                2015: -1.19, 2016: -1.65, 2017: -0.61, 2018: -1.05, 2019: -0.98, 2020: -1.77}
SHAO_WIND_POOLED = -1.28          # Table 3 / abstract, April 2009 - March 2021
SHAO_SOLAR_POOLED = -2.98         # Table 3 (year-by-year solar terms unstable, not used)

# Existing-plant efficiencies not in DUKES 5.10 (estimated): gas turbines and gas/oil steam.
ETA_OTHER_GAS = 0.30
ETA_OIL = 0.30


@lru_cache(maxsize=None)
def edition(year: int, cost_case: str = "medium"):
    """`fossil_fleet.edition`, cached here (it re-reads the cost workbooks on every call). cost_case: RB v1.1 O2."""
    return ff.edition(year, cost_case)


@lru_cache(maxsize=None)
def gas_price(year: int, month: int) -> float:
    return ff.gas_price_month(year, month)[0]


@lru_cache(maxsize=None)
def coal_price(year: int, month: int) -> float:
    return ff.coal_price_month(year, month)[0]


def shao_wind_coeff(year: int) -> tuple[float, str]:
    """£/MWh per GW for calendar year y: ¼ FY(y−1) + ¾ FY(y); pooled outside the sample."""
    a = SHAO_WIND_FY.get(year - 1, SHAO_WIND_POOLED)
    b = SHAO_WIND_FY.get(year, SHAO_WIND_POOLED)
    inside = year - 1 in SHAO_WIND_FY and year in SHAO_WIND_FY
    if not inside:
        fb.record("shao_wind_pooled", year, 0.25 * a + 0.75 * b, "£/MWh per GW",
                  "pooled coefficient for the out-of-sample part of the year")
    return 0.25 * a + 0.75 * b, "in sample" if inside else "pooled / partly out of sample"


# ----------------------------------------------------------------- existing plant (GB)

def _row57(table: str, gen: str, fuel: str) -> dict[int, float]:
    rows = list(openpyxl.load_workbook(DUKES_57, read_only=True, data_only=True)["5.7"].iter_rows(values_only=True))
    t = next(i for i, r in enumerate(rows) if r[0] and str(r[0]).startswith(table))
    hdr = rows[t + 1]
    r = next(r for r in rows[t + 2:t + 45] if r[0] and str(r[0]).startswith(gen)
             and r[1] and str(r[1]).startswith(fuel))
    return {int(h): float(v) for h, v in zip(hdr[2:], r[2:])
            if h and str(h).isdigit() and isinstance(v, (int, float))}


# Decision 0197 (Richard Lyon, 24 Oct 2026, msg 8982: "each year's own list"; card [internal ticket], v4-s43): the NI rows
# of DUKES 2002 and DUKES 2003 Table 5.10, entered from the archived editions (data/raw/dukes_s43/, checked cell by
# cell by work/[internal ticket]/check_ni_rows.py), and the NI stations of the 'DUKES 2004' and 'DUKES 2005' sheets of
# DUKES 5.11, which print no region column and are read by station name.
NI_LISTS_2002_2003 = RAW / "dukes_s43/dukes_5_10_ni_2002_2003.csv"
NI_STATIONS_BY_NAME = {"DUKES 2004": ("Kilroot", "Coolkeeragh", "Ballylumford"),
                       "DUKES 2005": ("Kilroot", "Coolkeeragh", "Ballylumford")}


def _ni_row(cells: list[str], mw: float | None) -> str | None:
    """Class one NI station row as its list prints it: None (not fossil, or Kilroot's coal, which is already out
    of coal_capacity_gb), 'ccgt' (labelled CCGT) or 'other' (any other fossil label, e.g. 'oil', 'oil/gas')."""
    low = " ".join(cells).lower()
    if any(k in low for k in ("wind", "solar", "hydro", "biomass", "landfill", "anaerobic",
                              "waste", "battery", "storage", "(mw)")):
        return None
    if "kilroot" in low and "coal" in low:
        return None                                        # NI coal: already out of coal_capacity_gb
    if mw is None:
        return None
    return "ccgt" if "ccgt" in low else "other"


@lru_cache(maxsize=None)
def ni_fossil_mw() -> dict[int, dict[str, float]]:
    """Northern Ireland fossil stations other than Kilroot's coal units, MW, by DUKES edition:
    {'ccgt': .., 'other': ..}, each year from its own edition's list, classed as that list prints it
    (decision 0197). DUKES 2002 and DUKES 2003 come from their Table 5.10, entered from the archived
    editions (data/raw/dukes_s43/dukes_5_10_ni_2002_2003.csv); DUKES 2004 and DUKES 2005 from their
    DUKES 5.11 sheets, read by station name as those sheets print no region column; DUKES 2006 on from
    the sheets' region column. Only a year before 2002 (the first list held) would take the nearest
    list, recorded as a fallback by `existing_plant`."""
    wb = openpyxl.load_workbook(DUKES_511, read_only=True, data_only=True)
    out = {}
    with NI_LISTS_2002_2003.open(newline="") as f:
        for row in csv.DictReader(f):
            y = int(row["edition"].removeprefix("DUKES "))
            acc = out.setdefault(y, {"ccgt": 0.0, "other": 0.0})
            mw = float(row["installed_capacity_mw"])
            k = _ni_row([row["company"], row["station"], row["fuel"]], mw)
            if k:
                acc[k] += mw
    for sn in wb.sheetnames:
        m = re.fullmatch(r"DUKES (\d{4})", sn)
        if not m:
            continue
        ccgt = other = 0.0
        for r in wb[sn].iter_rows(values_only=True):
            cells = [str(x).strip() for x in r if isinstance(x, str)]
            named = any(c in NI_STATIONS_BY_NAME.get(sn, ()) for c in cells)
            # A2 (input audit v4-s41, v4-s42, decision 0195): 'DUKES 2006' labels Coolkeeragh 'N. Ireland'
            if not named and not any((c.startswith("Northern Ireland") and len(c) < 20) or c == "N. Ireland"
                                     for c in cells):
                continue
            nums = [x for x in r if isinstance(x, (int, float)) and 5 < x < 3000]
            k = _ni_row(cells, nums[0] if nums else None)
            if k == "ccgt":
                ccgt += nums[0]
            elif k == "other":
                other += nums[0]
        # Input audit v4-s41 A1 (v4-s42, decision 0195): Kilroot's OCGT printed inside its coal/oil row
        # (DUKES 2010, 2011; split out in DUKES 2012) is NI 'other', not NI coal.
        other += ff.kilroot_split().get(int(m.group(1)), 0.0)
        if ccgt or other:
            out[int(m.group(1))] = {"ccgt": ccgt, "other": other}
    return dict(sorted(out.items()))


@dataclass(frozen=True)
class Existing:
    """Existing GB fossil plant for one year, MW installed (end of year, DUKES 5.7)."""
    coal: float
    ccgt: float
    other_gas: float
    oil: float
    basis: str

    def derated(self, der: dict[str, float]) -> dict[str, float]:
        return {"coal": self.coal * der["ccgt"], "ccgt": self.ccgt * der["ccgt"],
                "other_gas": self.other_gas * der["ocgt"], "oil": self.oil * der["ocgt"]}


@lru_cache(maxsize=None)
def existing_plant(year: int) -> Existing:
    """Coal: `coal_capacity_gb` (all generators, less Kilroot), as the old method. Gas and oil:
    major power producers only (5.7.A 'Gas fired', 'Oil fired', 'Mixed or dual fired'; CCGT
    from 5.7.B 'CCGT stations'), less Northern Ireland. Other generators' gas plant (mostly
    industrial CHP serving its own site) is left out: the conservative side (less existing
    plant, more new build)."""
    gas = _row57("Table 5.7.A", "Major power producers", "Gas fired")[year]
    oil = _row57("Table 5.7.A", "Major power producers", "Oil fired")[year]
    mixed = _row57("Table 5.7.A", "Major power producers", "Mixed or dual")[year]
    ccgt = _row57("Table 5.7.B", "Major power producers", "CCGT stations")[year]
    ni = ni_fossil_mw()
    nik = year if year in ni else (min(ni) if year < min(ni) else max(k for k in ni if k <= year))
    if nik != year:
        fb.record("ni_fossil_nearest_year", year, sum(ni[nik].values()), "MW", f"NI fossil MW of {nik}")
    coal, cb = ff.coal_capacity_gb(year)
    basis = "measured" if year in ni and cb == "measured" else "estimated"
    return Existing(coal=coal, ccgt=max(0.0, ccgt - ni[nik]["ccgt"]),
                    other_gas=max(0.0, gas - ccgt + mixed - ni[nik]["other"] * 0.5),
                    oil=max(0.0, oil - ni[nik]["other"] * 0.5), basis=basis)


# ----------------------------------------------------------------- s23b: the closure register

def closures_path() -> Path:
    """The register in use: `$GASCF_FOSSIL_CLOSURES` if set (tests, fixture runs), else
    `reference/fossil_closures.yaml` (strand s23a)."""
    return Path(os.environ.get(CLOSURES_ENV) or CLOSURES_FILE)


# Card [internal ticket] (measure only; Richard has not ruled): one switch per station for the Rugeley rule
# ("use the capacity in the DUKES 5.11 edition sheet listing last before closure") on five closures.
# All OFF by default. Turn on by naming switches in $GASCF_EDITION_MW (comma-separated), for example
# "barry,fawley,ironbridge_940,deeside,uskmouth_230". The register file is never edited by a switch.
EDITION_MW_FILE = ROOT / "reference/closures_edition_mw.yaml"
EDITION_MW_ENV = "GASCF_EDITION_MW"
# v4-s39 central (card [internal ticket]; corrections log C-009 under the Rugeley rule, 0159): the five switches ON
# when $GASCF_EDITION_MW is unset. Ironbridge (900 or 940) and Uskmouth (230 or 363) are still open for
# Richard (0164 item 3); set as C-009 has them. "none" reads the register as stored (the v4-s38 MW).
EDITION_MW_CENTRAL = "barry,deeside,fawley,ironbridge_900,uskmouth_230"
EDITION_MW_OFF = "none"


@lru_cache(maxsize=None)
def _edition_mw_table() -> dict[str, dict]:
    return {r["switch"]: r for r in json.loads(EDITION_MW_FILE.read_text())}


def edition_mw_switches() -> tuple[str, ...]:
    """The switches turned on, sorted: the central five when $GASCF_EDITION_MW is unset, () when it is
    "none" or empty."""
    raw = os.environ.get(EDITION_MW_ENV)
    if raw is None:
        raw = EDITION_MW_CENTRAL
    if raw.strip() == EDITION_MW_OFF:
        raw = ""
    sw = tuple(sorted({s.strip() for s in raw.split(",") if s.strip()}))
    table = _edition_mw_table()
    bad = [s for s in sw if s not in table]
    if bad:
        raise ValueError(f"{EDITION_MW_ENV}: unknown switch {bad}; one of {sorted(table)}")
    seen: dict[str, str] = {}
    for s in sw:
        st = table[s]["station"]
        if st in seen:
            raise ValueError(f"{EDITION_MW_ENV}: two switches for {st} ({seen[st]}, {s}); turn on one")
        seen[st] = s
    return sw


def _edition_mw_overrides(sw: tuple[str, ...]) -> dict[tuple[str, str], dict]:
    t = _edition_mw_table()
    return {(t[s]["station"], t[s]["unit"]): t[s] for s in sw}


def switched_mw(station: str, mw: float) -> float:
    """The MW to use for `station` (unit 'all') under the switches on: the edition MW if its
    switch is on, else `mw`. For scripts that type a station's MW by hand (the s24 margin script)."""
    o = _edition_mw_overrides(edition_mw_switches()).get((station, "all"))
    if o is None:
        return mw
    if mw != o["register_mw"]:
        raise ValueError(f"{EDITION_MW_ENV} {o['switch']}: {station} is {mw} MW here, the switch was "
                         f"written against {o['register_mw']} MW")
    return o["edition_mw"]


def _date(v, field: str, station: str) -> date | None:
    if v is None:
        return None
    try:
        return date.fromisoformat(str(v))
    except ValueError as e:
        raise ValueError(f"closure register: {station!r} {field} {v!r} is not an ISO date") from e


@lru_cache(maxsize=None)
def _read_closures(path: str, switches: tuple[str, ...] = ()) -> tuple[dict, ...]:
    p = Path(path)
    over = _edition_mw_overrides(switches)
    if not p.exists():
        raise FileNotFoundError(f"closure register {p} not found (strand s23a builds it; "
                                f"set {CLOSURES_ENV} to a fixture to run without it)")
    rows = json.loads(p.read_text())
    if not isinstance(rows, list) or not rows:
        raise ValueError(f"closure register {p}: expected a non-empty JSON list")
    out = []
    need = ("station", "unit", "fuel", "capacity_mw", "actual_close", "reason", "fossil_britain_close",
            "abatement_needed", "source_id", "url", "quote")
    for r in rows:
        miss = [k for k in need if k not in r]
        if miss:
            raise ValueError(f"closure register: {r.get('station')!r} lacks {miss}")
        st = f"{r['station']} / {r['unit']}"
        if r["fuel"] not in CLOSURE_FUELS:
            raise ValueError(f"closure register: {st} fuel {r['fuel']!r} not in {CLOSURE_FUELS}")
        if r["reason"] not in CLOSURE_REASONS:
            raise ValueError(f"closure register: {st} reason {r['reason']!r} not in {CLOSURE_REASONS}")
        if r["abatement_needed"] not in (None,) + ff.ABATEMENT:
            raise ValueError(f"closure register: {st} abatement {r['abatement_needed']!r}")
        act = _date(r["actual_close"], "actual_close", st)
        fb = _date(r["fossil_britain_close"], "fossil_britain_close", st)
        if act is None:
            raise ValueError(f"closure register: {st} has no actual_close")
        if r["reason"] in ("age", "air_pollution") and fb != act:
            raise ValueError(f"closure register: {st} closed for {r['reason']} but fossil-Britain date "
                             f"{fb} differs from actual {act} (brief rule)")
        if r["reason"] == "net_zero_policy" and fb is not None:
            raise ValueError(f"closure register: {st} closed for net-zero policy must stay open (null)")
        if fb is not None and fb < act:
            raise ValueError(f"closure register: {st} closes earlier in fossil Britain than it did")
        if r["capacity_mw"] is None and fb == act:
            # closes on the same date in both worlds: its capacity is never used here
            out.append({**r, "capacity_mw": 0.0, "actual_close_d": act, "fossil_britain_close_d": fb})
            continue
        if not isinstance(r["capacity_mw"], (int, float)):
            raise ValueError(f"closure register: {st} capacity {r['capacity_mw']!r} is not a number")
        mw = float(r["capacity_mw"])
        o = over.pop((r["station"], r["unit"]), None)
        if o is not None:
            if mw != o["register_mw"]:
                raise ValueError(f"{EDITION_MW_ENV} {o['switch']}: register {st} is {mw} MW, the switch was "
                                 f"written against {o['register_mw']} MW")
            mw = float(o["edition_mw"])
        if not 0 < mw < 5000:
            raise ValueError(f"closure register: {st} capacity {mw} MW out of range")
        out.append({**r, "capacity_mw": mw, "actual_close_d": act, "fossil_britain_close_d": fb})
    if over:
        raise ValueError(f"{EDITION_MW_ENV}: no register row for {sorted(over)} in {p}")
    return tuple(out)


def edition_mw_switches_for(path: str | Path) -> tuple[str, ...]:
    """The switches for the register at `path`. The central default (variable unset) belongs to the real register
    only: a fixture register read with the variable unset gets none (card [internal ticket]). A switch named in
    $GASCF_EDITION_MW applies to any register, strictly, as before."""
    if os.environ.get(EDITION_MW_ENV) is None and Path(path).resolve() != Path(CLOSURES_FILE).resolve():
        return ()
    return edition_mw_switches()


def read_closures(path: str | Path | None = None) -> tuple[dict, ...]:
    """The closure register (spec §15.3 item 2), schema-checked. One function, one place.
    Edition-MW switches ($GASCF_EDITION_MW; v4-s39: the C-009 five on the real register by default) apply here."""
    p = str(path or closures_path())
    return _read_closures(p, edition_mw_switches_for(p))


def _open_share(year: int, start: date, end: date | None) -> float:
    """Share of calendar year `year` after `start` (exclusive) and up to `end` (None = open)."""
    y0, y1 = date(year, 1, 1), date(year + 1, 1, 1)
    a = max(y0, start)
    b = min(y1, end) if end is not None else y1
    return max(0.0, (b - a).days) / (y1 - y0).days


def stack_kind(r: dict) -> str:
    """Which existing-plant block a kept station joins: coal; oil; gas stations as CCGT unless
    the register names an open-cycle or gas-turbine plant (then 'other_gas')."""
    if r["fuel"] == "coal":
        return "coal"
    if r["fuel"] == "oil":
        return "oil"
    txt = f"{r['station']} {r.get('unit', '')} {r.get('note', '')}".lower()
    return "other_gas" if ("ocgt" in txt or "gas turbine" in txt or "open cycle" in txt) else "ccgt"


def kept_open(year: int, closures: str = DEFAULT_CLOSURES, path: str | None = None, overrides: tuple = (),
              real_date_class: str = "none") -> dict:
    """Plant fossil Britain keeps that Britain actually closed, for calendar year `year`.

    MW are time-weighted over the year (a station kept from 1 April counts 0.75 of its MW).
    Returns {"mw": {coal, ccgt, other_gas, oil}, "abated_mw": {FGD, SCR, FGD+SCR}, "stations":
    [(station, unit, MW-share, reason)], "basis", "source_id"}. `closures="actual"`: nothing kept.

    overrides (RB v1.1 M2, `closure_test.fixed_point`): ((station, unit, close ISO or None, filters), ...) in place
    of the register's close and filters for the rows named; () = the register as written.
    real_date_class (RB v1.1 O1): the rows whose `closure_class` is this close on their real dates (after any
    overrides); 'none' = no class row.
    Cached on the edition-MW switches too ($GASCF_EDITION_MW), so a switch turned on in the same
    process is never served a stale result."""
    return _kept_open(year, closures, path, overrides, real_date_class, edition_mw_switches())


@lru_cache(maxsize=None)
def _kept_open(year: int, closures: str, path: str | None, overrides: tuple, real_date_class: str,
               switches: tuple[str, ...]) -> dict:
    if closures not in CLOSURE_RULES:
        raise ValueError(f"unknown closure rule {closures!r}; one of {CLOSURE_RULES}")
    if real_date_class != "none" and real_date_class not in CLOSURE_CLASSES:
        raise ValueError(f"unknown closure class {real_date_class!r}; one of {CLOSURE_CLASSES}")
    mw = {"coal": 0.0, "ccgt": 0.0, "other_gas": 0.0, "oil": 0.0}
    abated = {k: 0.0 for k in ff.ABATEMENT}
    stations = []
    if closures == "actual":
        return {"mw": mw, "abated_mw": abated, "stations": stations, "basis": "measured", "source_id": ""}
    rows = read_closures(path)
    if overrides:
        from gascf import closure_test as CT
        rows = CT.apply(rows, overrides)
    if real_date_class != "none":
        rows = tuple({**r, "fossil_britain_close_d": r["actual_close_d"], "abatement_needed": None}
                     if r.get("closure_class") == real_date_class else r for r in rows)
    for r in rows:
        sh = _open_share(year, r["actual_close_d"], r["fossil_britain_close_d"])
        if sh <= 0:
            continue
        k = stack_kind(r)
        mw[k] += r["capacity_mw"] * sh
        if r["abatement_needed"]:
            if k != "coal" and "FGD" in r["abatement_needed"]:
                raise ValueError(f"closure register: FGD on a {r['fuel']} station ({r['station']})")
            abated[r["abatement_needed"]] += r["capacity_mw"] * sh
        stations.append((r["station"], r["unit"], r["capacity_mw"] * sh, r["reason"]))
    fixture = any(r["source_id"] == "FIXTURE" for r in rows)
    return {"mw": mw, "abated_mw": abated, "stations": stations,
            "basis": "estimated" if fixture else "modelled",
            "source_id": "FIXTURE" if fixture else SRC_CLOSURES}


# ----------------------------------------------------------------- renewables in V

def _bst(d: date) -> bool:
    return rn._is_bst(d)


@lru_cache(maxsize=None)
def _neso_mix_all() -> dict[int, dict[tuple[date, int], dict[str, float]]]:
    """NESO historic generation mix, MW, by year then (local settlement date, period). The
    CSV is in UTC; settlement periods are UK local time (BST from the last Sunday in March)."""
    out: dict[int, dict] = {}
    if not NESO_MIX.exists():
        return out
    with NESO_MIX.open(newline="") as f:
        for r in csv.DictReader(f):
            t = datetime.fromisoformat(r["DATETIME"])
            loc = t + timedelta(hours=1) if _bst(t.date()) else t
            k = (loc.date(), loc.hour * 2 + loc.minute // 30 + 1)
            out.setdefault(loc.year, {})[k] = {c: float(r[c] or 0) for c in
                                               ("WIND", "HYDRO", "BIOMASS", "OTHER", "GAS", "COAL")}
    return out


def neso_mix_year(year: int) -> dict[tuple[date, int], dict[str, float]]:
    return _neso_mix_all().get(year, {})


def _dukes56_supplied(year: int, fuel: str) -> float:
    """GWh, DUKES 5.6.E electricity supplied (gross), all generating companies (UK)."""
    rows56 = rn._sheet(rn.DUKES_56, "5.6")
    t = rn._find(rows56, "Table 5.6.E Electricity supplied (gross)")
    hdr = rows56[t + 1]
    j = [i for i, h in enumerate(hdr) if h is not None and str(h) == str(year)][0]
    return float(rn._row(rows56, t, "All generating companies", fuel)[j])


@lru_cache(maxsize=None)
def hydro_year(year: int, rows_key: tuple) -> tuple[float, ...]:
    """Natural-flow hydro output, MW per half-hour (kept, zero running cost): NESO mix HYDRO
    from 2009, DUKES 5.6 hydro spread flat before."""
    mix = neso_mix_year(year)
    if len(mix) > 0.95 * len(rows_key):
        have = [v["HYDRO"] for v in mix.values()]
        mean = sum(have) / len(have)
        return tuple(mix[k]["HYDRO"] if k in mix else mean for k in rows_key)
    flat = _dukes56_supplied(year, "Hydro (natural flow)") * 1e3 / (0.5 * len(rows_key))
    return tuple(flat for _ in rows_key)


@lru_cache(maxsize=None)
def renewables_trans_year(year: int, rows_key: tuple) -> tuple[tuple[float, ...], str]:
    """Transmission-connected renewables inside V, MW per half-hour, aligned to `rows_key`
    ((day, period) of the residual rows). 2009 on: NESO mix
    WIND + BIOMASS, plus OTHER in half-hours where BIOMASS is blank (NESO's 'other' held
    Drax's converted biomass units before the 2017 split). Hydro is kept, not replaced.
    Before 2009: DUKES 5.6 total wind less NESO embedded wind, spread flat (estimated)."""
    mix = neso_mix_year(year)
    if len(mix) > 0.95 * len(rows_key):
        vals = [mix.get(k) for k in rows_key]
        have = [v for v in vals if v]
        mean = {c: sum(v[c] for v in have) / len(have) for c in have[0]}
        def ren(v):
            v = v or mean
            return v["WIND"] + v["BIOMASS"] + (v["OTHER"] if v["BIOMASS"] <= 0 else 0.0)
        return tuple(ren(v) for v in vals), "modelled"
    wind = _dukes56_supplied(year, "Total wind")
    emb_wind = sum(r.embedded_wind for r in rn.read_neso(year)) / 2e3          # GWh
    gwh = max(0.0, wind - emb_wind)
    flat = gwh * 1e3 / (0.5 * len(rows_key))
    return tuple(flat for _ in rows_key), "estimated"


@lru_cache(maxsize=None)
def embedded_split(year: int) -> dict[tuple[date, int], tuple[float, float]]:
    return {(r.day, r.period): (r.embedded_wind, r.embedded_solar) for r in rn.read_neso(year)}


# ----------------------------------------------------------------- the supply stack

COAL_ETA_MEAN_YEARS = range(2015, 2024)    # spec §15.5 item 3: DUKES 5.10.C 2015-2023 mean


@lru_cache(maxsize=None)
def _eta_dukes(kind: str, year: int) -> float:
    """DUKES 5.10.C fleet efficiency (GCV) as a fraction. DUKES prints 0 for coal in 2025 (no
    coal fleet; `prices.dukes_efficiency` drops it): coal kept open in fossil Britain then
    runs at the 2015-2023 coal-fleet mean (s27a; the s22 rule silently used 2024's 41.5%,
    one station). Any other missing year: the nearest printed year. Both are recorded."""
    e = prices.dukes_efficiency()[kind]
    if year in e:
        return e[year] / 100
    if kind == "coal" and year > max(e):
        v = sum(e[y] for y in COAL_ETA_MEAN_YEARS) / len(COAL_ETA_MEAN_YEARS)
        fb.record("coal_efficiency_mean_2015_2023", year, v, "% GCV",
                  f"mean of DUKES 5.10.C coal {COAL_ETA_MEAN_YEARS.start}-{COAL_ETA_MEAN_YEARS.stop - 1}")
        return v / 100
    y = min(e) if year < min(e) else max(e)
    fb.record("dukes_efficiency_nearest_year", year, e[y], "% GCV", f"{kind}: DUKES 5.10.C {y}")
    return e[y] / 100


@lru_cache(maxsize=None)
def oil_price_quarterly() -> dict[tuple[int, int], float]:
    return ff._qep_quarterly("Major power producers: Oil (pence per kWh)")


def _oil_price(year: int, month: int) -> float:
    q = oil_price_quarterly()
    k = (year, (month - 1) // 3 + 1)
    if k in q:
        return q[k]
    k0 = max(x for x in q if x <= k)
    fb.record("oil_price_carried_forward", year, q[k0], "p/kWh", f"QEP oil {k0[0]} Q{k0[1]} for Q{k[1]}")
    return q[k0]


def stack_month(year: int, month: int, avail: dict[str, float], tv: float = 0.0) -> list[tuple[str, float, float]]:
    """Existing plant as (name, MW available, SRMC £/MWh, no carbon), cheapest first."""
    return [(k, mw, c) for k, mw, c, _ in stack_parts(year, month, avail, tv)]


def stack_parts(year: int, month: int, avail: dict[str, float], tv: float = 0.0):
    """As `stack_month`, with the parts of each SRMC: (name, MW, SRMC, {fuel, transport, vom})."""
    ed = edition(year)
    gas = gas_price(year, month)
    coal = coal_price(year, month)
    eta = {"coal": _eta_dukes("coal", year), "ccgt": _eta_dukes("ccgt", year),
           "other_gas": ETA_OTHER_GAS, "oil": ETA_OIL}
    fuel = {"coal": coal, "ccgt": gas, "other_gas": gas, "oil": _oil_price(year, month)}
    tvk = {"coal": 0.0, "ccgt": tv, "other_gas": tv, "oil": 0.0}
    vom = {"coal": ff.vom_gbp_mwh(ed, "coal"), "ccgt": ff.vom_gbp_mwh(ed, "ccgt"),
           "other_gas": ff.vom_gbp_mwh(ed, "ocgt"), "oil": ff.vom_gbp_mwh(ed, "ocgt")}
    parts = {k: {"fuel": fuel[k] / eta[k], "transport": tvk[k] / eta[k], "vom": vom[k],
                 "burn": 1 / eta[k] if k in ("ccgt", "other_gas") else 0.0} for k in eta}
    out = [(k, avail[k], sum(v for q, v in parts[k].items() if q != "burn"), parts[k])
           for k in avail if avail[k] > 0]
    return sorted(out, key=lambda x: x[2])


def _walk(stack, x: float) -> tuple[float, float, dict[str, float]]:
    """Serve x MW from the stack: (cost £/h, marginal SRMC, MW by block). Beyond the stack's
    total the top block is stretched (the plant was there: it ran in the actual world)."""
    cost, left, gen, marg = 0.0, x, {}, stack[0][2] if stack else 0.0
    for i, (k, mw, c) in enumerate(stack):
        g = min(left, mw) if i < len(stack) - 1 else left
        if g <= 0:
            break
        cost += g * c
        gen[k] = gen.get(k, 0.0) + g          # a block may appear twice (coal under the LF cap)
        marg = max(marg, c)
        left -= g
    return cost, marg, gen


# ----------------------------------------------------------------- one year

@dataclass(frozen=True)
class LfLYear:
    year: int
    market_twh_m: dict           # actual fossil output by month, TWh (paid P − C)
    extra_twh: float
    extra_twh_m: dict            # extra running of existing plant by month, TWh
    energy_twh_m: dict           # whole fossil-Britain fleet output (= residual) by month, TWh
    extra_cost_gbp_m: float      # fuel + VOM of the extra running, no carbon
    extra_fuel_gbp_m: float
    extra_vom_gbp_m: float
    extra_transport_gbp_m: float
    extra_gas_burn_twh: float
    new: dict                    # new-build block (fossil_fleet FossilYear fields)
    firm_spare_gw: float
    firm_gross_gw: float
    unlisted_firm_gw: float      # de-rated firm the actual system had beyond the listed plant (CHP etc.)
    peak_actual_fossil_mw: float
    firm_spare_fixed_gbp_m: float
    moe_model_m: dict            # time-mean Δ by month, £/MWh (model)
    moe_vol_f_m: dict            # Σ Δ_model × actual existing-plant output, £ by month
    moe_vol_x_m: dict            # Σ Δ_model × extra running of existing plant, £ by month
    moe_vol_k_m: dict            # Σ Δ_model × kept output (nuclear, imports, PS, net), £ by month
    moe_model_mean: float
    moe_shao_mean: float
    shao_note: str
    wind_gw_mean: float
    solar_gw_mean: float
    renew_twh: float             # renewables' output replaced (capped at the residual each half-hour)
    renew_gross_twh: float       # renewables' output inside V before the cap
    actual_fossil_twh: float
    energy_twh: float
    peak_mw: float
    existing: Existing
    existing_derated_mw: float
    gas_burn_existing_twh: float
    basis: str
    # s23b (spec §15.3): plant fossil Britain keeps open, and the new-build rule
    kept: dict = None            # kept_open() result + costs (empty under closures="actual")
    kept_fixed_gbp_m: float = 0.0        # keep-open running costs + annualised filters, £m
    kept_abatement_vom_gbp_m: float = 0.0
    extra_gen_by_block_twh: dict = None  # extra running of existing (incl. kept) plant, by block
    closures: str = "actual"
    newbuild: str = "gas_only"
    nuclear_note: str = ""
    # s27a (spec §15.5): coal load-factor cap and the inputs of the gas-price feedback
    coal_lf: str = "none"
    own_firm_basis: str = "actual_residual"      # 0110 item 1.8: "gross_less_contracted" where contracted ICs apply
    peak_actual_fossil_gross_mw: float = 0.0     # peak of actual fossil + metered imports (0110 item 1.8)
    coal_lf_cap: float = 0.0             # cap applied (fraction), 0 = none
    coal_avail_scale: float = 1.0        # share of the coal block left in its merit place (rest moved to the top)
    coal_lf_achieved: float = 0.0        # coal block's yearly output / (nameplate x hours)
    coal_twh: float = 0.0                # coal block's yearly output, TWh (s29: >= actual)
    coal_ied: dict = None                # card [internal ticket]: IED hours-capped kept coal units and the block's limit
    existing_scr_fixed_gbp_m: float = 0.0   # card [internal ticket] "+ex_scr": SCR on existing WBA and Drax 5-6, £m
    existing_scr_vom_gbp_m: float = 0.0     # and its running cost, £m (year-y money)
    gas_marginal_share: float = 0.0      # share of half-hours gas plant is marginal at fossil Britain's load
    gas_marginal_share_actual_load: float = 0.0   # the same at the actual fossil load (check vs Zakeri 2023)
    extra_gas_fuel_gbp_m: float = 0.0    # gas commodity cost of the extra running of existing plant (no transport)
    kept_twh_m: dict = None              # kept output (nuclear, imports, PS, net) by month, TWh
    sizing: dict = None                 # 0057 D3/D4: the sizing peak and its parts (bills.sizing_peak)
    firm_req: dict = None               # 0065: firm requirement and margin-on-gross extra
    ic_called_twh_m: dict = None        # 0063: contracted imports called, TWh by month
    ic_called_twh: float = 0.0
    ic_called_halfhours: int = 0
    worked_hh: dict = None               # 44 (B20): 2024 residual peak and median half-hours, for the §4.3 worked example
    m8_sums: dict = None                 # RB v1.1 M8 (0163 R9 a): {month: {"f": MonthSums, "x": MonthSums}} when switched on
    hh_cap: tuple = None                 # 44 debug: every half-hour (residual, day, period, d, p0, p1, top), set only when keep_hh
    hh_debug: dict = None                # 44 debug: {stacks, new_srmc, new_ccgt_mw} for recomputing a half-hour step, set only when keep_hh
    closure_test: dict = None            # RB v1.1 M2: the road-load-factor loop (closure_test.fixed_point), None under 'register'
    # v4-s40 item 2 (0175): the coal block's own output by month, its uplift-weighted volume (sum of the half-hour uplift
    # x coal MW / 2, as moe_vol_f_m) and its M8 half-hour sums, so kept coal's margin reads the road's own revenue
    coal_twh_m: dict = None
    coal_moe_vol_m: dict = None
    coal_m8: dict = None


def _new_block(year, rows, short, transport, der, newbuild: str = "gas_only", cost_case: str = "medium"):
    """New plant on the half-hours where existing plant runs out, sized and costed by the
    edition in force (`fossil_fleet.cost_fossil_year`, margin 0: the firm margin is checked
    against the whole system below).

    newbuild="gas_only" (s22): CCGT/OCGT only (coal cap 0). "cheapest" (spec §15.3 item 4):
    each strip is the cheapest of CCGT, OCGT and coal with FGD and NOx control (the edition's
    coal plant plus its NOx retrofit), no cap on new coal; new nuclear only where
    `fossil_fleet.nuclear_available` allows it (it never does in 2002-2025; if it ever did, the
    engine stops rather than silently leaving nuclear out)."""
    if newbuild not in NEWBUILD_RULES:
        raise ValueError(f"unknown new-build rule {newbuild!r}; one of {NEWBUILD_RULES}")
    if max(short) <= 0:
        return None
    if newbuild == "cheapest" and ff.nuclear_available(year)[0]:
        raise NotImplementedError(f"{year}: new nuclear is available under the edition in force; "
                                  "costing it in the new-build choice is not built")

    @dataclass(frozen=True)
    class _R:
        day: date
        residual: float
    rr = [_R(r.day, s) for r, s in zip(rows, short)]
    if newbuild == "cheapest":
        return ff.cost_fossil_year(year, rr, 0.0, der, coal_cap_mw=float("inf"), transport=transport,
                                   coal_rule="allowed", cost_case=cost_case)
    return ff.cost_fossil_year(year, rr, 0.0, der, coal_cap_mw=0.0, transport=transport, cost_case=cost_case)


def kept_costs(year: int, kept: dict, cost_case: str = "medium") -> tuple[float, float, dict]:
    """(fixed £m, abatement £/MWh on kept coal, detail) for the plant fossil Britain keeps open:
    keep-open running costs of the edition's plant of that type (coal: coal; CCGT: CCGT; other
    gas and oil: OCGT) and the annualised filters, year-y money. No new capital for the station
    itself: it was built and paid for in the actual world. cost_case (RB v1.1 O2) reaches the filters' capital."""
    ed = edition(year, cost_case)
    tech = {"coal": "coal", "ccgt": "ccgt", "other_gas": "ocgt", "oil": "ocgt"}
    keep = {k: ff.keep_open_gbp_kw_yr(ed, tech[k]) * 1000 * mw / 1e6 for k, mw in kept["mw"].items()}
    ab_fixed, ab_vom_w, src = 0.0, 0.0, set()
    for need, mw in kept["abated_mw"].items():
        if mw <= 0:
            continue
        f, v, s = ff.abatement_gbp_kw_yr(ed, need)
        ab_fixed += f * 1000 * mw / 1e6
        ab_vom_w += v * mw
        src |= set(s)
    coal_kept = kept["mw"]["coal"]
    ab_vom = ab_vom_w / coal_kept if coal_kept else 0.0        # MW-weighted £/MWh over kept coal
    return sum(keep.values()) + ab_fixed, ab_vom, {"keep_open_gbp_m": keep, "abatement_fixed_gbp_m": ab_fixed,
                                                   "abatement_source_ids": sorted(src)}


@lru_cache(maxsize=None)
def dukes_coal_lf_by_year() -> dict[int, float]:
    """Coal fleet load factor by year, fraction: DUKES 5.10.B 'of which coal-fired stations'."""
    ws = openpyxl.load_workbook(prices.DUKES_510, read_only=True, data_only=True)["5.10.B and 5.10.C"]
    rows = list(ws.iter_rows(values_only=True))
    i = next(k for k, r in enumerate(rows) if r[0] and str(r[0]).startswith("Table 5.10.B"))
    h = rows[i + 1]
    yrs = {j: int(str(v)[:4]) for j, v in enumerate(h) if v and str(v)[:4].isdigit()}
    r = next(r for r in rows[i + 2:i + 10] if str(r[0] or "").strip().startswith("of which coal-fired"))
    return {y: float(r[j]) / 100 for j, y in yrs.items() if isinstance(r[j], (int, float)) and r[j] > 0}


def dukes_coal_lf_max() -> tuple[float, int]:
    """(max load factor as a fraction, year): DUKES 5.10.B 'of which coal-fired stations'."""
    vals = dukes_coal_lf_by_year()
    y = max(vals, key=vals.get)
    return vals[y], y


@lru_cache(maxsize=None)
def actual_coal_supplied_twh(year: int) -> float:
    """Coal power real Britain sent out, TWh: DUKES 5.6 'Annual summaries' (Table 5.6.J),
    major power producers, coal (thermal), 'Supplied (gross)' (generation less own use; the
    sent-out basis of the engine's stack). The floor under the coal load-factor cap (s29).
    Supplied (gross) is below Generation, so the floor is the smaller (conservative) figure."""
    ws = openpyxl.load_workbook(rn.DUKES_56, read_only=True, data_only=True)["5.6 Annual summaries"]
    rows = list(ws.iter_rows(values_only=True))
    title = f"Table 5.6.J {year} summary"
    i = next(k for k, r in enumerate(rows) if str(r[0] or "").startswith(title))
    h = rows[i + 1]
    jc = next(j for j, v in enumerate(h) if str(v or "").startswith("Coal"))
    for r in rows[i + 2:]:
        if str(r[0] or "").startswith("Table 5.6.J"):
            break
        if str(r[0] or "").startswith("Major power") and str(r[1] or "").strip() == "Supplied (gross)":
            return float(r[jc] or 0.0) / 1e3
    raise KeyError(f"DUKES 5.6.J has no major-power coal 'Supplied (gross)' row for {year}")


def coal_lf_cap(rule: str) -> float | None:
    if rule not in COAL_LF_RULES:
        raise ValueError(f"unknown coal load-factor rule {rule!r}; one of {COAL_LF_RULES}")
    return dukes_coal_lf_max()[0] if rule.startswith("dukes_max") else None


@lru_cache(maxsize=None)
def _ied_hours_file() -> dict:
    return json.loads(IED_HOURS_FILE.read_text())


def ied_capped_units(year: int, rule: str, path: str | None = None) -> list[tuple[str, str, float, float, float]]:
    """Kept coal units held to an IED hours cap in `year` under coal_lf `rule` (card [internal ticket]):
    [(station, unit, register MW, share of the year kept on or after the cap's start, cap hours a year)].
    Empty unless the rule is one of IED_HOURS_RULES. Dates are the closure register's, as `kept_open`."""
    coal_lf_cap(rule)                                   # rejects an unknown rule
    rule = rule.split("+ex_")[0]                        # card [internal ticket]: the existing-coal suffix is not kept coal
    if rule not in IED_HOURS_RULES:
        return []
    reading, hours = rule.removeprefix("dukes_max+ied_").removesuffix("+existing").rsplit("_", 1)
    f = _ied_hours_file()
    start = date.fromisoformat(f["from"])
    want = {(u["station"], u["unit"]) for u in f["units"] if reading in u["readings"]}
    out = []
    for r in read_closures(path):
        if r["fuel"] != "coal" or (r["station"], r["unit"]) not in want:
            continue
        sh = _open_share(year, max(r["actual_close_d"], start), r["fossil_britain_close_d"])
        if sh > 0:
            out.append((r["station"], r["unit"], r["capacity_mw"], sh, float(hours)))
    return out


def existing_ied_units(year: int, rule: str, path: str | None = None) -> list[tuple]:
    """Card [internal ticket] (measure only): EXISTING coal stations the IED reaches from 1 July 2020 under the rule's
    suffix: [(station, unit, register MW, share of the year, cap hours a year or None under "+ex_scr")].
    A station is existing in a year if it is in DUKES's end-of-year capacity (the engine's coal-block basis):
    every year before its real closure year. Empty without a suffix."""
    coal_lf_cap(rule)                                   # rejects an unknown rule
    if "+ex_" not in rule:
        return []
    mode = rule.split("+ex_")[1]
    hours = None if mode == "scr" else float(mode.removeprefix("cap_"))
    f = _ied_hours_file()
    start = date.fromisoformat(f["from"])
    want = {(u["station"], u["unit"]) for u in f["existing_jul2020"]}
    out = []
    for r in read_closures(path):
        if r["fuel"] != "coal" or (r["station"], r["unit"]) not in want:
            continue
        sh = _open_share(year, start, date(r["actual_close_d"].year, 1, 1))
        if sh > 0:
            out.append((r["station"], r["unit"], r["capacity_mw"], sh, hours))
    return out


def existing_scr_costs(year: int, rule: str, path: str | None = None) -> tuple[float, float]:
    """(fixed £m, filters' running cost £/MWh), year-y money: SCR on the existing stations under "+ex_scr",
    costed exactly as `kept_costs` costs SCR on kept coal. (0, 0) otherwise."""
    eq = sum(mw * sh for _, _, mw, sh, h in existing_ied_units(year, rule, path) if h is None)
    if eq <= 0:
        return 0.0, 0.0
    f, v, _ = ff.abatement_gbp_kw_yr(edition(year), "SCR")
    return f * 1000 * eq / 1e6, v


def _year_hours(year: int) -> float:
    return 8784.0 if year % 4 == 0 else 8760.0


def lcpd_hours_used(year: int, lf_by_year: dict[int, float]) -> float:
    """Card [internal ticket]: an LCPD opt-out station's operating hours used from the opt-out's start to the end of
    `year` - 1, at the coal block's achieved load factor each year (`lf_by_year`), never more than the total."""
    f = _ied_hours_file()["existing"]["lcpd"]
    total, used = float(f["total_hours"]), 0.0
    for t in range(date.fromisoformat(f["from"]).year, year):
        used += min(total - used, lf_by_year[t] * _year_hours(t))
    return used


def ied_existing_units(year: int, rule: str, path: str | None = None,
                       lcpd_used_h: float = 0.0,
                       kept_stations: tuple = ()) -> list[tuple[str, str, float, float, float]]:
    """Existing (not kept) coal held to its real hours cap in `year` (card [internal ticket]), as `ied_capped_units`:
    [(station, unit, register MW, share of the year capped, cap hours a year)]. Empty unless the rule ends in
    '+existing'. A station counts only while in the coal block on the engine's DUKES end-year capacity basis
    (register coal end on or after 31 December). LCPD opt-outs: hours left = total - `lcpd_used_h`, listed only
    when that is under the block cap x the year's hours (otherwise the cap does not bind). Permits: the reading's
    hours a year from the permits' start, pro rata.
    Card [internal ticket] (built OFF; ON by default in v4-s39 under 0166): with $GASCF_IED_LCPD_KEPT=on, an LCPD opt-out that is KEPT open
    past its real close (`kept_stations`, the `kept_open` list for the year, as M2's closure test sets it) is held to
    the same hours left for the share of the year it is kept: (station, unit, register MW, kept MW / register MW,
    hours left). Set the variable before the process starts (lfl_year is cached)."""
    cap = coal_lf_cap(rule)
    if rule not in IED_HOURS_RULES or not rule.endswith("+existing"):
        return []
    hours = float(rule.removesuffix("+existing").rsplit("_", 1)[1])
    f = _ied_hours_file()["existing"]
    lc, pm = f["lcpd"], f["permits"]
    lc_units = {(u["station"], u["unit"]) for u in lc["units"]}
    pm_units = {(u["station"], u["unit"]) for u in pm["units"]}
    lc0, lc1 = date.fromisoformat(lc["from"]), date.fromisoformat(lc["to"])
    pm0 = date.fromisoformat(pm["from"])
    left = max(0.0, float(lc["total_hours"]) - lcpd_used_h)
    out = []
    lcpd_binds = lc0.year <= year <= lc1.year and left < cap * _year_hours(year)
    kept_mw = {}
    if lcpd_kept_on():
        for s, u, mw, _ in kept_stations:
            kept_mw[(s, u)] = kept_mw.get((s, u), 0.0) + mw
    for r in read_closures(path):
        key = (r["station"], r["unit"])
        if r["fuel"] == "coal" and key in lc_units and lcpd_binds and kept_mw.get(key, 0.0) > 0:
            out.append((r["station"], r["unit"], r["capacity_mw"], kept_mw[key] / r["capacity_mw"], left))
            continue
        if r["fuel"] != "coal" or r["actual_close_d"] < date(year, 12, 31):
            continue
        if key in lc_units and lcpd_binds:
            out.append((r["station"], r["unit"], r["capacity_mw"], 1.0, left))
        elif key in pm_units:
            sh = _open_share(year, pm0, None)
            if sh > 0:
                out.append((r["station"], r["unit"], r["capacity_mw"], sh, hours))
    return out


def coal_energy_limit_mwh(year: int, coal_mw: float, hours: float, rule: str,
                          path: str | None = None, extra_units: list | tuple = ()) -> float | None:
    """The coal block's yearly energy limit, MWh, before the actual-output floor (s29). Off ('dukes_max'):
    cap x nameplate x hours. On (IED_HOURS_RULES): each capped unit at most its cap hours for the share of the
    year it is capped; the rest of the block at the 65.04% cap. None under 'none'. `extra_units`: existing coal
    held to its caps (card [internal ticket], `ied_existing_units`)."""
    cap = coal_lf_cap(rule)
    if cap is None:
        return None
    u = (list(ied_capped_units(year, rule, path)) + [x for x in existing_ied_units(year, rule, path) if x[4] is not None]
         + list(extra_units))
    capped_eq = sum(mw * sh for _, _, mw, sh, _ in u)
    return cap * (coal_mw - capped_eq) * hours + sum(mw * sh * h for _, _, mw, sh, h in u)


def _cap_stack(st, sc: float):
    """The stack with the coal block split under the load-factor cap: a share `sc` in its merit
    place, the rest (1 - sc) moved to the top of the stack, so it runs only in the half-hours
    when every other existing block is running (the peak). An annual energy limit, not a
    capacity limit: the full coal fleet stays available at the peak, as it was in 2003, the
    year the fleet ran at 65%."""
    if sc >= 1.0:
        return st
    out = [(k, mw * sc if k == "coal" else mw, c) for k, mw, c in st]
    top = [(k, mw * (1 - sc), c) for k, mw, c in st if k == "coal" and mw > 0]
    # just below the dearest block, which stays last: `_walk` stretches the last block when the
    # actual fossil output exceeds the listed plant, and that must not be counted as coal
    return out[:-1] + top + out[-1:] if out and out[-1][0] != "coal" else out + top


def _coal_scale(year, avail, coal_mw, cap, L, F, H, rows, tv, cap_mwh: float | None = None) -> tuple[float, float]:
    """(share of the coal block left in its merit place, achieved load factor). The coal
    block's yearly output in fossil Britain (actual + extra running, existing + kept coal) may
    not exceed cap x nameplate x hours. The share is lowered until it does (`_cap_stack`): the
    displaced coal output goes to the next blocks in the merit order. Peak availability, the
    existing-plant ceiling and the firm test are unchanged (s27a fix: the s26a design scaled
    the coal MW themselves, which also cut peak capacity and built new OCGT in 2003, a year
    the actual fleet met the same load with the same plant)."""
    hours = len(rows) / 2
    if coal_mw <= 0 or avail.get("coal", 0.0) <= 0:
        return 1.0, 0.0
    A = sum(avail.values())
    base = {m: stack_month(year, m, avail, tv) for m in range(1, 13)}

    def out(sc):
        st = {m: _cap_stack(base[m], sc) for m in base}
        return sum(_walk(st[r.day.month], max(0.0, min(l, max(A, f)) - h))[2].get("coal", 0.0) / 2
                   for r, l, f, h in zip(rows, L, F, H))

    o = out(1.0)
    if cap is None:
        return 1.0, o / (coal_mw * hours)
    # s29: never below the coal real Britain actually sent out that year: fossil Britain has
    # the same coal plant with carbon taken off, so it cannot burn less coal than reality.
    # The engine's end-of-year coal MW is below DUKES's capacity basis, so 65% of it fell
    # below the actual output in 2003, 2005-07, 2012 and 2013. Scaling the cap to DUKES's
    # basis instead was rejected: it runs the engine's smaller fleet at 77-83% and raises the
    # gap (spec §15.1 conservative rule).
    # card [internal ticket]: `cap_mwh` (coal_energy_limit_mwh) replaces cap x nameplate x hours when IED hours caps are on
    # v4-s40 item 1 (0175): a permit-derived limit (cap_mwh) is absolute unless $GASCF_COAL_CEILING=floored (v4-s39)
    if cap_mwh is not None and coal_ceiling_reading() == "absolute":
        lim = cap_mwh
    else:
        lim = max(cap * coal_mw * hours if cap_mwh is None else cap_mwh, actual_coal_supplied_twh(year) * 1e6)
    lo, hi = 0.0, 1.0
    if o <= lim * (1 + 1e-4):
        return 1.0, o / (coal_mw * hours)
    for _ in range(40):                        # output rises with the share: bisect
        mid = (lo + hi) / 2
        om = out(mid)
        if om > lim:
            hi = mid
        else:
            lo, o = mid, om
        if hi - lo < 1e-4:
            break
    return lo, out(lo) / (coal_mw * hours)


@lru_cache(maxsize=None)
def road_lf_closures(residual: str, margin: str, transport_key, newbuild: str, closures_file: str | None,
                     coal_lf: str, ic_basis: str, real_date_class: str = "none", cost_case: str = "medium") -> dict:
    """RB v1.1 M2 (decision 0163 R1-R3 (a)): the 'mixed' rows' closes by the margin test at the road's own coal load
    factor, repeated until the list stops changing (`closure_test.fixed_point`). The load factor is the road's under
    the same settings as the run (so a sensitivity that moves the road's coal output re-runs the test)."""
    from gascf import closure_test as CT
    path = closures_file or str(closures_path())
    rows = [r for r in read_closures(path) if r["reason"] == "mixed"]
    for r in rows:
        if "closure_test" not in r:
            raise ValueError(f"closure register: mixed row {r['station']} / {r['unit']} has no closure_test data")

    def lf_of(y: int, state: tuple) -> float:
        return lfl_year(y, residual, margin, transport_key, "legitimate", newbuild, closures_file, coal_lf, ic_basis,
                        False, "register", state, real_date_class, cost_case).coal_lf_achieved

    def margin_of(r: dict, y: int, lf: float, state: tuple) -> tuple[float, float]:
        m = CT.station_margin(r["closure_test"]["margin_name"], y, lf)
        return m["margin_before_filters_gbp_m"], m["margin_after_fgdscr_gbp_m"]

    return {"rule": "road_lf", **CT.fixed_point(rows, lf_of, margin_of)}


@lru_cache(maxsize=None)
def _mid_year(year: int) -> dict | None:
    """Elexon MID, APX rows, as held by the engine (data/raw/elexon/MID_APX_YYYY.csv, `gascf.mid`); None for a year
    with no file (the engine holds 2018 on; Elexon's free MID starts 12 Sep 2016)."""
    from gascf.mid import OUT, read_mid
    return read_mid(year) if (OUT / f"MID_APX_{year}.csv").exists() else None


def market_index_price(day, sp: int) -> float | None:
    """RB v1.1 F7 (decision 0163, R25 (a)): Elexon's market index price (MID, APX), £/MWh cash, for one settlement
    period; None where the engine holds no MID for that year or period (never a filled value). It is a short-term
    market index, not the day-ahead auction price, and is labelled so."""
    from datetime import date as _date
    d = _date.fromisoformat(day) if isinstance(day, str) else day
    m = _mid_year(d.year)
    if m is None or (d, int(sp)) not in m:
        return None
    return m[(d, int(sp))][0]


def _worked_hh(cap: list, factor: tuple) -> dict | None:
    """44 (B20): the 2024 worked half-hours. peak = largest residual (earliest on ties); median = the lower
    middle of the residual sorted ascending, rank (n-1)//2 (earliest on ties). `cap` is chronological, so the
    first row attaining a value is the earliest. `factor` is the year's (mult, flat) MOE pair from moe_factor;
    scaled step = raw x mult + flat. RB v1.1 F7 (0163 R25 (a)): each carries Elexon's market index price for
    that half-hour (`market_index_price`), None where none is held."""
    if not cap:
        return None
    mult, flat = factor

    def entry(res, day, period, d, p0, p1, top, *_):
        return {"date": day.isoformat(), "sp": period, "residual_mw": res,
                "raw_step_gbp_mwh": d, "scaled_step_gbp_mwh": d * mult + flat,
                "p0_gbp_mwh": p0, "p1_gbp_mwh": p1, "rb_cost_gbp_mwh": p1, "top_block": top,
                "market_index_price_gbp_mwh": market_index_price(day, period)}

    peak_res = max(c[0] for c in cap)
    peak = next(c for c in cap if c[0] == peak_res)
    asc = sorted(cap, key=lambda c: c[0])
    med_res = asc[(len(asc) - 1) // 2][0]
    median = next(c for c in cap if c[0] == med_res)
    return {"peak": entry(*peak), "median": entry(*median)}


@lru_cache(maxsize=None)
def lfl_year(year: int, residual: str = "neso", margin: str = "central", transport_key=None,
             closures: str = DEFAULT_CLOSURES, newbuild: str = DEFAULT_NEWBUILD,
             closures_file: str | None = None, coal_lf: str = DEFAULT_COAL_LF,
             ic_basis: str = "metered", keep_hh: bool = False, closure_test: str = "register",
             overrides: tuple = (), real_date_class: str = "none", cost_case: str = "medium") -> LfLYear:
    """cost_case (RB v1.1 O2, `fossil_fleet.COST_CASES`): the plant-cost case for new build, firm spare plant and the
    kept coal's filters; 'medium' = central. The merit-order stack uses running costs only, which no case moves.
    rb-v4-s39-rc merge: under closure_test='road_lf' the fixed point runs at the same cost_case (the road's load
    factor 'under the same settings as the run', road_lf_closures)."""
    from gascf import bills as B
    from gascf import closure_test as CT
    if closure_test not in CT.RULES:
        raise ValueError(f"unknown closure_test rule {closure_test!r}; one of {CT.RULES}")
    ct_detail = None
    if closure_test == "road_lf" and closures == "legitimate" and real_date_class != "mixed":   # O1 mixed row: all real dates
        if overrides:
            raise ValueError("closure_test='road_lf' sets its own overrides")
        ct_detail = road_lf_closures(residual, margin, transport_key, newbuild, closures_file, coal_lf, ic_basis,
                                     real_date_class, cost_case)
        overrides = ct_detail["overrides"]
    if residual == "fuelhh" and year >= 2016:
        from gascf.residual import residual_year
        rows = residual_year(year, replace_other=True)[0]
    else:
        rows = rn.residual_neso_year(year)
    tr = B.transport_for(year) if transport_key else None
    der = capacity.read_derating()
    ex0 = existing_plant(year)
    kept = kept_open(year, closures, closures_file or str(closures_path()) if closures == "legitimate" else None,
                     overrides, real_date_class)
    km = kept["mw"]
    ex = Existing(coal=ex0.coal + km["coal"], ccgt=ex0.ccgt + km["ccgt"], other_gas=ex0.other_gas + km["other_gas"],
                  oil=ex0.oil + km["oil"], basis=ex0.basis if not kept["stations"] else
                  ("estimated" if "estimated" in (ex0.basis, kept["basis"]) else ex0.basis))
    avail = ex.derated(der)
    A_actual = sum(ex0.derated(der).values())     # the plant the actual system really had
    keys = tuple((r.day, r.period) for r in rows)
    rtr, rbasis = renewables_trans_year(year, keys)
    L = [max(0.0, r.residual) for r in rows]
    R = [t + r.embedded for t, r in zip(rtr, rows)]
    F = [max(0.0, l - x) for l, x in zip(L, R)]                 # actual fossil output
    tv = (B.transport_for(year).commodity_gbp_mwh_gas if transport_key else 0.0)
    H = hydro_year(year, keys)
    cap = coal_lf_cap(coal_lf)
    cpath = closures_file or str(closures_path()) if closures == "legitimate" else None
    ied_u = ied_capped_units(year, coal_lf, cpath) if kept["stations"] else []
    ex_jul = existing_ied_units(year, coal_lf, cpath)        # card [internal ticket]: existing coal under the IED, Jul 2020
    lcpd_used = 0.0
    lc_from = date.fromisoformat(_ied_hours_file()["existing"]["lcpd"]["from"]).year
    if coal_lf.endswith("+existing") and year > lc_from:
        # card [internal ticket]: LCPD hours used to the end of last year, at this rule's own block load factors.
        # rb-v4-s39-rc merge: the earlier years run with this run's own closure list, class row and cost case
        # (closure_test resolved to its overrides above), so under road_lf the hours are the road's own.
        lcpd_used = lcpd_hours_used(year, {t: lfl_year(t, residual, margin, transport_key, closures, newbuild,
                                                       closures_file, coal_lf, ic_basis, keep_hh, "register",
                                                       overrides, real_date_class, cost_case).coal_lf_achieved
                                           for t in range(lc_from, year)})
    ex_u = ied_existing_units(year, coal_lf, cpath, lcpd_used, tuple(kept["stations"]))
    cap_mwh = (coal_energy_limit_mwh(year, ex.coal, len(rows) / 2, coal_lf, cpath, ex_u)
               if (ied_u or ex_u or any(u[4] is not None for u in ex_jul)) and ex.coal > 0 else None)
    scale, lf_got = _coal_scale(year, avail, ex.coal, cap, L, F, H, rows, tv, cap_mwh)
    capped_eq = sum(mw * sh for _, _, mw, sh, _ in ied_u)
    coal_ied = {"rule": coal_lf, "capped_mw_eq": capped_eq,
                "capped_cap_twh": sum(mw * sh * h for _, _, mw, sh, h in ied_u) / 1e6,
                "limit_twh": cap_mwh / 1e6 if cap_mwh is not None else None,
                "units": [{"station": s, "unit": u, "mw": mw, "capped_share": sh, "cap_hours": h}
                          for s, u, mw, sh, h in ied_u],
                "existing": [{"station": s, "unit": u, "mw": mw, "share": sh, "cap_hours": h}
                             for s, u, mw, sh, h in ex_jul],
                "existing_mw_eq": sum(mw * sh for _, _, mw, sh, _ in ex_u),
                # v4-s40 item 1: whether the s29 actual-output floor applies over the limit (False where a
                # permit-derived limit binds absolutely; True under the DUKES cap alone or the v4-s39 reading)
                "actual_floor": cap is not None and (cap_mwh is None or coal_ceiling_reading() == "floored"),
                "ceiling_reading": coal_ceiling_reading(),
                "existing_cap_twh": sum(mw * sh * h for _, _, mw, sh, h in ex_u) / 1e6,
                "lcpd_hours_used_before": lcpd_used,
                "existing_units": [{"station": s, "unit": u, "mw": mw, "capped_share": sh, "cap_hours": h}
                                   for s, u, mw, sh, h in ex_u]}
    A = sum(avail.values())                       # full coal MW: the cap limits energy, not capacity
    serv = [min(l, max(A, f)) for l, f in zip(L, F)]            # existing plant in fossil Britain
    short = [l - s for l, s in zip(L, serv)]
    new = _new_block(year, rows, short, tr, der, newbuild, cost_case)
    ed = edition(year, cost_case)      # O2: firm spare fixed cost; efficiencies and VOM are the same in every case
    stacks = {m: _cap_stack(stack_month(year, m, avail, tv), scale) for m in range(1, 13)}
    sparts = {m: {k: p for k, _, _, p in stack_parts(year, m, avail, tv)} for m in range(1, 13)}
    xparts = {"fuel": 0.0, "transport": 0.0, "vom": 0.0, "burn": 0.0}
    xgen = {k: 0.0 for k in avail}
    xgas_fuel = 0.0
    n_gas_marg = n_gas_marg0 = 0
    kept_m = {m: 0.0 for m in range(1, 13)}
    new_srmc = {}
    for m in range(1, 13):
        gas = gas_price(year, m)
        new_srmc[m] = {k: (gas + tv) / getattr(ed, k).efficiency_gcv + ff.vom_gbp_mwh(ed, k) for k in ("ccgt", "ocgt")}
    new_ccgt_mw = (new.ccgt_gw + new.coal_gw) * 1e3 if new else 0.0     # baseload new plant
    extra_cost = 0.0
    gas_burn = 0.0
    mkt_m = {m: 0.0 for m in range(1, 13)}
    xm = {m: 0.0 for m in range(1, 13)}
    lm = {m: 0.0 for m in range(1, 13)}
    hh_cap: list = []
    moe_sum = {m: 0.0 for m in range(1, 13)}
    moe_f = {m: 0.0 for m in range(1, 13)}
    moe_x = {m: 0.0 for m in range(1, 13)}
    moe_k = {m: 0.0 for m in range(1, 13)}
    n_m = {m: 0 for m in range(1, 13)}
    from gascf import hh_pricing as H8
    m8 = {m: {"f": H8.MonthSums(), "x": H8.MonthSums()} for m in range(1, 13)} if H8.active() else None
    coal_m = {m: 0.0 for m in range(1, 13)}
    coal_moe = {m: 0.0 for m in range(1, 13)}
    coal_m8 = {m: H8.MonthSums() for m in range(1, 13)} if m8 is not None else None
    for r, l, f, s, sh, h in zip(rows, L, F, serv, short, H):
        m = r.day.month
        if m8 is not None:
            p8, g8 = H8.mid_at(r.day, r.period), H8.is_gap(r.day, r.period)
            m8[m]["f"].add(p8, f, g8)
            m8[m]["x"].add(p8, s - f, g8)
        st = stacks[m]
        w0, w1 = max(0.0, f - h), max(0.0, s - h)        # hydro runs first, at no running cost
        c0, p0, g0 = _walk(st, w0)
        c1, p1, g1 = _walk(st, w1)
        for k in g1:
            dg = (g1[k] - g0.get(k, 0.0)) / 2                # MWh of extra running, this block
            xgen[k] += dg
            for q in xparts:
                xparts[q] += dg * sparts[m][k][q]
            if k in ("ccgt", "other_gas"):
                xgas_fuel += dg * gas_price(year, m) * sparts[m][k]["burn"]     # commodity only
        # gas marginal at fossil Britain's load: the top existing block that ran is gas, or
        # new plant (CCGT/OCGT; new coal only under 'cheapest', counted as gas: small) is running
        if sh > 0 or (g1 and list(g1)[-1] in ("ccgt", "other_gas")):
            n_gas_marg += 1
        if g0 and list(g0)[-1] in ("ccgt", "other_gas"):
            n_gas_marg0 += 1
        kept_m[m] += r.kept / 2e6
        if sh > 0:
            p1 = max(p1, new_srmc[m]["ocgt"] if sh > new_ccgt_mw + 1e-6 else new_srmc[m]["ccgt"])
        extra_cost += (c1 - c0) / 2
        d = max(0.0, p1 - p0)
        top_block = (("new_ocgt" if sh > new_ccgt_mw + 1e-6 else "new_ccgt") if sh > 1e-9
                     else (list(g1)[-1] if g1 else None))
        hh_cap.append((r.residual, r.day, r.period, d, p0, p1, top_block, m, w0, w1, sh))
        moe_sum[m] += d
        moe_f[m] += d * f / 2
        moe_x[m] += d * (s - f) / 2
        moe_k[m] += d * r.kept / 2
        g1c = g1.get("coal", 0.0)                       # v4-s40 item 2: the coal block this half-hour, MW
        coal_m[m] += g1c / 2e6
        coal_moe[m] += d * g1c / 2
        if coal_m8 is not None:
            coal_m8[m].add(p8, g1c, g8)
        n_m[m] += 1
        mkt_m[m] += f / 2e6
        xm[m] += (s - f) / 2e6
        lm[m] += l / 2e6
        gas_burn += sum(g * sparts[m][k]["burn"] for k, g in g1.items()) / 2e6
    extra_twh = (sum(serv) - sum(F)) / 2e6
    extra_vom = xparts["vom"] / 1e6
    extra_fuel = xparts["fuel"] / 1e6
    extra_transport = xparts["transport"] / 1e6
    extra_burn = xparts["burn"] / 1e6
    # firm capacity (like for like): the actual system's own adequacy is taken as given; fossil
    # Britain buys only the extra shortfall that taking renewables out creates. Need = peak ×
    # (1 + margin) against existing plant de-rated + the new block de-rated, less the same test
    # on the actual fossil residual (which the real system met with plant this module does not
    # list: industrial CHP, hydro, other generators).
    peak = max(L)
    peak_f = max(F)
    mf = capacity.margin_fraction(margin)
    new_der = ((new.ccgt_gw + new.coal_gw) * der["ccgt"] + (new.ocgt_gw + new.spare_gw) * der["ocgt"]) * 1e3 if new else 0.0
    sp = B.sizing_peak(year, rows, ic_basis, residual)
    if sp["contracted"]:
        # 0057 option B / 0065: requirement = gross peak × (1 + m) − contracted interconnectors
        fr = B.firm_requirement(sp["gross_peak_mw"] / 1e3, mf, sp["imports_contracted_gw"])
        gross = max(0.0, fr["firm_requirement_gw"] * 1e3 - A - new_der)
    else:
        fr = {"firm_requirement_gw": peak * (1 + mf) / 1e3, "margin_on_gross_extra_gw": 0.0}
        gross = max(0.0, peak * (1 + mf) - A - new_der)
    own = max(0.0, peak_f * (1 + mf) - A_actual)   # the actual system's own (unlisted) firm plant
    own_basis, peak_fg = "actual_residual", peak_f
    if sp["contracted"]:
        # 0110 item 1.8 (Richard Lyon, 19 Oct 2026, option b): the actual system's own firm test on the same basis
        # as Reliable Britain's, so one basis in one subtraction: (actual fossil + metered imports) peak x (1 + m)
        # - contracted imports - the plant the actual system had. Imports matched by day and period, as sizing_peak.
        imp = {(h.day, h.period): h.imports for h in rn.residual_neso_year(year)}
        peak_fg = max(f + imp.get(k, 0.0) for f, k in zip(F, keys))
        own = max(0.0, peak_fg * (1 + mf) - sp["imports_contracted_gw"] * 1e3 - A_actual)
        own_basis = "gross_less_contracted"
    spare = max(0.0, gross - own) / der["ocgt"]
    # 0063: contracted imports called where the metered residual exceeds what the fleet can generate
    new_mw = ((new.ccgt_gw + new.coal_gw + new.ocgt_gw + new.spare_gw) * 1e3) if new else 0.0
    called_m = {m: 0.0 for m in range(1, 13)}
    n_called = 0
    for r_, l, f_ in zip(rows, L, F):
        c = max(0.0, l - (max(A, f_) + new_mw + spare))
        if c > 0:
            called_m[r_.day.month] += c / 2e6
            n_called += 1
    fc = ff.fixed_gbp_kw_yr(ed, "ocgt") + (tr.fixed_gbp_kw_e_yr(ed.ocgt.efficiency_gcv) if tr else 0.0)
    # Shao cross-check: mean wind and solar GW removed
    emb = embedded_split(year)
    mix = neso_mix_year(year)
    wind_tr = [mix[k]["WIND"] if k in mix else None for k in keys]
    if rbasis == "estimated" or sum(v is None for v in wind_tr) > 0.05 * len(keys):
        wind_gw = _dukes56_supplied(year, "Total wind") / (len(keys) / 2)
        fb.record("wind_gw_dukes_flat", year, wind_gw, "GW mean", "DUKES 5.6.E wind spread flat")
    else:
        wv = [v for v in wind_tr if v is not None]
        wind_gw = (sum(wv) / len(wv) + sum(emb.get(k, (0, 0))[0] for k in keys) / len(keys)) / 1e3
    solar_gw = sum(emb.get(k, (0, 0))[1] for k in keys) / len(keys) / 1e3
    b, note = shao_wind_coeff(year)
    shao = -b * wind_gw - SHAO_SOLAR_POOLED * solar_gw
    moe_m = {m: moe_sum[m] / n_m[m] if n_m[m] else 0.0 for m in range(1, 13)}
    model_mean = sum(moe_sum.values()) / max(1, sum(n_m.values()))
    basis = "estimated" if "estimated" in (rbasis, ex.basis) else "modelled"
    kfix, kvom, kdet = kept_costs(year, kept, cost_case) if kept["stations"] else (0.0, 0.0, {})
    # filters' running cost on kept coal's extra output (kept coal's share of the coal block)
    k_ab_vom = kvom * xgen.get("coal", 0.0) * (km["coal"] / ex.coal if ex.coal else 0.0) / 1e6
    # card [internal ticket] "+ex_scr": SCR on existing West Burton A and Drax 5-6, fixed as kept SCR, and its running
    # cost on the two stations' MW share of the coal block's whole output (actual + extra)
    ex_scr_fixed, ex_scr_v = existing_scr_costs(year, coal_lf, cpath)
    coal_twh = lf_got * ex.coal * (len(rows) / 2) / 1e6
    ex_scr_eq = sum(mw * sh for _, _, mw, sh, h in ex_jul if h is None)
    ex_scr_vom = ex_scr_v * coal_twh * (ex_scr_eq / ex.coal if ex.coal else 0.0)
    y = LfLYear(
        kept={**kept, **kdet}, kept_fixed_gbp_m=kfix, kept_abatement_vom_gbp_m=k_ab_vom, closure_test=ct_detail,
        extra_gen_by_block_twh={k: v / 1e6 for k, v in xgen.items()}, closures=closures, newbuild=newbuild,
        nuclear_note=ff.nuclear_available(year)[1],
        coal_lf=coal_lf, coal_lf_cap=cap or 0.0, coal_avail_scale=scale, coal_lf_achieved=lf_got, coal_twh=coal_twh, coal_ied=coal_ied,
        existing_scr_fixed_gbp_m=ex_scr_fixed, existing_scr_vom_gbp_m=ex_scr_vom,
        gas_marginal_share=n_gas_marg / max(1, len(rows)), gas_marginal_share_actual_load=n_gas_marg0 / max(1, len(rows)), extra_gas_fuel_gbp_m=xgas_fuel / 1e6, kept_twh_m=kept_m,
        year=year, market_twh_m=mkt_m, extra_twh=extra_twh, extra_twh_m=xm, energy_twh_m=lm, extra_cost_gbp_m=extra_cost / 1e6,
        extra_fuel_gbp_m=extra_fuel, extra_vom_gbp_m=extra_vom,
        extra_transport_gbp_m=extra_transport, extra_gas_burn_twh=extra_burn,
        new={} if new is None else {"ccgt_gw": new.ccgt_gw, "ocgt_gw": new.ocgt_gw, "spare_gw": new.spare_gw,
                                     "coal_gw": new.coal_gw,
                                     "energy_twh": new.energy_twh, "gen_twh": dict(new.gen_twh),
                                     "fixed_gbp_m": new.fixed_gbp_m, "fuel_gbp_m": sum(new.fuel_gbp_m.values()),
                                     "vom_gbp_m": new.vom_gbp_m, "total_gbp_m": new.total_gbp_m,
                                     "gas_fuel_gbp_m": new.fuel_gbp_m.get("ccgt", 0.0) + new.fuel_gbp_m.get("ocgt", 0.0)
                                     - tv * (new.gen_twh["ccgt"] / ed.ccgt.efficiency_gcv
                                             + new.gen_twh["ocgt"] / ed.ocgt.efficiency_gcv),
                                     "gas_burn_twh": new.gen_twh["ccgt"] / ed.ccgt.efficiency_gcv
                                     + new.gen_twh["ocgt"] / ed.ocgt.efficiency_gcv,
                                     "source_id": new.source_id},
        firm_spare_gw=spare / 1e3, firm_gross_gw=gross / der["ocgt"] / 1e3, unlisted_firm_gw=own / 1e3, peak_actual_fossil_mw=peak_f, own_firm_basis=own_basis, peak_actual_fossil_gross_mw=peak_fg, firm_spare_fixed_gbp_m=fc * 1000 * spare / 1e6,
        moe_model_m=moe_m, moe_vol_f_m=moe_f, moe_vol_x_m=moe_x, moe_vol_k_m=moe_k, moe_model_mean=model_mean, moe_shao_mean=shao,
        shao_note=note, wind_gw_mean=wind_gw, solar_gw_mean=solar_gw, renew_twh=sum(min(l, x) for l, x in zip(L, R)) / 2e6, renew_gross_twh=sum(R) / 2e6,
        actual_fossil_twh=sum(F) / 2e6, energy_twh=sum(L) / 2e6, peak_mw=peak, existing=ex,
        existing_derated_mw=A, gas_burn_existing_twh=gas_burn, basis=basis,
        sizing=sp, firm_req=fr, ic_called_twh_m=called_m, ic_called_twh=sum(called_m.values()),
        ic_called_halfhours=n_called, worked_hh=None, m8_sums=m8,
        coal_twh_m=coal_m, coal_moe_vol_m=coal_moe, coal_m8=coal_m8,
        hh_cap=tuple(hh_cap) if keep_hh else None,
        hh_debug={"stacks": stacks, "new_srmc": new_srmc, "new_ccgt_mw": new_ccgt_mw} if keep_hh else None)
    return replace(y, worked_hh=_worked_hh(hh_cap, moe_factor(y, "shao")))


def moe_factor(y: LfLYear, rule: str) -> tuple[float, float]:
    """(multiplier on the model's Δ shape, flat add-on £/MWh) for the MOE rule."""
    if rule not in MOE_RULES:
        raise ValueError(f"unknown MOE rule {rule!r}; one of {MOE_RULES}")
    mod, sh = y.moe_model_mean, max(0.0, y.moe_shao_mean)
    if rule == "none":
        return 0.0, 0.0
    if rule == "model":
        return 1.0, 0.0
    target = {"shao": sh, "larger": max(mod, sh), "smaller": min(mod, sh)}[rule]
    if mod > 1e-9:
        return target / mod, 0.0
    return 0.0, target


# ----------------------------------------------------------------- s27a: gas-price feedback

@lru_cache(maxsize=None)
def gb_gas_demand_twh() -> dict[int, float]:
    """GB (UK) natural gas total demand, TWh by year: Energy Trends 4.1 'Annual (GWh)',
    row 'Total demand' (stored data/raw/s26a/s26a-ET_4.1_SEP_26.xlsx)."""
    ws = openpyxl.load_workbook(ET_41, read_only=True, data_only=True)["Annual (GWh)"]
    rows = list(ws.iter_rows(values_only=True))
    h = next(r for r in rows if r[0] == "Column1")
    yrs = {j: int(str(v)[:4]) for j, v in enumerate(h) if v and str(v)[:4].isdigit()}
    r = next(r for r in rows if str(r[0] or "").strip() == "Total demand")
    return {y: float(r[j]) / 1e3 for j, y in yrs.items() if isinstance(r[j], (int, float))}


@lru_cache(maxsize=None)
def actual_power_gas_burn_twh(year: int) -> float:
    """Actual gas burnt by major power producers, TWh: DUKES 5.6.A 'Fuel used in generation'."""
    rows56 = rn._sheet(rn.DUKES_56, "5.6")
    t = next(i for i, r in enumerate(rows56[:-1]) if str(r[0] or "").startswith("Table 5.6.A")
             and str(rows56[i + 1][0] or "").startswith("Generation type"))
    hdr = rows56[t + 1]
    j = [i for i, x in enumerate(hdr) if x is not None and str(x)[:4] == str(year)][0]
    r = next(r for r in rows56[t + 2:] if str(r[0] or "").startswith("Major power producers")
             and str(r[1] or "").strip() == "Gas")
    return float(r[j]) / 1e3


# Card [internal ticket]: M6 (0163 R7 (a)) both ways is the default; the v4-s37/v4-s38 floored dQ is an explicit setting,
# $GASCF_M6_FEEDBACK=floored, for guards that reproduce frozen releases cut before v4-s39. Not a range or methods row.
M6_ENV, M6_OPTIONS, M6_DEFAULT = "GASCF_M6_FEEDBACK", ("both", "floored"), "both"


def m6_reading() -> str:
    v = os.environ.get(M6_ENV, M6_DEFAULT)
    if v not in M6_OPTIONS:
        raise ValueError(f"${M6_ENV}={v!r}: one of {M6_OPTIONS}")
    return v


def gas_feedback(y: LfLYear, rule: str = DEFAULT_GAS_FEEDBACK) -> dict:
    """Spec §15.5 item 2: fossil Britain burns more gas, so its gas price is higher.

    dP/P = dQ / ((e_s + e_d) x Q) (decision 0144 item 2A; Albrizio et al. 2022 equation 5): by rule
    (`GAS_ELASTICITY_PAIRS`): central = Albrizio's own pair, supply 0.06 (Krichene (2005) Table 5 (printed p.20,
    PDF p.21), short-run, as Albrizio use it) + demand 0.169 (Albrizio Table 2, EU-weighted) = 0.229;
    krichene_lr = Krichene (2005) Table 2 world long-run pair 0.16 + 0.45 = 0.61 (the sensitivity row); high =
    0.06 supply only, kept for case B (0048, open); see the comment above GAS_FEEDBACK_RULES. Q = EU historic demand
    4,000 TWh (Neuhoff Table 2) + GB gas demand that year (ET 4.1), dQ = fossil Britain's gas
    burn for power (existing plant at fossil Britain's load + new plant) less the actual gas
    burnt by major power producers (DUKES 5.6.A), signed: negative where fossil Britain burns less gas, so its
    gas price is lower (RB v1.1 M6, decision 0163 R7 (a); was floored at 0 to v4-s38) (s27c fix: was measured against
    the engine's coal-first dispatch of the actual load, which double-counted). Applied to (a) the at-cost gas fuel (existing
    plant's extra running and the new plant), and (b) every market-priced MWh (actual fossil
    output and the kept nuclear, imports and pumped storage) x the share of half-hours gas
    plant is marginal x gas price / CCGT efficiency. The dispatch is not re-run with the
    dearer gas (second order; stated). Returns £m and the inputs."""
    if rule not in GAS_FEEDBACK_RULES:
        raise ValueError(f"unknown gas-feedback rule {rule!r}; one of {GAS_FEEDBACK_RULES}")
    # s27c fix: dQ against the ACTUAL gas burn (DUKES 5.6.A, major power producers, fuel TWh), as
    # the s26a design and the s25a review sized it. The engine's own dispatch of the actual fossil
    # load (g0) runs fossil Britain's kept coal first and burns far less gas than the real system
    # did (2025: 52.9 TWh vs DUKES 168.7), so `extra_gas_burn_twh` (g1 - g0) also counts coal
    # displacing gas inside the actual load; it is right for the extra-running COST (the actual
    # output is market-priced) but not for the change in gas DEMAND the market sees.
    fossil_burn = y.gas_burn_existing_twh + (y.new or {}).get("gas_burn_twh", 0.0)
    actual_burn = actual_power_gas_burn_twh(y.year)
    # RB v1.1 M6 (decision 0163, R7 (a), Richard Lyon 21 Oct 2026): the feedback works both ways. dQ is the signed
    # difference; where fossil Britain burns less gas than the actual system, its gas price is lower (was floored at 0).
    m6 = m6_reading()
    dq = fossil_burn - actual_burn
    if m6 == "floored":          # v4-s37/v4-s38 reading, explicit only (card [internal ticket])
        dq = max(0.0, dq)
    d = gb_gas_demand_twh()
    yy = y.year if y.year in d else (min(d) if y.year < min(d) else max(d))
    if yy != y.year:
        fb.record("gb_gas_demand_nearest_year", y.year, d[yy], "TWh", f"ET 4.1 total demand {yy}")
    q = EU_GAS_DEMAND_TWH + d[yy]
    e_s, e_d = GAS_ELASTICITY_PAIRS.get(rule, (None, None))
    out = {"rule": rule, "elasticity": GAS_ELASTICITY.get(rule), "elasticity_supply": e_s, "elasticity_demand": e_d,
           "dq_twh": dq, "q_twh": q,
           "actual_power_gas_twh": actual_burn, "fossil_power_gas_twh": fossil_burn,
           "dq_vs_engine_actual_dispatch_twh": y.extra_gas_burn_twh + (y.new or {}).get("gas_burn_twh", 0.0),
           "gas_marginal_share": y.gas_marginal_share,
           "gas_marginal_share_actual_load": y.gas_marginal_share_actual_load,
           "source_ids": [SRC_NEUHOFF, SRC_ET41, SRC_56]}
    if rule == "none":
        return {**out, "price_rise_frac": 0.0, "at_cost_gbp_m": 0.0, "market_gbp_m": 0.0, "kept_gbp_m": 0.0}
    frac = dq / (GAS_ELASTICITY[rule] * q)
    at_cost = frac * (y.extra_gas_fuel_gbp_m + (y.new or {}).get("gas_fuel_gbp_m", 0.0))
    eta = _eta_dukes("ccgt", y.year)
    per_mwh = {m: frac * y.gas_marginal_share * gas_price(y.year, m) / eta for m in range(1, 13)}
    mkt = sum(per_mwh[m] * y.market_twh_m[m] for m in range(1, 13))
    kept = sum(per_mwh[m] * (y.kept_twh_m or {}).get(m, 0.0) for m in range(1, 13))
    return {**out, "price_rise_frac": frac, "at_cost_gbp_m": at_cost, "market_gbp_m": mkt, "kept_gbp_m": kept,
            "market_mean_gbp_mwh": mkt / y.actual_fossil_twh if y.actual_fossil_twh else 0.0}
