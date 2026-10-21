"""The "Actual" column of the household electricity receipt, every Ofgem cap period.

What this gives you: for each price-cap period from April 2015 to October-December 2026,
the GB-average annual electricity bill of a single-rate, Direct Debit ("Other Payment
Method") household at Ofgem's own benchmark consumption, split into the spec §6 Layer 1
lines, in £ per household per year, plus the published cap level the lines must add up to.

Everything is read from Ofgem's own workbooks in ``data/raw/ofgem/``; nothing is typed in:

* Annex 9 v1.11 (levelisation, published 2026-08), tab ``1b Historical level tables``:
  Ofgem's GB-average cap level by component, Apr 2017 onwards, *after* the levelisation
  allowance. This is the published Direct Debit cap level, so it is the reconciliation target.
* Cap level model v1.31, tabs ``ElecSingle_Other_Benchmark`` / ``ElecSingle_Other_Nil``
  rows 225-239 ("GB average"): the same components for the four 2015-17 periods, which
  Ofgem shows "for illustration only" (no cap was in force).
* Cap level model tab ``3a DF``: the direct fuel component split into direct fuel,
  backwardation allowance and CfD allowance (regional rows; GB = mean of 14 regions).
* Annex 4 v1.25 tab ``1a Policy Cost Allowance`` rows 53-60: GB scheme-by-scheme policy
  costs (RO, FiT, ECO, WHD, AAHEDC, NCC = Energy Intensive Industries levy, nRAB = Sizewell C).
* Annex 3 v1.25 tabs ``2a TNUoS``, ``2b DUoS``, ``2c BSUoS``: the network allowance split.

Basis tags (review B2). Every cap-component line is tagged ``allowance``: an Ofgem *ex-ante*
allowance set before the period, not an outturn. In 2022-23 the CfD allowance even had the
opposite sign to the levy LCCC actually charged. So each receipt also carries the LCCC
outturn for CfD (``cfd_outturn``: reconciled daily levy, latest run, £/MWh of eligible demand,
from goods.cfd_outturn), for the staircase to use in history. VAT on Annex 9 periods is
Ofgem's published inc-VAT level, so it is ``allowance`` too; for 2015-17 it is our statutory
5% (``estimated``).

Residuals (review S2). Two lines are fitted to Ofgem's totals: AAHEDC absorbs the sub-penny
gap between Annex 4's GB scheme estimates and the cap's PC, and DUoS absorbs the sub-penny gap
between Annex 3 v1.25 and the cap model's NC. Each fitted amount is kept on the receipt in
``residuals`` and tested to be under 1p.

Terms: "nil" = Ofgem's cap at zero consumption (the standing-charge part, £/yr);
"typical" = the cap at Ofgem's benchmark consumption Q (MWh/yr). The unit-rate part of a
line is typical − nil, and scales with consumption. "GB average" is Ofgem's unweighted
mean of its 14 charge-restriction regions (the cap model's GB rows are literally
``AVERAGE`` of the 14 regional rows).
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field, replace
from functools import lru_cache
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

RAW = Path(__file__).resolve().parents[2] / "data/raw/ofgem"
CAP_MODEL = RAW / "Default-tariff-cap-level-v1.31.xlsx"
ANNEX9 = RAW / "Annex-9-Levelisation-allowance-methodology-and-levelised-cap-levels-v1.11.xlsx"
ANNEX4 = RAW / "Annex-4-policy-cost-allowance-methodology-v1.25.xlsx"
ANNEX3 = RAW / "Annex-3-network-cost-allowance-methodology-electricity-v1.25.xlsx"

SRC_A9 = "ofgem-annex9-v1.11"
SRC_CAP = "ofgem-cap-model-v1.31"
SRC_A4 = "ofgem-annex4-v1.25"
SRC_A3 = "ofgem-annex3-elec-v1.25"

REGIONS = (
    "Eastern", "East Midlands", "London", "N Wales and Mersey", "Midlands", "Northern",
    "North West", "Southern", "South East", "South Wales", "Southern Western", "Yorkshire",
    "Southern Scotland", "Northern Scotland",
)

# Ofgem component codes in tab 1b, in sheet order (Annex 9 adds "Levelisation ").
CODES = ("DF", "CM", "AA", "PC", "NC", "OC", "SMNCC", "IC", "PAAC", "PAP", "CO", "DRC",
         "EBIT", "HAP")

# Receipt lines: key -> (label, spec §6 group). Order is the order on the receipt.
LINES = {
    "direct_fuel": ("Wholesale: direct fuel", "wholesale"),
    "backwardation": ("Wholesale: backwardation allowance", "wholesale"),
    "cfd": ("Contracts for Difference (in Ofgem's wholesale allowance)", "policy"),
    "capacity_market": ("Capacity Market", "wholesale"),
    # Ofgem's own words (cap model tab 3c AA, cell B3): "Adjustment Allowance ... calculated
    # according to Annex 8 to the licence conditions". No gloss (review S7).
    "adjustment_allowance": ("Adjustment Allowance (Ofgem cap component AA, licence Annex 8)", "other"),
    "ro": ("Renewables Obligation", "policy"),
    "fit": ("Feed-in Tariffs", "policy"),
    "eco": ("Energy Company Obligation (ECO)", "policy"),
    "whd": ("Warm Home Discount", "policy"),
    "aahedc": ("AAHEDC (Hydro Benefit Replacement Scheme)", "policy"),
    "eii_levy": ("Energy Intensive Industries levy (Ofgem 'NCC')", "policy"),
    "sizewell_c": ("Sizewell C nuclear RAB levy (Ofgem 'nRAB')", "policy"),
    "tnuos": ("Network: transmission (TNUoS)", "network"),
    "duos": ("Network: distribution (DUoS)", "network"),
    "bsuos": ("Network: balancing services (BSUoS)", "network"),
    "operating_costs": ("Operating Cost Allowance (to Jun 2025)", "supplier"),
    "core_operating_costs": ("Core Operating Cost Allowance (from Jul 2025)", "supplier"),
    "smart_meter": ("Smart Metering Net Cost Change", "supplier"),
    "industry_charges": ("Industry Charges", "supplier"),
    "payment_adjustment_cost": ("Payment Method Adjustment Additional Cost", "supplier"),
    "payment_adjustment_pct": ("Payment Method Adjustment Percentage", "supplier"),
    "debt_related_costs": ("Debt Related Cost Allowance", "supplier"),
    "ebit": ("EBIT (supplier margin)", "supplier"),
    "headroom": ("Headroom allowance", "supplier"),
    "levelisation": ("Levelisation allowance (Ofgem Annex 9, from Apr 2024)", "other"),
    "vat": ("VAT", "tax"),
}

# Lines Ofgem reports directly in tab 1b (key -> 1b code).
DIRECT = {
    "capacity_market": "CM", "adjustment_allowance": "AA", "operating_costs": "OC",
    "smart_meter": "SMNCC", "industry_charges": "IC", "payment_adjustment_cost": "PAAC",
    "payment_adjustment_pct": "PAP", "core_operating_costs": "CO",
    "debt_related_costs": "DRC", "ebit": "EBIT", "headroom": "HAP",
}

AA_ANCHOR = "calculated according to Annex 8 to the licence conditions"
BASIS_ALLOWANCE = "allowance"   # Ofgem ex-ante allowance, not outturn (review B2)
BASES = ("allowance", "measured", "modelled", "estimated")

_MON = {m: i + 1 for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}


@dataclass(frozen=True)
class Line:
    key: str
    label: str
    group: str
    typical: float          # £/household/yr at Ofgem's consumption Q
    nil: float              # £/household/yr at zero consumption (the fixed part)
    basis: str              # allowance | measured | modelled | estimated
    source_id: str
    cell: str               # where it was read


@dataclass(frozen=True)
class Receipt:
    period: tuple[int, int, int, int]   # (start year, start month, end year, end month)
    label: str                          # Ofgem's label, normalised
    consumption_mwh: float              # Ofgem's benchmark consumption Q
    vat_rate: float
    lines: tuple[Line, ...]
    published_ex_vat: float             # Ofgem's GB-average total, ex VAT
    published_inc_vat: float
    published_note: str                 # what "published" means for this period
    nil_published_ex_vat: float
    cap_in_force: bool                  # True for Oct 2018-Mar 2019: half-year containing the
                                        # first cap quarter (cap from 1 Jan 2019); see published_note
    epg_in_force: bool                  # Energy Price Guarantee, Oct 2022 - Jun 2023
    notes: tuple[str, ...] = field(default=())
    # LCCC CfD outturn for the period (goods.cfd_outturn), None if any day is not yet priced
    cfd_outturn: dict | None = None
    # amounts fitted into a plug line so lines add to Ofgem's totals, £/yr at Ofgem's Q
    residuals: dict[str, float] = field(default_factory=dict)

    def cfd_outturn_typical(self) -> float | None:
        """LCCC CfD outturn, £/household/yr at this receipt's consumption (no loss uplift)."""
        if self.cfd_outturn is None:
            return None
        return self.cfd_outturn["gbp_mwh"] * self.consumption_mwh

    def line(self, key: str) -> Line:
        for ln in self.lines:
            if ln.key == key:
                return ln
        raise KeyError(key)

    @property
    def total_ex_vat(self) -> float:
        return sum(ln.typical for ln in self.lines if ln.key != "vat")

    @property
    def total(self) -> float:
        return sum(ln.typical for ln in self.lines)

    def unit_rate(self, key: str) -> float:
        """£/MWh of the unit-rate part of a line at Ofgem's consumption."""
        ln = self.line(key)
        return (ln.typical - ln.nil) / self.consumption_mwh


# --------------------------------------------------------------------------- helpers

def _norm(s) -> str:
    return " ".join(str(s).replace("–", "-").replace("\n", " ").split())


def period_key(label) -> tuple[int, int, int, int]:
    """(y1, m1, y2, m2) from any of Ofgem's period labels."""
    s = _norm(label).lower()
    found = re.findall(r"(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*(\d{4})?", s)
    if len(found) < 2:
        raise ValueError(f"not a period label: {label!r}")
    (m1, y1), (m2, y2) = found[0], found[-1]
    y2 = int(y2)
    m1, m2 = _MON[m1], _MON[m2]
    y1 = int(y1) if y1 else (y2 if m1 <= m2 else y2 - 1)
    return (y1, m1, y2, m2)


def label_of(key: tuple[int, int, int, int]) -> str:
    names = list(_MON)
    y1, m1, y2, m2 = key
    return f"{names[m1 - 1].title()} {y1} - {names[m2 - 1].title()} {y2}"


def _num(v) -> float:
    """A component cell: Ofgem writes '-' (or leaves blank) when a component does not exist."""
    if isinstance(v, (int, float)):
        return float(v)
    if v is None or (isinstance(v, str) and v.strip() == "-"):
        return 0.0
    raise ValueError(f"expected a number or '-', found {v!r}")


def _expect(actual, expected, where):
    if _norm(actual) != _norm(expected):
        raise ValueError(f"layout moved: {where} holds {actual!r}, expected {expected!r}")


def _cols(row, start=0) -> dict[tuple, int]:
    """Period key -> column index, for a header row of period labels (first wins)."""
    out = {}
    for j, v in enumerate(row):
        if j < start or not isinstance(v, str):
            continue
        try:
            k = period_key(v)
        except ValueError:
            continue
        out.setdefault(k, j)
    return out


def _book(path: Path):
    if not path.exists():
        raise FileNotFoundError(f"{path} missing - copy it into data/raw/ofgem/")
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def _rows(wb, tab, max_col=60):
    return list(wb[tab].iter_rows(values_only=True, max_col=max_col))


# --------------------------------------------------------------------------- readers

@lru_cache(maxsize=1)
def _annex9():
    """Annex 9 tab 1b: Other Payment Method, single-rate, nil (rows 13-29) and typical (32-48)."""
    R = _rows(_book(ANNEX9), "1b Historical level tables", 35)
    _expect(R[7][1], "Other Payment Method", "A9 1b B8")
    _expect(R[9][1], "Electricity: Single-Rate Metering Arrangement", "A9 1b B10")
    out = {}
    for hdr, name in ((11, "nil"), (30, "typical")):
        _expect(R[hdr][1], "Nil consumption" if name == "nil" else "Typical consumption",
                f"A9 1b B{hdr + 1}")
        codes = [_norm(R[hdr + 1 + k][1]) for k in range(17)]
        want = list(CODES) + ["Levelisation", "Total_GB average", "Total inc VAT"]
        if codes != want:
            raise ValueError(f"A9 1b {name} rows moved: {codes}")
        for k, j in _cols(R[hdr], 2).items():
            if j > 30:
                continue
            vals = {c: R[hdr + 1 + i][j] for i, c in enumerate(want)}
            out.setdefault(k, {})[name] = vals
            out[k][name + "_col"] = j
    return out


@lru_cache(maxsize=1)
def _cap_model():
    wb = _book(CAP_MODEL)
    res = {}
    for tab, name in (("ElecSingle_Other_Benchmark", "typical"), ("ElecSingle_Other_Nil", "nil")):
        R = _rows(wb, tab, 50)
        _expect(R[11][4] if R[11][4] else R[11][1], "28AD Charge Restriction Period:", f"{tab} row 12")
        names = [_norm(R[224 + i][2]) for i in range(15)]
        want = ["DF", "CM", "AA", "PC", "NC", "OC", "SMNCC", "IC", "PAAC", "PAP", "CO", "DRC",
                "E", "H", "Total_GB average"]
        if names != want or _norm(R[224][3]) != "GB average":
            raise ValueError(f"{tab} GB-average rows 225-239 moved: {names}")
        for k, j in _cols(R[11], 6).items():
            res.setdefault(k, {})[name] = {c: R[224 + i][j] for i, c in enumerate(want)}
    # benchmark consumption table (Notes rows 163-164)
    N = _rows(wb, "Notes", 8)
    _expect(N[161][1], "Typical consumption (MWh) effective dates", "Notes B162")
    _expect(N[163][1], "Electricity - Single-Rate Metering Arrangement", "Notes B164")
    cons = {"hdr": [N[162][j] for j in (2, 3, 4)], "mwh": [float(N[163][j]) for j in (2, 3, 4)]}
    # 3a DF: direct fuel (rows 11-24), backwardation (48-61), CfD (85-98), combined (119-132)
    D = _rows(wb, "3a DF", 58)
    _expect(D[9][1], "4. Direct Fuel Cost Component at Benchmark Annual Consumption Level m kWh (typical consumption)", "3a DF B10")
    _expect(D[46][1], "Backwardation allowance (Typical consumption)", "3a DF B47")
    _expect(D[83][1], "CFD allowance", "3a DF B84")
    split = {}
    for name, first, hdr in (("direct", 10, 6), ("backwardation", 47, 43), ("cfd", 84, 80),
                             ("combined", 118, 115)):
        _expect(D[first][1], "Electricity", f"3a DF B{first + 1}")
        _expect(D[first][2], "Single-Rate Metering Arrangement", f"3a DF C{first + 1}")
        regs = [_norm(D[first + i][3]) for i in range(14)]
        if sorted(regs) != sorted(REGIONS):
            raise ValueError(f"3a DF {name} regions moved: {regs}")
        for k, j in _cols(D[hdr], 7).items():
            vals = [D[first + i][j] for i in range(14)]
            if all(isinstance(v, (int, float)) for v in vals):
                split.setdefault(k, {})[name] = sum(vals) / 14
    A = _rows(wb, "3c AA", 4)
    _expect(A[1][1], "Adjustment Allowance", "3c AA B2")
    if AA_ANCHOR not in _norm(A[2][1]):
        raise ValueError("3c AA B3 no longer cites Annex 8; check the adjustment_allowance label")
    res["_cons"] = cons
    res["_split"] = split
    return res


@lru_cache(maxsize=1)
def _annex4():
    """GB scheme-by-scheme policy costs, Annex 4 tab 1a rows 53-60, plus the consumption cell used."""
    wb = _book(ANNEX4)
    R = _rows(wb, "1a Policy Cost Allowance", 45)
    _expect(R[44][1][:42], "2. Scheme by scheme estimates (GB average)", "A4 1a B45")
    _expect(R[52][1], "Electricity - Single-Rate Metering Arrangement", "A4 1a B53")
    schemes = [(_norm(R[52 + i][2]), _norm(R[52 + i][3])) for i in range(8)]
    want = [("RO", "£/MWh supplied"), ("FiT", "£/MWh supplied"), ("ECO", "£/MWh supplied"),
            ("WHD", "£/customer"), ("WHD", "£/MWh supplied"), ("AAHEDC (GB average)", "£/MWh supplied"),
            ("NCC (GB average)", "£/MWh supplied"), ("nRAB (GB average)", "£/MWh supplied")]
    if schemes != want:
        raise ValueError(f"A4 1a scheme rows moved: {schemes}")
    out = {}
    # row 12 carries the 28AD period labels (row 50's copy has a typo, "January 2023 - March 2022")
    for k, j in _cols(R[11], 6).items():
        out[k] = {"RO": R[52][j], "FiT": R[53][j], "ECO": R[54][j], "WHD_cust": R[55][j],
                  "WHD_mwh": R[56][j], "AAHEDC": R[57][j], "NCC": R[58][j], "nRAB": R[59][j]}
    # consumption: the '3a Demand' cell each row-15 formula multiplies by
    F = list(openpyxl.load_workbook(ANNEX4, read_only=True)["1a Policy Cost Allowance"]
             .iter_rows(min_row=15, max_row=15, max_col=45, values_only=True))[0]
    dem = {r: R2 for r, R2 in _demand(wb).items()}
    cons = {}
    for k, j in _cols(R[11], 6).items():
        m = re.search(r"'3a Demand'!\$?([A-Z]+)\$?(\d+)", str(F[j]))
        if m:
            cons[k] = dem[m.group(1) + m.group(2)]
    return out, cons


def _demand(wb) -> dict[str, float]:
    R = _rows(wb, "3a Demand", 14)
    _expect(R[8][1], "Electricity - Single-Rate Metering Arrangement", "A4 3a Demand B9")
    return {"C9": float(R[8][2]), "H9": float(R[8][7]), "L9": float(R[8][11])}


@lru_cache(maxsize=1)
def _annex3():
    """GB-mean TNUoS / DUoS (nil and typical) and BSUoS (typical) per period."""
    wb = _book(ANNEX3)
    out = {}
    for tab, key in (("2a TNUoS", "tnuos"), ("2b DUoS", "duos"), ("2c BSUoS", "bsuos")):
        R = _rows(wb, tab, 58)
        _expect(R[7][6], "28AD Charge Restriction Period:", f"A3 {tab} G8")
        blocks = {"typical": 10} if key == "bsuos" else {"nil": 10, "typical": 24}
        for name, first in blocks.items():
            _expect(R[first][1] if name == "nil" or key == "bsuos" else R[10][1],
                    "Single-Rate Metering Arrangement", f"A3 {tab} B{first + 1}")
            lvl = _norm(R[first][2])
            if (name == "nil") != (lvl == "Nil"):
                raise ValueError(f"A3 {tab} row {first + 1} is {lvl!r}, expected {name}")
            regs = [_norm(R[first + i][5]) for i in range(14)]
            if regs != list(REGIONS):
                raise ValueError(f"A3 {tab} regions moved: {regs}")
            for k, j in _cols(R[7], 8).items():
                vals = [R[first + i][j] for i in range(14)]
                out.setdefault(k, {}).setdefault(key, {})[name] = (
                    sum(_num(v) for v in vals) / 14)
    return out


# --------------------------------------------------------------------------- assembly

def _lookup(d: dict, k: tuple, what: str):
    """Value for cap period k: exact column if there is one, else the narrowest source
    column whose period contains k (Annex 3 network charges are set per charging year, so
    e.g. Oct-Dec 2022 sits in Annex 3's "October 2022 - March 2023" column)."""
    if k in d:
        return d[k]
    start, end = (k[0], k[1]), (k[2], k[3])
    hits = [c for c in d if isinstance(c, tuple) and (c[0], c[1]) <= start and end <= (c[2], c[3])]
    if not hits:
        raise KeyError(f"{what}: no column covers {label_of(k)}")
    return d[min(hits, key=lambda c: (c[2] - c[0]) * 12 + c[3] - c[1])]


def _consumption_for(k, cons_table) -> float:
    """Ofgem's benchmark consumption for a period (cap model Notes rows 163-164)."""
    hdr, mwh = cons_table["hdr"], cons_table["mwh"]
    # "April 2018 - December 2025" | "January 2026 -June 2026" | "July 2026 onwards"
    end1, end2 = period_key(hdr[0]), period_key(hdr[1])
    start = (k[0], k[1])
    if start <= (end1[2], end1[3]):
        return mwh[0]         # cap model C7 note: all periods before P15b use 3,100 kWh
    if start <= (end2[2], end2[3]):
        return mwh[1]
    return mwh[2]


def _vat_rate(ex, inc) -> float:
    r = round(inc / ex - 1, 4)
    if r not in (0.0, 0.05):
        raise ValueError(f"unexpected VAT rate {r}")
    return r


def _build(k) -> Receipt:
    cm, a9 = _cap_model(), _annex9()
    a4, a4cons = _annex4()
    a3 = _annex3()
    a3_by = {key: {kk: v[key] for kk, v in a3.items() if key in v} for key in ("tnuos", "duos", "bsuos")}
    split = _lookup(cm["_split"], k, "3a DF")
    Q = _consumption_for(k, cm["_cons"])
    if k in a4cons and abs(a4cons[k] - Q) > 1e-9:
        raise ValueError(f"{label_of(k)}: Annex 4 uses {a4cons[k]} MWh, cap model {Q}")
    notes = []
    if k in a9:
        col = a9[k]["typical_col"]
        typ = {c: _num(v) for c, v in a9[k]["typical"].items()}
        nil = {c: _num(v) for c, v in a9[k]["nil"].items()}
        src, where = SRC_A9, f"Annex 9 1b col {get_column_letter(col + 1)}"
        pub_ex, pub_inc = typ["Total_GB average"], typ["Total inc VAT"]
        nil_ex = nil["Total_GB average"]
        cap_in_force = k >= (2019, 1, 2019, 3) or k == (2018, 10, 2019, 3)
        if k == (2018, 10, 2019, 3):
            note = ("Ofgem GB-average levelised cap, Direct Debit, single rate; half-year "
                    "containing the first cap quarter (cap in force from 1 Jan 2019 only)")
        elif cap_in_force:
            note = "Ofgem GB-average levelised cap, Direct Debit, single rate"
        else:
            note = "Ofgem's notional cap-basis level (no cap in force before Jan 2019)"
        vat = _vat_rate(pub_ex, pub_inc)
    else:
        t, n = cm[k]["typical"], cm[k]["nil"]
        ren = {"E": "EBIT", "H": "HAP"}
        typ = {ren.get(c, c): _num(v) for c, v in t.items()}
        nil = {ren.get(c, c): _num(v) for c, v in n.items()}
        typ["Levelisation"] = nil["Levelisation"] = 0.0
        src, where = SRC_CAP, "cap model ElecSingle_Other_* rows 225-239"
        pub_ex = typ["Total_GB average"]
        nil_ex = nil["Total_GB average"]
        vat = 0.05
        pub_inc = pub_ex * (1 + vat)
        cap_in_force = False
        note = ("Ofgem's illustrative cap-basis level, 'historical examples for illustration "
                "only' (no cap in force); VAT 5% applied by us")
        notes.append("VAT rate 5% is the statutory reduced rate on domestic fuel (VATA 1994 "
                     "Sch 7A); Ofgem's workbook does not state VAT for these periods.")
    lines: list[Line] = []

    def add(key, typical, nil_, basis, source, cell):
        label, group = LINES[key]
        lines.append(Line(key, label, group, float(typical), float(nil_), basis, source, cell))

    # wholesale: split DF into direct / backwardation / CfD (3a DF, GB mean)
    if abs(split["combined"] - typ["DF"]) > 1e-6:
        raise ValueError(f"{label_of(k)}: 3a DF combined {split['combined']} != 1b DF {typ['DF']}")
    parts = {n: split.get(n, 0.0) for n in ("direct", "backwardation", "cfd")}
    if abs(sum(parts.values()) - typ["DF"]) > 1e-6:
        raise ValueError(f"{label_of(k)}: DF split does not add up")
    A = BASIS_ALLOWANCE
    add("direct_fuel", parts["direct"], 0, A, SRC_CAP, "3a DF rows 11-24")
    add("backwardation", parts["backwardation"], 0, A, SRC_CAP, "3a DF rows 48-61")
    add("cfd", parts["cfd"], 0, A, SRC_CAP, "3a DF rows 85-98 (Ofgem ex-ante allowance)")
    add("capacity_market", typ["CM"], nil["CM"], A, src, where + " CM")
    add("adjustment_allowance", typ["AA"], nil["AA"], A, src, where + " AA (Annex 8)")

    # policy: Annex 4 GB scheme estimates x Q, WHD per customer is the nil part
    p = {n: _num(v) for n, v in _lookup(a4, k, "Annex 4").items()}
    pol = {
        "ro": p["RO"] * Q, "fit": p["FiT"] * Q, "eco": p["ECO"] * Q,
        "whd": p["WHD_cust"] + p["WHD_mwh"] * Q, "aahedc": p["AAHEDC"] * Q,
        "eii_levy": p["NCC"] * Q, "sizewell_c": p["nRAB"] * Q,
    }
    gap = typ["PC"] - sum(pol.values())
    pol_basis = A
    residuals = {"policy_into_aahedc": gap}
    if abs(gap) > 0.01:
        raise ValueError(f"{label_of(k)}: Annex 4 schemes {sum(pol.values()):.4f} != PC {typ['PC']:.4f}")
    # the GB-average schemes differ from the 14-region mean of PC only by AAHEDC/NCC/nRAB
    # losses-factor averaging; put the (sub-penny) residual into AAHEDC so lines add exactly.
    pol["aahedc"] += gap
    if abs(gap) > 1e-9:
        notes.append(f"policy: {gap:+.5f} £ regional-averaging residual folded into AAHEDC")
    whd_nil = p["WHD_cust"]
    if abs(whd_nil - nil["PC"]) > 1e-6:
        raise ValueError(f"{label_of(k)}: WHD £/customer {whd_nil} != nil PC {nil['PC']}")
    for key in ("ro", "fit", "eco", "whd", "aahedc", "eii_levy", "sizewell_c"):
        add(key, pol[key], whd_nil if key == "whd" else 0.0, pol_basis, SRC_A4,
            "Annex 4 1a rows 53-60 x Q")

    # network: Annex 3 GB means
    n3 = {key: _lookup(a3_by[key], k, f"Annex 3 {key}") for key in ("tnuos", "duos", "bsuos")}
    net_t = {"tnuos": n3["tnuos"]["typical"], "duos": n3["duos"]["typical"],
             "bsuos": n3["bsuos"]["typical"]}
    net_n = {"tnuos": n3["tnuos"]["nil"], "duos": n3["duos"]["nil"], "bsuos": 0.0}
    for d, tot, nm in ((net_t, typ["NC"], "typical"), (net_n, nil["NC"], "nil")):
        resid = tot - sum(d.values())
        # Ofgem's Annex 3 v1.25 and cap model v1.31 disagree by <1p in one period
        # (Apr-Jun 2023, 3e NC-Elec 206.197 vs Annex 3 1a 206.189 for Eastern). The cap
        # model is what sets the cap, so the sub-penny residual goes into DUoS, noted.
        if abs(resid) > 0.01:
            raise ValueError(f"{label_of(k)}: Annex 3 {nm} {sum(d.values())} != NC {tot}")
        residuals[f"network_{nm}_into_duos"] = resid
        if abs(resid) > 1e-9:
            d["duos"] += resid
            notes.append(f"network ({nm}): {resid:+.5f} £ Annex 3 v1.25 vs cap model v1.31 "
                         "difference folded into DUoS")
    for key in ("tnuos", "duos", "bsuos"):
        add(key, net_t[key], net_n[key], A, SRC_A3, f"Annex 3 2{'abc'['tnuos duos bsuos'.split().index(key)]} GB mean")

    for key, code in DIRECT.items():
        if key == "capacity_market" or key == "adjustment_allowance":
            continue
        add(key, typ[code], nil[code], A, src, f"{where} {code}")
    add("levelisation", typ["Levelisation"], nil["Levelisation"], A, src,
        f"{where} Levelisation")
    order = list(LINES)
    lines.sort(key=lambda ln: order.index(ln.key))
    ex = sum(ln.typical for ln in lines)
    ex_nil = sum(ln.nil for ln in lines)
    add("vat", ex * vat, ex_nil * vat, A if k in a9 else "estimated",
        src, f"{where} Total inc VAT / Total" if k in a9 else "statutory 5%")
    return Receipt(
        period=k, label=label_of(k), consumption_mwh=Q, vat_rate=vat, lines=tuple(lines),
        published_ex_vat=pub_ex, published_inc_vat=pub_inc, published_note=note,
        nil_published_ex_vat=nil_ex, cap_in_force=cap_in_force,
        epg_in_force=(2022, 10) <= (k[0], k[1]) <= (2023, 4), notes=tuple(notes),
        cfd_outturn=_cfd_outturn(k), residuals=residuals)


def period_dates(k) -> tuple:
    """(first day, last day) of cap period k."""
    from datetime import date, timedelta
    y2, m2 = (k[2] + 1, 1) if k[3] == 12 else (k[2], k[3] + 1)
    return date(k[0], k[1], 1), date(y2, m2, 1) - timedelta(days=1)


def _cfd_outturn(k) -> dict | None:
    """LCCC outturn beside Ofgem's CfD allowance (review B2). The LCCC levy is £/MWh of
    eligible demand; Ofgem's allowance also uplifts for losses, which we do not apply (this
    keeps the outturn line, and so the renewables side, smaller)."""
    from .goods import cfd_outturn
    return cfd_outturn(*period_dates(k))


@lru_cache(maxsize=1)
def receipts() -> tuple[Receipt, ...]:
    """Every cap period, Apr 2015 to Oct-Dec 2026, in date order."""
    cm, a9 = _cap_model(), _annex9()
    keys = set(a9)
    # 2015-17 periods: in the cap model but before Annex 9's first column (Apr 2017)
    first = min(a9)
    keys |= {k for k in cm if isinstance(k, tuple) and k < first}
    # the cap model also has "Jan 2019 - Mar 2019" (first cap quarter) inside Oct 2018 - Mar 2019;
    # Ofgem's published table uses the half-year, so we do too.
    last = max(a9)
    return tuple(_build(k) for k in sorted(keys) if k <= last)


def receipt(period) -> Receipt:
    k = period if isinstance(period, tuple) else period_key(period)
    for r in receipts():
        if r.period == k:
            return r
    raise KeyError(f"no cap period {period!r}")


def _restated(basis: str, f: float) -> str:
    """An allowance restated at another consumption by Ofgem's own Annex 9 tab 1c method is
    still an allowance; anything else restated becomes `modelled`."""
    return basis if f == 1 or basis == BASIS_ALLOWANCE else "modelled"


def at_consumption(period, kwh: float) -> Receipt:
    """Restate a period's bill at another annual consumption (e.g. 2,500 kWh).

    Each line keeps its nil (standing-charge) part and scales its unit-rate part
    (typical − nil) by kwh / (Ofgem's Q). VAT is recomputed at the period's rate. This is
    how Ofgem's own Annex 9 tab 1c restates the cap at 2,500 kWh.
    """
    r = receipt(period)
    f = (kwh / 1000.0) / r.consumption_mwh
    new = []
    for ln in r.lines:
        if ln.key == "vat":
            continue
        new.append(replace(ln, typical=ln.nil + (ln.typical - ln.nil) * f,
                           basis=_restated(ln.basis, f)))
    ex = sum(ln.typical for ln in new)
    vat = r.line("vat")
    new.append(replace(vat, typical=ex * r.vat_rate, basis=_restated(vat.basis, f)))
    return replace(r, lines=tuple(new), consumption_mwh=kwh / 1000.0,
                   published_ex_vat=float("nan"), published_inc_vat=float("nan"),
                   published_note=f"restated by us at {kwh:,.0f} kWh; no published level")
