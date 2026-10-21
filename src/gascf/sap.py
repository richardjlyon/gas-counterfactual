"""Gas System Average Price (SAP), daily, p/kWh gross CV (spec §4, §8).

Two official copies of one National Gas series (PUBOB603, "SAP, Actual Day"):
- National Gas operational-data API: from 1 Dec 2021 only.
- ONS real-time indicators, 2021 edition (published 23 Dec 2021): 1 Jan 2018 to 19 Dec 2021,
  "Source: National Grid".
The overlap (1-19 Dec 2021) must agree to the published 4 decimal places (test).
No free official daily SAP has been found before 2018; the daily validation window
therefore starts in 2018.

Usage (refresh the API copy): uv run python -m gascf.sap
"""
from __future__ import annotations

import csv
import json
import urllib.request
from datetime import date, timedelta
from pathlib import Path

import openpyxl

RAW = Path(__file__).parents[2] / "data/raw"
API = "https://api.nationalgas.com/operationaldata/v1/publications/gasday"
API_FROM = date(2021, 12, 1)
API_CSV = RAW / "nationalgas" / "SAP_PUBOB603.csv"
ONS_XLSX = RAW / "ons" / "sap_2021.xlsx"


def fetch_api(start: date = API_FROM, end: date | None = None) -> Path:
    end = end or date.today()
    out: dict[str, str] = {}
    a = start
    while a <= end:
        b = min(a + timedelta(days=364), end)
        body = json.dumps({"fromDate": a.isoformat(), "toDate": b.isoformat(),
                           "publicationIds": ["PUBOB603"], "latestValue": "Y"}).encode()
        req = urllib.request.Request(API, data=body, headers={"Content-Type": "application/json", "User-Agent": "gascf/0.1 (richardlyon.net)"})
        with urllib.request.urlopen(req, timeout=120) as r:
            for p in json.load(r)[0]["publications"]:
                out[p["applicableFor"]] = p["value"]
        a = b + timedelta(days=1)
    API_CSV.parent.mkdir(parents=True, exist_ok=True)
    with API_CSV.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["gasDay", "sap_p_per_kwh"])
        for k in sorted(out):
            w.writerow([k, out[k]])
    return API_CSV


def read_api() -> dict[date, float]:
    with API_CSV.open(newline="") as f:
        return {date.fromisoformat(r["gasDay"]): float(r["sap_p_per_kwh"]) for r in csv.DictReader(f)}


def read_ons() -> dict[date, float]:
    wb = openpyxl.load_workbook(ONS_XLSX, read_only=True, data_only=True)
    ws = wb["System Average Price"]
    hdr = [r for r in ws.iter_rows(min_row=1, max_row=3, values_only=True)]
    if hdr[0][1] != "SAP, Actual Day" or hdr[1][1] != "Value (p/kWh)":
        raise ValueError("ONS SAP layout changed")
    return {r[0].date(): float(r[1]) for r in ws.iter_rows(values_only=True)
            if hasattr(r[0], "year") and isinstance(r[1], (int, float))}


def sap_gbp_per_mwh() -> dict[date, float]:
    """Daily SAP in £/MWh (gross CV): ONS before 1 Dec 2021, National Gas API from then."""
    ons, api = read_ons(), read_api()
    out = {d: v * 10 for d, v in ons.items() if d < API_FROM}
    out.update({d: v * 10 for d, v in api.items()})
    return dict(sorted(out.items()))


if __name__ == "__main__":
    p = fetch_api()
    print(p, sum(1 for _ in p.open()) - 1, "days")
