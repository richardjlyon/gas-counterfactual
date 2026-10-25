"""Readers for DUKES Table 5.10 and DESNZ Electricity Generation Costs 2025, Annex A.

DUKES 5.10.B: plant load factor (%), by generator type and year.
DUKES 5.10.C: thermal efficiency, gross CV basis (%), by generator type and year.
DESNZ Annex A, 'Technical and Cost Assumptions': per-technology inputs, by commissioning year.
DESNZ Annex A, 'Additional Estimates YYYY': DESNZ's own levelised cost rows (£/MWh, 2024 prices).
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl

from . import workbook_cache

DUKES_SHEET = "5.10.B and 5.10.C"
CCGT = "Combined cycle gas turbine stations"


def _years_row(rows, title: str) -> tuple[int, dict[int, int]]:
    """Index of the 'Generator Type' header row after `title`, and {year: column}."""
    for i, r in enumerate(rows):
        if r[0] and str(r[0]).startswith(title):
            h = rows[i + 1]
            if str(h[0]).strip() != "Generator Type":
                raise ValueError(f"DUKES layout changed under {title!r}")
            return i + 1, {int(v): j for j, v in enumerate(h) if v and str(v).isdigit()}
    raise ValueError(f"DUKES layout changed: {title!r} not found")


def _series(rows, hdr: int, years: dict[int, int], label: str) -> dict[int, float]:
    for r in rows[hdr + 1: hdr + 12]:
        if r[0] and str(r[0]).strip() == label:
            return {y: float(r[j]) for y, j in years.items() if isinstance(r[j], (int, float))}
    raise ValueError(f"DUKES: row {label!r} not found")


@dataclass(frozen=True)
class Dukes510:
    ccgt_load_factor: dict[int, float]   # per cent, Table 5.10.B
    ccgt_efficiency: dict[int, float]    # per cent, gross CV, Table 5.10.C
    published: str                        # from the cover sheet


def read_dukes_510(path: Path) -> Dukes510:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows = list(wb[DUKES_SHEET].iter_rows(values_only=True))
    hb, yb = _years_row(rows, "Table 5.10.B Plant load factor")
    hc, yc = _years_row(rows, "Table 5.10.C Thermal efficiency, gross calorific value basis")
    cover = [r[0] for r in wb["Cover Sheet"].iter_rows(values_only=True) if r[0]]
    pub = next(str(c) for c in cover if str(c).startswith("These data were published"))
    return Dukes510(
        ccgt_load_factor=_series(rows, hb, yb, CCGT),
        ccgt_efficiency=_series(rows, hc, yc, CCGT),
        published=pub.split("\n")[0],
    )


# ---------------------------------------------------------------- DESNZ Annex A

TECH_SHEET = "Technical and Cost Assumptions"


@dataclass(frozen=True)
class Plant:
    """One DESNZ Annex A technology column, medium case unless stated."""
    name: str
    commissioning: int
    size_mw: float
    net_load_factor: float          # Annex A's own; the '30%'/'93%' labels are not the LF used
    efficiency_gcv: float
    life_years: int
    predev_gbp_kw: dict[str, float]          # low / medium / high
    predev_phasing: tuple[float, ...]
    construction_gbp_kw: dict[str, float]
    construction_phasing: tuple[float, ...]
    infrastructure_gbp_k: dict[str, float]   # £'000 per plant
    fixed_om_gbp_mw_yr: float
    variable_om_gbp_mwh: float
    insurance_gbp_mw_yr: float
    connection_uos_gbp_mw_yr: float
    hurdle: dict[str, float]


def read_desnz_plants(path: Path, prefix: str = "Gas ") -> dict[tuple[str, int], Plant]:
    """Technology columns whose name starts with `prefix` (default: unabated and CCUS gas).

    Only gas columns are read: some renewable columns carry phasing that does not sum
    to 1, which this engine has no use for and does not try to interpret."""
    t = workbook_cache.sheet_rows(path, TECH_SHEET)
    labels = [(str(r[0]).strip() if r[0] else "", str(r[1]).strip() if r[1] else "") for r in t]

    def row(name: str, sub: str = "") -> int:
        cur = ""
        for i, (a, b) in enumerate(labels):
            cur = a or cur
            if cur.startswith(name) and (not sub or b == sub):
                return i
        raise ValueError(f"DESNZ Annex A layout changed: {name!r}/{sub!r}")

    def phasing(name: str, j: int) -> tuple[float, ...]:
        i = row(name)
        vals = [t[k][j] for k in range(i, i + 6)]
        while vals and not vals[-1]:
            vals.pop()
        if abs(sum(vals) - 1) > 1e-6:
            raise ValueError(f"DESNZ {name} phasing does not sum to 1 at column {j}")
        return tuple(float(v) for v in vals)

    def lmh(name: str, j: int) -> dict[str, float]:
        return {k: float(t[row(name, k.title())][j]) for k in ("low", "medium", "high")}

    plants = {}
    tech = None
    for j in range(3, len(t[0])):
        tech = str(t[0][j]).strip() if t[0][j] else tech
        year = t[1][j]
        if not tech or not isinstance(year, int) or not tech.startswith(prefix):
            continue
        plants[(" ".join(tech.split()), year)] = Plant(
            name=" ".join(tech.split()), commissioning=year,
            size_mw=float(t[row("Reference plant size")][j]),
            net_load_factor=float(t[row("Average load factor")][j]),
            efficiency_gcv=float(t[row("Average fuel efficiency")][j]),
            life_years=int(t[row("Operating lifetime")][j]),
            predev_gbp_kw=lmh("Pre-development Cost", j),
            predev_phasing=phasing("Pre-development phasing", j),
            construction_gbp_kw=lmh("Construction Cost", j),
            construction_phasing=phasing("Construction phasing", j),
            infrastructure_gbp_k=lmh("Infrastructure", j),
            fixed_om_gbp_mw_yr=float(t[row("Fixed O&M")][j]),
            variable_om_gbp_mwh=float(t[row("Variable O&M")][j]),
            insurance_gbp_mw_yr=float(t[row("Insurance")][j]),
            connection_uos_gbp_mw_yr=float(t[row("Connection and use of system charges")][j]),
            hurdle=lmh("Hurdle Rate", j),
        )
    return plants


def read_desnz_levelised(path: Path, commissioning: int = 2030) -> dict[str, dict[str, float | None]]:
    """DESNZ's printed levelised cost rows, £/MWh. '<1' is returned as None."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    e = list(wb[f"Additional Estimates {commissioning}"].iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(e) if r[0] == "Levelised cost breakdown")
    out = {}
    for j, tech in enumerate(e[hdr]):
        if j < 2 or not tech:
            continue
        out[" ".join(str(tech).split())] = {
            str(e[i][0]).rstrip("0123456789").strip(): (float(e[i][j]) if isinstance(e[i][j], (int, float)) else None)
            for i in range(hdr + 1, len(e)) if e[i][0]
        }
    return out
