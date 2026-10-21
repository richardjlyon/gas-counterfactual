"""Wholesale electricity and gas prices, by month, 2002–2025 (spec §3.1 NZ1, §5.1, §15.1 decision 3).

Terms: *day-ahead baseload* = the price of a flat block of power for every hour of the next day;
*NBP* = National Balancing Point, the GB wholesale gas hub; *SRMC* = short-run marginal cost,
the fuel plus carbon cost of one more MWh from a plant; *GCV* = gross calorific value.

Three electricity series, joined month by month:

1. **Measured, Jun 2010 →**: Ofgem Wholesale Market Indicators (WMI), "Electricity prices:
   day-ahead baseload contracts – monthly average (Great Britain)", `ofgem/s18a-WMI_elec_…js`.
2. **Ofgem route (a), Jan 2003 – May 2010**: Ofgem, *Wholesale Energy Markets in 2015*
   (Sep 2015), Figure 8 "Monthly average wholesale power prices (baseload Day-ahead),
   2003-2015", source ICIS. Ofgem prints no table, but the PDF chart is vector: the line is
   one path of 149 points whose coordinates are read from the stored PDF and converted with
   the chart's own gridlines. On the 60 months where it overlaps the WMI series (Jun 2010 –
   May 2015) it agrees to within £0.06/MWh, so the read is exact to plotting precision.
   Figure 5 of the same report gives NBP day-ahead gas from Jan 2001 the same way.
3. **Modelled, route (c)** (decision 3): fuel + carbon cost of the marginal plant, coal or gas,
   plus a market margin calibrated on Jun 2010 – Dec 2017; used for Jan–Dec 2002 (before the
   Ofgem chart starts) and exported for every month as the check against routes 1 and 2.

Which series is used before Jun 2010 is a switch (`route`): "ofgem_chart" (central, Ofgem's
own published price, with 2002 modelled), "model" (route c throughout), or "lower" (the lower
of the two each month: Richard's conservative rule).
"""
from __future__ import annotations

import json
import hashlib

import statistics
from datetime import date
from functools import lru_cache
from pathlib import Path
from gascf.pathutil import rel_to_root

import openpyxl
import pymupdf

from gascf import carbon

RAW = Path(__file__).parents[2] / "data/raw"
WMI_ELEC = RAW / "ofgem/s18a-WMI_elec_dayahead_monthly_n9WSDfgxc.js"
WMI_GAS = RAW / "ofgem/s18a-WMI_gas_dayahead_monthly_YFsS4g98P.js"
WEM2015 = RAW / "ofgem/s20c-Wholesale_Energy_Markets_in_2015.pdf"
QEP = RAW / "desnz_qep/s18a-QEP_table_321.xlsx"
DUKES_510 = RAW / "DUKES_5.10.xlsx"

SRC_WMI_ELEC = "ofgem-wmi-elec-dayahead-baseload-monthly"
SRC_WMI_GAS = "ofgem-wmi-gas-nbp-dayahead-monthly"
SRC_FIG8 = "ofgem-wholesale-energy-markets-2015-fig8-vector"
SRC_FIG5 = "ofgem-wholesale-energy-markets-2015-fig5-vector"
SRC_MODEL = "s20c-srmc-model-calibrated-2010-2017"
SRC_QEP = "desnz-qep-3.2.1"

KWH_PER_THERM = 29.3071                     # 1 therm = 100,000 Btu = 29.3071 kWh
PTHERM_PER_GBP_MWH = KWH_PER_THERM / 10     # p/therm ÷ 2.93071 = £/MWh

CALIBRATION = ((2010, 6), (2017, 12))      # WMI start to the last year before the 2018 MID cross-check
WMI_FROM = (2010, 6)

Month = tuple[int, int]


def _months(a: Month, b: Month) -> list[Month]:
    out, (y, m) = [], a
    while (y, m) <= b:
        out.append((y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


# ------------------------------------------------------------------ Ofgem WMI (everviz JS)

def read_wmi(path: Path) -> dict[Month, float]:
    """Monthly values from an Ofgem WMI everviz chart file (CSV string inside the JS)."""
    t = path.read_text()
    i = t.index('"csv":"') + 6
    csv_text, _ = json.JSONDecoder().raw_decode(t[i:])
    lines = csv_text.splitlines()
    if lines[0].replace('"', "").split(",")[0] != "Category":
        raise ValueError(f"{path.name}: header {lines[0]!r}")
    out = {}
    for ln in lines[1:]:
        d, v = ln.split(",")
        mm, _, yy = d.strip('"').split("/")
        out[(int(yy), int(mm))] = float(v)
    return out


@lru_cache(maxsize=None)
def wmi_elec() -> dict[Month, float]:
    """Day-ahead baseload, £/MWh."""
    return read_wmi(WMI_ELEC)


@lru_cache(maxsize=None)
def wmi_gas_ptherm() -> dict[Month, float]:
    """NBP day-ahead, p/therm."""
    return read_wmi(WMI_GAS)


def ptherm_to_gbp_mwh(p: float) -> float:
    return p / PTHERM_PER_GBP_MWH


# ------------------------------------------------------------------ Ofgem WEM 2015 chart vectors

def read_chart_line(pdf: Path, title: str, first_month: Month) -> dict[Month, float]:
    """Read a single-line monthly chart from the vector paths of the page carrying `title`.

    The data line is the thick black path (width > 1.5) with the most segments; the y scale
    is fitted from the numeric axis labels left of the plot and their gridlines. Points are
    one per month from `first_month`. Points left of the plot area (hidden by the chart's clip,
    so not shown by Ofgem) are dropped.
    """
    doc = pymupdf.open(pdf)
    page = next((p for p in doc if title in p.get_text()), None)
    if page is None:
        raise ValueError(f"{pdf.name}: chart {title!r} not found")
    lines = [d for d in page.get_drawings()
             if d.get("color") == (0.0, 0.0, 0.0) and (d.get("width") or 0) > 1.5]
    path = max(lines, key=lambda d: len(d["items"]))
    if any(it[0] != "l" for it in path["items"]):
        raise ValueError("chart line is not a polyline")
    pts = [path["items"][0][1]] + [it[2] for it in path["items"]]
    hlines = [it for d in page.get_drawings() for it in d["items"]
              if it[0] == "l" and abs(it[1].y - it[2].y) < 0.01 and abs(it[2].x - it[1].x) > 100
              and it[1].y < 400]
    grid = sorted({round(it[1].y, 2) for it in hlines})
    plot_left = min(min(it[1].x, it[2].x) for it in hlines)   # points left of this are clipped
    labels = [(float(w[4]), (w[1] + w[3]) / 2) for w in page.get_text("words")
              if w[4].isdigit() and w[0] < 150 and 100 < w[1] < 345]
    snapped = [(v, min(grid, key=lambda g: abs(g - yc))) for v, yc in labels]
    (v0, y0), (v1, y1) = min(snapped), max(snapped)
    scale = (v1 - v0) / (y0 - y1)
    for v, y in snapped:                                  # every label sits on its gridline
        if abs(v0 + (y0 - y) * scale - v) > 0.01 * (v1 - v0):
            raise ValueError(f"chart axis inconsistent at label {v}")
    dx = [b.x - a.x for a, b in zip(pts, pts[1:])]
    if max(dx) - min(dx) > 0.2:
        raise ValueError("chart points not evenly spaced by month")
    out = {}
    for k, p in zip(_months(first_month, (2100, 1)), pts):
        if p.x >= plot_left - 1.0:                         # drop points hidden by the plot clip
            out[k] = v0 + (y0 - p.y) * scale
    return out


@lru_cache(maxsize=None)
def fig8_elec() -> dict[Month, float]:
    """Ofgem WEM 2015 Fig 8, day-ahead baseload £/MWh, Jan 2003 – May 2015."""
    return read_chart_line(WEM2015, "Figure 8: Monthly average wholesale power prices", (2003, 1))


@lru_cache(maxsize=None)
def fig5_gas_ptherm() -> dict[Month, float]:
    """Ofgem WEM 2015 Fig 5, NBP day-ahead p/therm, Dec 2000 – May 2015. The title says
    2003-2015 but the plotted axis runs from 2001; the path in the file starts at Jan 2000
    (padding and points hidden by the plot clip, dropped). Month alignment is proved by the
    overlap test against WMI (68 months within 0.04 p/therm)."""
    return read_chart_line(WEM2015, "Figure 5: Monthly average wholesale gas prices", (2000, 1))


def overlap_error(chart: dict[Month, float], wmi: dict[Month, float]) -> dict[str, float]:
    d = [chart[k] - wmi[k] for k in chart if k in wmi]
    return {"n": len(d), "mean": statistics.mean(d), "max_abs": max(abs(x) for x in d)}


# ------------------------------------------------------------------ fuel inputs for the model

_Q = {"Jan to Mar": (1, 2, 3), "Apr to Jun": (4, 5, 6), "Jul to Sep": (7, 8, 9), "Oct to Dec": (10, 11, 12)}


@lru_cache(maxsize=None)
def qep_quarterly() -> dict[Month, dict[str, float]]:
    """DESNZ QEP 3.2.1 quarterly price paid by major power producers, £/MWh of fuel (GCV),
    spread to each month of the quarter: {'coal': …, 'gas': …} (None where '..')."""
    ws = openpyxl.load_workbook(QEP, read_only=True, data_only=True)["3.2.1"]
    rows = list(ws.iter_rows(values_only=True))
    h = next(r for r in rows if r[0] == "Year" and r[1] == "Quarter")
    jc = h.index("Major power producers: Coal (pence per kWh) [Note 1]")
    jg = h.index("Major power producers: Natural gas (pence per kWh) [Note 4]")
    out = {}
    for r in rows:
        if isinstance(r[0], int) and r[1] in _Q:
            c = r[jc] * 10 if isinstance(r[jc], (int, float)) else None
            g = r[jg] * 10 if isinstance(r[jg], (int, float)) else None
            for m in _Q[r[1]]:
                out[(r[0], m)] = {"coal": c, "gas": g}
    return out


@lru_cache(maxsize=None)
def dukes_efficiency() -> dict[str, dict[int, float]]:
    """DUKES 5.10.C thermal efficiency (GCV, per cent): {'ccgt': {...}, 'coal': {...}}."""
    ws = openpyxl.load_workbook(DUKES_510, read_only=True, data_only=True)["5.10.B and 5.10.C"]
    rows = list(ws.iter_rows(values_only=True))
    i = next(k for k, r in enumerate(rows) if r[0] and str(r[0]).startswith("Table 5.10.C"))
    h = rows[i + 1]
    # headers are '1996' … but '2004 [note 11]' carries a note: take the leading year
    yrs = {j: int(str(v)[:4]) for j, v in enumerate(h) if v and str(v)[:4].isdigit()}
    if 2004 not in yrs.values():
        raise ValueError("DUKES 5.10.C: 2004 column not found")
    out = {}
    for r in rows[i + 2:i + 8]:
        lab = str(r[0] or "").strip()
        key = {"Combined cycle gas turbine stations": "ccgt", "Coal fired stations": "coal"}.get(lab)
        if key:
            out[key] = {y: float(r[j]) for j, y in yrs.items() if isinstance(r[j], (int, float)) and r[j] > 0}
    if set(out) != {"ccgt", "coal"}:
        raise ValueError("DUKES 5.10.C rows moved")
    return out


def nbp_gbp_mwh() -> dict[Month, tuple[float, str]]:
    """NBP day-ahead £/MWh by month: WMI from Oct 2009, Ofgem Fig 5 before (Dec 2000 →)."""
    out = {k: (ptherm_to_gbp_mwh(v), SRC_FIG5) for k, v in fig5_gas_ptherm().items()}
    out.update({k: (ptherm_to_gbp_mwh(v), SRC_WMI_GAS) for k, v in wmi_gas_ptherm().items()})
    return out


# ------------------------------------------------------------------ route (c): the model

@lru_cache(maxsize=None)
def srmc_inputs(years: tuple[int, int] = (2002, 2025)) -> dict[Month, dict[str, float]]:
    """SRMC of gas (NBP) and coal (QEP) plant at the margin, £/MWh elec, by month, carbon included."""
    eff = dukes_efficiency()
    gas = nbp_gbp_mwh()
    qep = qep_quarterly()
    allow = carbon.allowance_monthly_v4(range(years[0], years[1] + 1))
    out = {}
    for k in _months((years[0], 1), (years[1], 12)):
        if k not in gas or qep.get(k, {}).get("coal") is None:
            continue
        d = date(k[0], k[1], 1)
        a = allow[k][0]
        eg, ec = eff["ccgt"][k[0]] / 100, eff["coal"][k[0]] / 100
        out[k] = {"gas": gas[k][0] / eg + carbon.c_marg_gbp_per_mwh_elec(a, d, eg),
                  "coal": qep[k]["coal"] / ec + carbon.c_marg_coal_gbp_per_mwh_elec(a, d, ec)}
    return out


def _ols(xs: list[list[float]], y: list[float]) -> list[float]:
    """Least squares by the normal equations (no numpy): returns coefficients."""
    n = len(xs[0])
    a = [[sum(r[i] * r[j] for r in xs) for j in range(n)] for i in range(n)]
    b = [sum(r[i] * t for r, t in zip(xs, y)) for i in range(n)]
    for i in range(n):                                    # Gauss–Jordan
        p = a[i][i]
        a[i] = [v / p for v in a[i]]; b[i] /= p
        for k in range(n):
            if k != i:
                f = a[k][i]
                a[k] = [vk - f * vi for vk, vi in zip(a[k], a[i])]; b[k] -= f * b[i]
    return b


MODEL_FORMS = ("gas_only", "dearer_marginal", "cheaper_marginal", "ols_gas_coal")


@lru_cache(maxsize=None)
def calibrate(form: str) -> dict:
    """Fit a model form on Jun 2010 – Dec 2017 against WMI; return params and errors.

    - gas_only:         P = SRMC_gas + μ           (§15 1(c) as first written)
    - dearer_marginal:  P = max(SRMC_gas, SRMC_coal) + μ
    - cheaper_marginal: P = min(SRMC_gas, SRMC_coal) + μ   (coal sets the price when it is cheaper)
    - ols_gas_coal:     P = α + βg·SRMC_gas + βc·SRMC_coal   (share of each fuel at the margin)
    """
    s = srmc_inputs()
    w = wmi_elec()
    ks = [k for k in _months(*CALIBRATION) if k in s and k in w]
    y = [w[k] for k in ks]
    if form == "ols_gas_coal":
        coef = _ols([[1.0, s[k]["gas"], s[k]["coal"]] for k in ks], y)
        params = {"alpha": coef[0], "beta_gas": coef[1], "beta_coal": coef[2]}
    else:
        base = [_base(form, s[k]) for k in ks]
        params = {"mu": statistics.mean(t - b for t, b in zip(y, base))}
    pred = [model_price(form, s[k], params) for k in ks]
    err = [p - t for p, t in zip(pred, y)]
    by_year = {}
    for k, e in zip(ks, err):
        by_year.setdefault(k[0], []).append(e)
    return {"form": form, "params": params, "n": len(ks),
            "mean_error": statistics.mean(err), "rmse": statistics.mean(e * e for e in err) ** 0.5,
            "mae": statistics.mean(abs(e) for e in err),
            "annual_mean_error": {yy: statistics.mean(v) for yy, v in by_year.items()}}


def _base(form: str, s: dict[str, float]) -> float:
    if form == "gas_only":
        return s["gas"]
    if form == "dearer_marginal":
        return max(s["gas"], s["coal"])
    if form == "cheaper_marginal":
        return min(s["gas"], s["coal"])
    raise ValueError(form)


def model_price(form: str, s: dict[str, float], params: dict) -> float:
    if form == "ols_gas_coal":
        return params["alpha"] + params["beta_gas"] * s["gas"] + params["beta_coal"] * s["coal"]
    return _base(form, s) + params["mu"]


def backcast_error(form: str) -> dict:
    """Out-of-sample error of a calibrated form against Ofgem Fig 8, Jan 2003 – May 2010."""
    s, f8 = srmc_inputs(), fig8_elec()
    p = calibrate(form)["params"]
    ks = [k for k in _months((2003, 1), (2010, 5)) if k in s and k in f8]
    err = [model_price(form, s[k], p) - f8[k] for k in ks]
    by_year = {}
    for k, e in zip(ks, err):
        by_year.setdefault(k[0], []).append(e)
    return {"n": len(ks), "mean_error": statistics.mean(err),
            "rmse": statistics.mean(e * e for e in err) ** 0.5,
            "annual_mean_error": {yy: statistics.mean(v) for yy, v in by_year.items()}}


COAL_FORMS = ("dearer_marginal", "cheaper_marginal", "ols_gas_coal")   # forms that let coal set the price
CALIBRATION_RMSE_GATE = 5.0      # £/MWh, the same width as the validate.py ±£5 spread gate
MODEL_ONLY = ((2002, 1), (2002, 12))   # months with neither WMI nor the Ofgem chart


@lru_cache(maxsize=None)
def choose_model_form() -> str:
    """Decision 3 as a rule: among forms that allow coal to set the price, keep those whose
    2010–2017 calibration RMSE is within £5/MWh, then take the one giving the **lower** mean
    price over the model-only months (2002): Richard's conservative rule. On the stored data
    this is 'ols_gas_coal' (calibration RMSE £3.29; 2002 mean £12.0/MWh)."""
    ok = [f for f in COAL_FORMS if calibrate(f)["rmse"] <= CALIBRATION_RMSE_GATE]
    if not ok:
        raise ValueError("no coal-aware model form passes the calibration gate")
    ks = _months(*MODEL_ONLY)
    return min(ok, key=lambda f: statistics.mean(modelled_elec(f)[k] for k in ks))


def model_only_range() -> dict[str, float]:
    """Mean 2002 price under every calibrated form (the decision-1 range for the page)."""
    ks = _months(*MODEL_ONLY)
    return {f: statistics.mean(modelled_elec(f)[k] for k in ks) for f in MODEL_FORMS}


def modelled_elec(form: str | None = None) -> dict[Month, float]:
    """Route (c) price by month, £/MWh, for every month with fuel inputs (2002 →)."""
    form = form or choose_model_form()
    p = calibrate(form)["params"]
    return {k: model_price(form, v, p) for k, v in srmc_inputs().items()}


# ------------------------------------------------------------------ the joined series

ROUTES = ("ofgem_chart", "model", "lower")


def elec_monthly(route: str = "ofgem_chart", years: range = range(2002, 2026)) -> dict[Month, dict]:
    """P_elec(m), £/MWh, with basis and source_id, Jan 2002 – Dec 2025."""
    if route not in ROUTES:
        raise ValueError(route)
    w, f8, mod = wmi_elec(), fig8_elec(), modelled_elec()
    out = {}
    for k in _months((years[0], 1), (years[-1], 12)):
        if k >= WMI_FROM:
            out[k] = {"gbp_per_mwh": w[k], "basis": "measured", "source_id": SRC_WMI_ELEC}
        elif k in f8 and route == "ofgem_chart":
            out[k] = {"gbp_per_mwh": f8[k], "basis": "measured", "source_id": SRC_FIG8}
        elif k in f8 and route == "lower" and f8[k] <= mod[k]:
            out[k] = {"gbp_per_mwh": f8[k], "basis": "measured", "source_id": SRC_FIG8}
        else:
            out[k] = {"gbp_per_mwh": mod[k], "basis": "modelled", "source_id": SRC_MODEL}
    return out


def elec_annual_mean(route: str = "ofgem_chart") -> dict[int, dict]:
    """Simple mean of the monthly prices by calendar year (for display; NZ1 weights by V)."""
    m = elec_monthly(route)
    out = {}
    for (y, _), v in m.items():
        out.setdefault(y, []).append(v)
    return {y: {"gbp_per_mwh": statistics.mean(x["gbp_per_mwh"] for x in v),
                "basis": "measured" if all(x["basis"] == "measured" for x in v) else "modelled",
                "source_id": "+".join(sorted({x["source_id"] for x in v}))} for y, v in out.items()}


def nz1(volume_twh: dict[Month, float], route: str = "ofgem_chart") -> dict[int, dict]:
    """NZ1, £m by calendar year: Σ months P_elec(m) × V(m), V in TWh from the residual strand."""
    p = elec_monthly(route, range(min(volume_twh)[0], max(volume_twh)[0] + 1))
    out: dict[int, dict] = {}
    for k, v in volume_twh.items():
        r = out.setdefault(k[0], {"gbp_m": 0.0, "basis": "measured", "source_id": set()})
        r["gbp_m"] += p[k]["gbp_per_mwh"] * v
        r["source_id"].add(p[k]["source_id"])
        if p[k]["basis"] != "measured":
            r["basis"] = "modelled"
    for r in out.values():
        r["source_id"] = "+".join(sorted(r["source_id"]))
    return out


# ------------------------------------------------------------------ QEP vs day-ahead ratio (ticket 47, EIC ruling 17 Oct 2026 pass-4 point 3)

QEP_RATIO_YEARS = (2010, 2025)


@lru_cache(maxsize=None)
def qep_annual_gas_p_per_kwh(path: Path = QEP) -> dict[int, float]:
    """DESNZ QEP 3.2.1 (Annual), major power producers' natural gas price, p/kWh of fuel (GCV), by calendar year.
    Read by the column's header text from the Annual sheet (DESNZ's own annual figure, not the month-averaged
    quarterly sheet); unrounded. `path` lets a test point at a temporary copy; the engine always uses QEP."""
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["3.2.1 (Annual)"]
    rows = list(ws.iter_rows(values_only=True))
    h = next(r for r in rows if r and r[0] == "Year")
    jg = h.index("Major power producers: Natural gas (pence per kWh) [Note 4]")
    return {r[0]: float(r[jg]) for r in rows if isinstance(r[0], int) and isinstance(r[jg], (int, float))}


@lru_cache(maxsize=None)
def _wmi_gas_ptherm_from(path: Path = WMI_GAS) -> dict:
    """WMI NBP day-ahead gas, p/therm, monthly, from `path` (defaults to the registered WMI_GAS)."""
    return read_wmi(path)


def _sha256(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def qep_dayahead_ratio_2010_25(qep_path: Path = QEP, wmi_path: Path = WMI_GAS) -> dict:
    """Ratio of the QEP major power producers' gas price to the Ofgem WMI NBP day-ahead price, by calendar year
    2010-25, with the plain mean of the 16 yearly ratios and the lowest and highest (ticket 47). Both prices in
    £/MWh of fuel (GCV): QEP p/kWh x 10; WMI the year's 12-month mean p/therm / 2.93071 (ptherm_to_gbp_mwh).
    Unrounded throughout: the ratio uses the unrounded prices, not the derivation CSV's 2 dp columns. Requires all
    12 months of every year 2010-25 in each input; raises otherwise. Not the same as qep_vs_nbp_annual (spec
    §12.2), which averages the quarterly sheet against nbp_gbp_mwh; leave that and its test untouched."""
    qep = qep_annual_gas_p_per_kwh(qep_path)
    wmi = _wmi_gas_ptherm_from(wmi_path)
    a, b = QEP_RATIO_YEARS
    by_year = {}
    for y in range(a, b + 1):
        if y not in qep:
            raise ValueError(f"QEP annual gas price missing for {y}")
        ks = [(y, m) for m in range(1, 13)]
        if not all(k in wmi for k in ks):
            raise ValueError(f"WMI day-ahead missing a month in {y}")
        qep_gbp_mwh = qep[y] * 10
        wmi_mean_ptherm = statistics.mean(wmi[k] for k in ks)
        wmi_gbp_mwh = ptherm_to_gbp_mwh(wmi_mean_ptherm)
        by_year[y] = {"qep_p_per_kwh": qep[y], "qep_gbp_mwh": qep_gbp_mwh, "wmi_months": len(ks),
                      "wmi_mean_p_per_therm": wmi_mean_ptherm, "wmi_gbp_mwh": wmi_gbp_mwh,
                      "ratio": qep_gbp_mwh / wmi_gbp_mwh}
    ratios = {y: by_year[y]["ratio"] for y in by_year}
    ymin = min(ratios, key=ratios.get)
    ymax = max(ratios, key=ratios.get)
    return {"by_year": by_year, "mean": statistics.mean(ratios.values()),
            "min": {"year": ymin, "value": ratios[ymin]}, "max": {"year": ymax, "value": ratios[ymax]},
            "inputs": [{"path": rel_to_root(qep_path), "sha256": _sha256(qep_path)},
                       {"path": rel_to_root(wmi_path), "sha256": _sha256(wmi_path)}]}


# ------------------------------------------------------------------ gas price join (§12.2 item 12)

def qep_vs_nbp_annual() -> dict[int, dict[str, float]]:
    """QEP gas-for-power price against NBP day-ahead, £/MWh gas, calendar-year means."""
    q, n = qep_quarterly(), nbp_gbp_mwh()
    out = {}
    for y in range(2001, 2026):
        ks = [(y, m) for m in range(1, 13)]
        if all(k in n for k in ks) and all(q.get(k, {}).get("gas") for k in ks):
            qq = statistics.mean(q[k]["gas"] for k in ks)
            nn = statistics.mean(n[k][0] for k in ks)
            out[y] = {"qep": qq, "nbp": nn, "ratio": qq / nn}
    return out
