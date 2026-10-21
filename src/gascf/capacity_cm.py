"""Backup capacity: what Capacity Market (CM) cost an all-gas Britain would face against
today's (report only; spec §3.3 is NOT changed by this module).

The CM pays power stations, storage, interconnectors and demand response to be available at
winter peak. Each auction buys a "target capacity" set by NESO's Electricity Capacity Report
(ECR): roughly, peak demand plus a margin, less the de-rated capacity expected to be there
anyway *outside* the CM. Subsidised renewables are outside it by rule: a unit may take part
only if it is "not in receipt of a benefit from an excluded low carbon scheme (e.g. Contracts
for Difference)" (Ofgem, CM annual report 2025/26, para 1.8). "De-rated" means scaled down to
what the unit is expected to deliver at peak; for wind NESO uses an equivalent firm capacity
(EFC) of 0.182 for 2027/28 (ECR 2026 data workbook, Figure 24).

So subsidised wind *reduces* what the CM must buy. NESO says so in its own words for the 2026
recommendation: "a small increase in Contracts for Difference (CfD) capacity, i.e. more
CM-ineligible capacity leading to a reduction in the recommendation (−0.3GW)". In an all-gas
Britain that wind is absent and the CM would have to buy its de-rated contribution from gas.

What this module computes (all read from stored files):

* ``wind_outside_cm_derated_gw``: ECR Figure B4 Base Case CM-ineligible wind nameplate for the
  delivery year x the Figure 24 wind EFC (2027/28 column, the nearest given). Wind only: the
  ECR does not split solar out of its non-wind ineligible figure (Figure B5), so this is a
  floor on the de-rated capacity subsidised renewables supply.
* ``clearing_price``: the T-4 clearing price for the delivery year (Ofgem CM annual report
  2025/26, Table 1; DY 2026/27 £63.00/kW/yr, matching the DESNZ Auction Monitor report).
* ``extra_cost_gbp``: the extra CM an all-gas Britain would buy = GW x price. Per average
  household by domestic share of electricity / GB households (CM is charged on winter-peak
  demand, where the domestic share is higher, so this understates the household part).

Whether today's CM cost is *higher* because of renewables: no primary source attributes any
part of the clearing price to renewables. The volume effect NESO does publish runs the other
way (renewables shrink the target). Recommendation (in notes for the parent): keep §3.3.
Fleet A already builds and pays in full the spare plant an all-gas Britain needs, so holding
the CM equal does not hide a net-zero cost; if anything an all-gas CM would be larger, which
would raise step 1 and shrink the net-zero cost (the cautious direction already holds).
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

import openpyxl
import pymupdf

RAW = Path(__file__).parents[2] / "data/raw"
ECR_WB = RAW / "desnz_cm/NESO_ECR_2026_Data_Workbook.xlsx"
ECR_PDF = RAW / "desnz_cm/NESO_ECR_2026.pdf"
OFGEM_CM = RAW / "ofgem_cm/Ofgem_CM_annual_report_2025-26.pdf"
MONITOR_T4_2627 = RAW / "desnz_cm/auction_monitor_T4_DY2026-27.pdf"
SRC_ECR = "neso_ecr_2026_data_workbook"
SRC_OFGEM_CM = "ofgem_cm_annual_report_2025-26"
SRC_MONITOR = "desnz_auction_monitor_t4_dy2026-27"

QUOTE_ELIGIBILITY = ("not in receipt of a benefit from an excluded low carbon scheme (e.g. "
                     "Contracts for Difference)")
QUOTE_NESO_CFD = ("a small increase in Contracts for Difference (CfD) capacity, i.e. more "
                  "CM-ineligible capacity leading to a reduction in the recommendation (−0.3GW)")


@lru_cache(maxsize=4)
def _text(path: Path) -> str:
    return re.sub(r"\s+", " ", " ".join(p.get_text() for p in pymupdf.open(path)))


def _sheet(name: str) -> list[tuple]:
    wb = openpyxl.load_workbook(ECR_WB, read_only=True, data_only=True)
    if name not in wb.sheetnames:
        raise ValueError(f"ECR workbook layout changed: no sheet {name!r}")
    # the sheets have blank leading columns: keep only filled cells, in order
    return [tuple(c for c in r if c is not None) or (None,) for r in wb[name].iter_rows(values_only=True)]


def _dy(fy: int) -> str:
    return f"{fy}/{(fy + 1) % 100:02d}"


@lru_cache(maxsize=1)
def read_wind_outside_cm() -> dict[int, float]:
    """ECR Figure B4, 'Base Case 2026' row: CM-ineligible wind, GW nameplate, by delivery year."""
    rows = _sheet("Figure B4")
    if "Wind capacity not eligible for the CM" not in str(rows[0][0]):
        raise ValueError("ECR Figure B4 title changed")
    hdr = next(r for r in rows if r[0] == "Capacity (GW)")
    base = next((r for r in rows if r[0] == "Base Case 2026"), None)
    if base is None:
        raise ValueError("ECR Figure B4: 'Base Case 2026' row not found")
    return {int(str(h)[:4]): float(v) for h, v in zip(hdr[1:], base[1:])
            if h and re.fullmatch(r"\d{4}/\d{2}", str(h)) and isinstance(v, (int, float))}


@lru_cache(maxsize=1)
def read_wind_efc() -> dict[int, float]:
    """ECR Figure 24: wind fleet EFC by delivery year (2027/28 and 2030/31 are given)."""
    rows = _sheet("Figure 24")
    hdr = next(r for r in rows if r[0] == "Technology Class")
    wind = next((r for r in rows if r[0] == "Wind"), None)
    if wind is None:
        raise ValueError("ECR Figure 24: 'Wind' row not found")
    return {int(str(h)[:4]): float(v) for h, v in zip(hdr[1:], wind[1:])
            if h and isinstance(v, (int, float))}


@lru_cache(maxsize=1)
def read_t4_clearing_prices() -> dict[int, float]:
    """Ofgem CM annual report 2025/26, Table 1: T-4/T-3 clearing price by delivery year, £/kW/yr."""
    t = _text(OFGEM_CM)
    i = t.find("Table 1: Capacity Market T-4/T-3 Auctions")
    j = t.find("T-1 Auctions", i)
    if i < 0 or j < 0:
        raise ValueError("Ofgem CM report layout changed: Table 1 not found")
    return {int(y): float(p) for y, p in re.findall(r"(\d{4})/\d{2} .*?£(\d+\.\d\d)/kW/year", t[i:j])}


def read_monitor_t4_2627() -> tuple[float, float]:
    """DESNZ Auction Monitor, 2022 T-4 (DY 2026/27): (clearing £/kW/yr, awarded MW)."""
    t = _text(MONITOR_T4_2627)
    m = re.search(r"Clearing Price of the 2022 T-4 Auction was equal to £(\d+\.\d\d), with the "
                  r"aggregate Bidding Capacity of Capacity Agreements of ([\d,]+\.\d+)MW", t)
    if not m:
        raise ValueError("Auction Monitor report layout changed")
    return float(m.group(1)), float(m.group(2).replace(",", ""))


def check_quotes() -> dict[str, bool]:
    return {"eligibility": QUOTE_ELIGIBILITY in _text(OFGEM_CM),
            "neso_cfd": QUOTE_NESO_CFD in _text(ECR_PDF)}


@dataclass(frozen=True)
class CmComparison:
    delivery_year: str
    wind_outside_cm_gw: float
    wind_efc: float
    wind_efc_year: str
    wind_outside_cm_derated_gw: float
    clearing_price: float          # £/kW/yr, nominal, T-4 for this delivery year
    extra_cost_gbp: float          # nominal £/yr an all-gas CM would add
    extra_per_avg_household: float | None
    basis: str
    source_id: str
    note: str


def all_gas_cm_extra(fy: int) -> CmComparison:
    """Extra CM an all-gas Britain would buy in delivery year fy/fy+1 (October to September)."""
    from .system_costs import domestic_share_cy, gb_households_series, window_weight
    wind = read_wind_outside_cm()
    if fy not in wind:
        raise KeyError(f"UNSOURCED: ECR Figure B4 has no {_dy(fy)}")
    efc = read_wind_efc()
    efc_year = min(efc, key=lambda y: (abs(y - fy), y))
    prices = read_t4_clearing_prices()
    if fy not in prices:
        raise KeyError(f"UNSOURCED: no T-4 clearing price for DY {_dy(fy)}")
    gw = wind[fy] * efc[efc_year]
    cost = gw * 1e6 * prices[fy]             # GW -> kW
    w = {fy: 0.25, fy + 1: 0.75}        # CM delivery year 1 Oct fy to 30 Sep fy+1: 3/12 + 9/12
    try:
        dom, _ = window_weight(domestic_share_cy(), w)
        hh, _ = window_weight({y: v for y, (v, _) in gb_households_series().items()}, w)
        per_hh = cost * dom / hh
    except KeyError:
        per_hh = None
    return CmComparison(
        _dy(fy), wind[fy], efc[efc_year], _dy(efc_year), gw, prices[fy], cost, per_hh, "estimated",
        f"{SRC_ECR}; {SRC_OFGEM_CM}",
        f"ECR Fig B4 Base Case wind outside CM {wind[fy]:.1f} GW x Fig 24 wind EFC "
        f"{efc[efc_year]:.3f} ({_dy(efc_year)}) = {gw:.2f} GW de-rated, at the DY {_dy(fy)} T-4 "
        f"price £{prices[fy]:.2f}/kW/yr; wind only (solar not split out), per household by "
        "domestic share of electricity (understates: CM is charged at winter peak)")
