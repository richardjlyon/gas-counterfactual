"""UK emissions-trading auction revenue as traced, s36 (ticket 20; decisions 0056, 0067, 0068).

Phase II 2008-12: the 30 UK auctions (DMO results; DECC Trust Statements), the published £ as they
stand, by calendar year of the auction date. The derived columns (volume x price x rate) are not read."""
from __future__ import annotations

import csv
import hashlib
from datetime import datetime
from functools import lru_cache
from pathlib import Path

import openpyxl

RAW = Path(__file__).resolve().parents[2] / "data/raw"
PHASE2 = RAW / "s36/ets/uk-ets-phase2-auctions.csv"
PHASE2_SHA256 = "0ccf58033ab5c9e6a611ebf438f6cf3769dd063609e49a07221cf2faa1806cbf"   # data/raw/SHA256SUMS
SRC_PHASE2 = "uk-ets-phase2-auctions-dmo-decc"


def _sha(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _checked_rows(path: Path, sha256: str, n: int) -> list[dict]:
    if _sha(path) != sha256:
        raise ValueError(f"{Path(path).name}: sha256 differs from the registered file")
    rows = _rows(path)
    if len(rows) != n:
        raise ValueError(f"{Path(path).name}: {len(rows)} rows, expected {n} rows")
    return rows


def _rows(path: Path) -> list[dict]:
    with Path(path).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def phase2_revenue_cy(path: Path = PHASE2, sha256: str = PHASE2_SHA256) -> dict[int, float]:
    """{calendar year: UK Phase II auction revenue, cash £m}, 2008-12 (DECC published £'000).
    Refuses a file whose sha256 differs or that has other than 30 rows."""
    out: dict[int, float] = {}
    for r in _checked_rows(path, sha256, 30):
        y = datetime.strptime(r["date"], "%d %B %Y").year
        out[y] = out.get(y, 0.0) + float(r["revenue_gbp_thousand_published_decc_trust_statement"]) / 1e3
    return out


PHASE3 = RAW / "s36/ets/uk-ets-phase3-2012-auctions.csv"
PHASE3_SHA256 = "188c6f18d243c2654640dcbf46dd6f3b4eedd1bf89c3325ced235f1a165b31d3"   # data/raw/SHA256SUMS


def phase3_2012(path: Path = PHASE3, sha256: str = PHASE3_SHA256) -> dict:
    """0067: the two early Phase III auctions of Nov-Dec 2012 (DECC ARA 2012-13, HC 18, Trust Statement
    Note 2.2), cash £m, kept apart from the Phase II series and added only in tax collected."""
    rows = _checked_rows(path, sha256, 2)
    for r in rows:
        d = datetime.strptime(r["date"], "%d %b %Y")
        if not (d.year == 2012 and d.month in (11, 12)):
            raise ValueError(f"{Path(path).name}: {r['date']} is outside Nov-Dec 2012")
    return {"gbp_m": sum(float(r["revenue_gbp_thousand_published_decc_trust_statement"]) for r in rows) / 1e3,
            "rows": len(rows), "allowances": sum(int(r["allowances_auctioned"]) for r in rows),
            "phase": "III", "year": 2012, "source_id": rows[0]["source_id"]}


# ------------------------------------------------------------------ power share: the 0055/0068 stand-in
GHG_TABLES = RAW / "s36/ets/final-greenhouse-gas-emissions-tables-2024.xlsx"
EEA_REGISTRY = RAW / "s36/ets/ETS_Database_September_2026.xlsx"
SRC_GHG = "desnz-final-ghg-1990-2024-t1.2"
SRC_EEA = "eea-ets-union-registry-v82-sep2026"
# UK EU ETS verified emissions: GB stationary (20-99), GB aviation (10, from 2012) and Northern Ireland (XI 20-99)
EEA_UK_PARTS = (("GB", "20-99"), ("GB", "10"), ("XI", "20-99"))
STAND_IN_LOW_FACTOR = 0.957   # 0055 band low: the lowest calibration year (2023, UK ETS power / inventory less MSW)


@lru_cache(maxsize=None)
def _inventory_power_less_msw() -> dict[int, float]:
    """DESNZ GHG Table 1.2, MtCO2e: 'Electricity supply total' less 'Power stations - municipal solid waste'."""
    rows = list(openpyxl.load_workbook(GHG_TABLES, read_only=True, data_only=True)["1.2"].iter_rows(values_only=True))
    hdr = next(r for r in rows if r[0] == "TES sector")
    tot = next(r for r in rows if r[0] == "Electricity supply total")
    msw = next(r for r in rows if r[2] == "Power stations - municipal solid waste")
    return {int(h): tot[j] - msw[j] for j, h in enumerate(hdr) if h and str(h).isdigit()}


@lru_cache(maxsize=None)
def _eea_verified() -> dict[tuple[str, str, int], float]:
    """EEA Union Registry v82: {(country, activity, year): verified emissions, t}."""
    it = openpyxl.load_workbook(EEA_REGISTRY, read_only=True, data_only=True)["Sheet1"].iter_rows(values_only=True)
    h = next(it)
    ix = {k: h.index(k) for k in ("country_code", "main_activity_code", "citl_information", "year", "value")}
    out = {}
    for r in it:
        if (r[ix["citl_information"]] == "2. Verified emissions" and r[ix["country_code"]] in ("GB", "XI")
                and str(r[ix["year"]]).isdigit()):          # skips the trading-period total rows
            out[(r[ix["country_code"]], r[ix["main_activity_code"]], int(r[ix["year"]]))] = r[ix["value"]] or 0.0
    return out


def stand_in_share(y: int) -> dict:
    """0055/0068: power's share of UK EU ETS emissions = inventory power stations less waste (DESNZ 1.2)
    / UK verified EU ETS emissions (EEA), both for year y."""
    num = _inventory_power_less_msw()[y]
    ver = _eea_verified()
    den = sum(ver.get((c, a, y), 0.0) for c, a in EEA_UK_PARTS) / 1e6
    act20 = sum(ver.get((c, "20", y), 0.0) for c in ("GB", "XI")) / 1e6
    return {"value": num / den, "numerator_mtco2": num, "denominator_mtco2": den,
            "low": STAND_IN_LOW_FACTOR * num / den, "high": act20 / den, "rule": "0055 stand-in",
            "numerator_source_id": SRC_GHG, "denominator_source_id": SRC_EEA}


SHARE_2008 = RAW / "s36/ets/ets-power-share-2008.csv"
SHORTFALL_RULE = "0058 shortfall rule, option (b) (0110)"


@lru_cache(maxsize=None)
def _eea_free() -> dict[tuple[str, str, int], float]:
    """EEA Union Registry v82: {(country, activity, year): freely allocated allowances, t}."""
    it = openpyxl.load_workbook(EEA_REGISTRY, read_only=True, data_only=True)["Sheet1"].iter_rows(values_only=True)
    h = next(it)
    ix = {k: h.index(k) for k in ("country_code", "main_activity_code", "citl_information", "year", "value")}
    out = {}
    for r in it:
        if (r[ix["citl_information"]] == "1.1 Freely allocated allowances" and r[ix["country_code"]] in ("GB", "XI")
                and str(r[ix["year"]]).isdigit()):
            out[(r[ix["country_code"]], r[ix["main_activity_code"]], int(r[ix["year"]]))] = r[ix["value"]] or 0.0
    return out


def shortfall_share(y: int) -> dict:
    """0058's shortfall-rule sensitivity, computed by option (b) (decision 0110 item 1.19, Richard Lyon, 19 Oct 2026):
    the 0055 power proxy (DESNZ GHG 1.2 power stations less MSW) / UK verified EU ETS emissions less free
    allocation (EEA, the UK parts), 2013-20 only. A methods-page sensitivity on the split; never central."""
    if not 2013 <= y <= 2020:
        raise ValueError(f"shortfall rule: 2013-20 only, not {y}")
    s = stand_in_share(y)
    ver, free = _eea_verified(), _eea_free()
    den = sum(ver.get((c, a, y), 0.0) - free.get((c, a, y), 0.0) for c, a in EEA_UK_PARTS) / 1e6
    return {"value": s["numerator_mtco2"] / den, "numerator_mtco2": s["numerator_mtco2"],
            "denominator_mtco2": den, "rule": SHORTFALL_RULE, "numerator_source_id": SRC_GHG,
            "denominator_source_id": SRC_EEA}


def share_2008_record(path: Path = SHARE_2008) -> dict:
    """0068: the 2008 power share as entered, with its status, numerator, denominator and source ids."""
    (row,) = _rows(path)
    return row


UKETS_SHARE_2021_24 = RAW / "s36/ets/ukets-power-share-2021-24.csv"


def ukets_shares(path: Path = UKETS_SHARE_2021_24) -> dict[int, dict]:
    """2021-24 (0055): DESNZ printed UK ETS power and covered totals (2021: the ruled 0.428), refused unless SOURCED."""
    from gascf.pending import require_sourced
    return {int(r["year"]): r for r in require_sourced(_rows(path), "ukets_power_share_2021_24")}


def share_calibration() -> dict[int, float]:
    """0055's printed calibration: UK ETS power (DESNZ) / inventory power stations less MSW, 2022-24."""
    inv = _inventory_power_less_msw()
    return {y: float(r["power_tco2e"]) / 1e6 / inv[y] for y, r in ukets_shares().items() if y >= 2022}


def power_share(y: int, share_2008_file: Path = SHARE_2008) -> dict:
    """Power's share of UK auction revenue for year y (0058: its share of verified emissions).
    2008 (0068): the entered record, refused unless SOURCED (02's guard) and equal to the stand-in
    recomputed from the registered tables; 2009-20: the 0055 stand-in with its band; 2021-24: DESNZ
    printed (0055); 2025: the 2024 share, provisional (0079)."""
    if y >= 2021:
        r = ukets_shares()[min(y, 2024)]
        # 2022-24 the exact printed tonnes ratio (the file's value column is that ratio to 6 dp);
        # 2021 the ruled 0.428 (its power figure is a chart reading, 0055)
        v = float(r["value"]) if y == 2021 else float(r["power_tco2e"]) / float(r["total_covered_tco2e"])
        out = {"value": v, "rule": "0055 DESNZ printed", "source_id": r["source_id"],
               "numerator_source_id": r["source_id"], "denominator_source_id": r["source_id"],
               "status": r["status"]}
        if y == 2025:
            out |= {"rule": "0079: 2024 share", "status": "provisional"}
        return out
    s = stand_in_share(y)
    if y == 2008:
        from gascf.pending import require_sourced
        (rec,) = require_sourced([share_2008_record(share_2008_file)], "ets_power_share_2008")
        if abs(float(rec["value"]) - s["value"]) > 1e-9:
            raise ValueError(f"ets_power_share_2008: entered {rec['value']} != tables {s['value']:.9f}")
        s = s | {"status": rec["status"], "decision": "0068"}
    return s


# ------------------------------------------------------------------ UK ETS 2021-25 by calendar year (ticket 21)
UKETS_2021_24 = RAW / "s36/ets/ukets-auctions-2021-24.csv"
UKETS_2021_24_SHA256 = "8c2e0d47681f74201af20cf5528c410821dca5290374ef6bbe31edab66b581a9"   # data/raw/SHA256SUMS
UKETS_AUCTIONS_PER_YEAR = {2021: 16, 2022: 25, 2023: 25, 2024: 25}


def ukets_auction_rows(path: Path = UKETS_2021_24, sha256: str = UKETS_2021_24_SHA256) -> list[dict]:
    """The printed per-auction rows of the DESNZ reports for 2021-24, one row per auction."""
    if _sha(path) != sha256:
        raise ValueError(f"{Path(path).name}: sha256 differs from the registered file")
    return _rows(path)


def ukets_revenue_cy(path: Path = UKETS_2021_24, sha256: str = UKETS_2021_24_SHA256) -> dict[int, dict]:
    """0012/0055: {calendar year: {gbp, auctions}} UK ETS auction revenue, cash £, summed from the printed
    rows by the calendar year of the auction date. Refuses a file whose rows per year differ."""
    out: dict[int, dict] = {}
    seen: set[str] = set()
    for r in ukets_auction_rows(path, sha256):
        y = int(r["date"][:4])
        if r["date"] in seen:
            raise ValueError(f"{Path(path).name}: {r['date']} appears twice")
        if y != int(r["scheme_year"]):
            raise ValueError(f"{Path(path).name}: {r['date']} lies outside scheme year {r['scheme_year']}")
        seen.add(r["date"])
        o = out.setdefault(y, {"gbp": 0, "auctions": 0})
        o["gbp"] += int(r["revenue_gbp"])
        o["auctions"] += 1
    for y, n in UKETS_AUCTIONS_PER_YEAR.items():
        if out.get(y, {}).get("auctions") != n:
            raise ValueError(f"{Path(path).name}: {y} has {out.get(y, {}).get('auctions')} auctions, expected {n}")
    return out


UKETS_2025 = RAW / "s36/ets-2025/ice-report-278-2025.csv"
UKETS_2025_SHA256 = "4aab8a4a85d3389ee7c00623a6bd926a002deb41d2824e6b815df3a1154c9806"   # data/raw/SHA256SUMS
SRC_UKETS_2025 = "ice-report-278-ukets-auctions-2025"
NOTE_2025 = ("provisional (0079): calendar-2025 auction revenue from ICE report 278, to be replaced by DESNZ's "
             "figure (about 30 Nov 2026)")


def ukets_revenue_2025(path: Path = UKETS_2025, sha256: str = UKETS_2025_SHA256) -> dict:
    """0079: calendar-2025 UK ETS auction revenue (25 auctions, each sold out), cash £, provisional."""
    rows = _checked_rows(path, sha256, 25)
    dates = [r["date"] for r in rows]
    if len(set(dates)) != len(dates) or any(not d.startswith("2025-") for d in dates):
        raise ValueError(f"{Path(path).name}: a date twice or outside 2025")
    return {"gbp": round(sum(float(r["notional_gbp"]) for r in rows)), "auctions": len(rows),
            "allowances": sum(int(r["sold"]) for r in rows), "status": "provisional", "note": NOTE_2025,
            "source_id": SRC_UKETS_2025}


UK_REVENUE_2013_20 = RAW / "s36/ets/uk-auction-revenue-2013-20-calendar.csv"   # 0110 item 1.19: 2019-20 calendar
SRC_UKETS_2021_24 = "desnz-functioning-ukcm-2021-24-auction-tables"


def uk_revenue_2013_20(path: Path = UK_REVENUE_2013_20) -> dict[int, dict]:
    """2013-20 UK auction revenue, cash £m, as the engine has used it (EC platform reports 2013-18; BEIS
    Trust Statement printed FY2019-20 and FY2020-21 put on calendar 2019 = 0 and 2020 = 2,414.5 by decision 0110),
    from a registered s36 file with its period, not the Clock column."""
    from gascf.pending import require_sourced
    return {int(r["year"]): r | {"gbp_m": float(r["uk_revenue_gbp_m"])}
            for r in require_sourced([r | {"value": r["uk_revenue_gbp_m"]} for r in _rows(path)], "uk_revenue_2013_20")}
