"""System costs on the bill: constraint payments, offshore transmission, GB households and
the domestic share of electricity.

Spec §6, Layer 1, step 4(b). Four things live here because the tax and goods layers need the
household count and the domestic share too (one definition of each for the whole engine):

1. **Constraint payments.** What NESO (the National Energy System Operator, formerly
   National Grid ESO) pays generators to switch off or on because the transmission network
   cannot carry the power where it is wanted. Measured outturn by *financial year*
   (1 April to 31 March), read from NESO's "Constraint Breakdown" daily files, which give a
   cost and a volume (MWh) per day for each category NESO defines. NESO groups its four
   categories as "Inertia (reducing largest loss and increasing system inertia), Voltage, and
   Thermal". Inertia is stability, which spec §6 lists under "other renewables-driven
   balancing: not counted", so step 4(b) counts **thermal and voltage** only; the inertia
   group is read and reported (``by_type``) but kept out.

   **All-gas baseline (central, volume basis).** For each counted category, the 2017/18
   (earliest NESO year) volume is priced at *that year's* £/MWh for the category (cost /
   sum of |daily net volume|). Constraint £/MWh follows the gas price (offers come from gas
   plant, wind bids are priced against lost subsidy), so an all-gas Britain would have paid
   the crisis price on its own constraint volume; only the extra *volume* is counted against
   renewables. Wind was already on the system in 2017/18, so the baseline volume is if
   anything too high for an all-gas Britain, which shrinks the renewables line. A category
   with no usable price in a year (|volume| under 1% of its 2017/18 volume, or a cost <= 0)
   keeps its 2017/18 cost in real terms (GDP deflator) instead; this is our stated rule.
   **Alternative (P6 bracket):** the old basis, 2017/18 cost in real terms by the GDP
   deflator, returned as ``*_real_terms`` lines. It puts the 2021-23 gas-price effect into
   4(b), so it is the high end, not a cautious one.

   **Allocation to households.** Before 1 April 2023 BSUoS (Balancing Services Use of System,
   the charge that recovers balancing and constraint costs) was "recovered equally (50:50)
   from demand ... and generation" (CMP201 report, para 2.1, stored under data/raw/neso/).
   Only the demand half reached bills through BSUoS; the generator half sat in wholesale
   prices, which all-gas Britain replaces with its own fleet cost, so counting it here would
   count it twice. From 1 April 2023 BSUoS is charged on final demand only (Ofgem CMP308
   decision), so the demand share is 1. Per household two ways:
     * ``per_2500kwh``: £ per MWh of GB final consumption x 2.5 MWh (the receipt's basis);
       GB = DUKES 5.1 UK final consumption x the GB share of DESNZ subnational consumption
       (Table 7.1, GB / (GB + NI));
     * ``per_avg_household``: domestic share x cost / GB households.
   Financial year *t*/t+1 takes 9/12 of calendar year t and 3/12 of t+1 for demand,
   households and the GDP deflator.

2. **Offshore transmission (OFTO) revenue, demand share** (labelled estimated). OFTOs own
   the cables from offshore wind farms to shore. NESO's Final TNUoS tariff tables give the
   total OFTO revenue passed through (Table 33) and how much of it offshore generators pay
   directly through offshore local tariffs (Table 23, ``O``). The rest is recovered from the
   general pot, which demand pays in proportion D.R / (MAR - all local charges). Domestic
   customers pay the domestic band's share of the demand residual (Table B consumption
   proportion), and NESO's own domestic site count (Table B) turns it into £ per household.
   Local charges paid by offshore generators are not counted here (they reach bills, if at
   all, through wholesale prices); an upper bound counting them is returned alongside.

3. **GB household count** by year: ``gb_households`` reads the engine's one count,
   ``households.gb_parts`` (ONS LFS April-June for England, Wales and Scotland; s30).

4. **Domestic share of electricity**: DUKES 2026 Table 5.2 "Domestic" public-distribution
   row / Table 5.5.D "UK electricity sales" (a share of electricity *sold*, the volume the
   levies and BSUoS are charged on). ``domestic_share_cy`` is the engine's only definition.

Every reader checks its anchor labels and raises if a layout moves.
"""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from datetime import date, timedelta
from functools import lru_cache
from pathlib import Path

import openpyxl

from .deflator import read_deflator

ROOT = Path(__file__).parents[2]
RAW = ROOT / "data/raw"
CONSTRAINT_DIR = RAW / "neso/constraint_breakdown"
DAILY_COST_DIR = RAW / "neso/daily_balancing_costs"
ABC_WORKBOOK = RAW / "neso/NESO_2025_Annual_Balancing_Costs_Data_Workbook.xlsx"
TNUOS_DIR = RAW / "neso/tnuos"
DUKES_51 = RAW / "dukes/DUKES_5.1.xlsx"
DUKES_52 = RAW / "dukes/DUKES_5.2.xlsx"
DUKES_55 = RAW / "dukes/DUKES_5.5.xlsx"
SUBNATIONAL = RAW / "desnz_subnational/subnational-electricity-and-gas-consumption-summary-report-2024.pdf"
CMP201 = RAW / "neso/CMP201_document_129056.pdf"
CMP201_QUOTE = ("which is recovered equally (50:50) from demand (represented by Suppliers) and "
                "generation (represented by Generators)")

# NESO's four categories (cost and volume columns). NESO groups them as "Inertia (reducing
# largest loss and increasing system inertia), Voltage, and Thermal" (dataset resource
# description, constraint_breakdown_package_show.json). Inertia is stability, "not counted"
# in spec §6, so only voltage and thermal enter step 4(b); the inertia group is reported.
CATEGORIES = ("Reducing largest loss", "Increasing system inertia", "Voltage constraints",
              "Thermal constraints")
COUNTED = ("Voltage constraints", "Thermal constraints")
NOT_COUNTED = ("Reducing largest loss", "Increasing system inertia")
NESO_GROUPING_QUOTE = ("Inertia (reducing largest loss and increasing system inertia), Voltage, "
                       "and Thermal")
PACKAGE_SHOW = RAW / "neso/constraint_breakdown_package_show.json"
COST_COLUMNS = tuple(f"{c} cost" for c in CATEGORIES)
VOLUME_COLUMNS = tuple(f"{c} volume" for c in CATEGORIES)
BASELINE_FY = 2017  # earliest financial year on the NESO data portal (2017/18)
MIN_PRICE_VOLUME = 0.01  # a category's year price is usable if |volume| >= 1% of its 2017/18 volume
BSUOS_DEMAND_SHARE_PRE_CMP308 = 0.5   # CMP201 para 2.1 "recovered equally (50:50)"
CMP308_FROM_FY = 2023                 # BSUoS on final demand only from 1 April 2023

SRC_CONSTRAINT = "neso_constraint_breakdown"
SRC_DUKES51 = "desnz_dukes_2026_t5.1"
SRC_SUBNAT = "desnz_subnational_consumption_2024_t7.1"
SRC_CMP201 = "neso_cmp201_report"
SRC_TNUOS = "neso_final_tnuos_tables_{fy}"


def fy_label(fy: int) -> str:
    """2017 -> '2017/18'."""
    return f"{fy}/{(fy + 1) % 100:02d}"


@dataclass(frozen=True)
class Line:
    """One output figure with its provenance."""
    value: float
    unit: str
    basis: str            # allowance | measured | modelled | estimated
    source_id: str
    note: str = ""


# --------------------------------------------------------------------------- constraints

@dataclass(frozen=True)
class ConstraintYear:
    """One NESO financial year. ``cost_gbp`` is the *counted* categories only (no inertia)."""
    fy: int
    cost_gbp: float
    days: int
    complete: bool
    by_type: dict[str, float] = field(default_factory=dict)      # £ by cost column, all four
    volume_mwh: dict[str, float] = field(default_factory=dict)   # sum of |daily net| MWh, by category
    missing_dates: tuple[date, ...] = ()                          # days absent from NESO's file

    @property
    def cost_all_categories(self) -> float:
        return sum(self.by_type.values())

    def category_cost(self, cat: str) -> float:
        return self.by_type[f"{cat} cost"]

    def price(self, cat: str) -> float:
        """£ per MWh of |daily net volume| for one category (nan if no volume)."""
        v = self.volume_mwh[cat]
        return self.category_cost(cat) / v if v else float("nan")


def _parse_day(s: str) -> date:
    s = s.strip()[:10]
    if "/" in s:
        d, m, y = s.split("/")
        return date(int(y), int(m), int(d))
    y, m, d = s.split("-")
    return date(int(y), int(m), int(d))


def _fy(d: date) -> int:
    return d.year if d.month >= 4 else d.year - 1


@lru_cache(maxsize=4)
def read_constraint_breakdown(directory: Path = CONSTRAINT_DIR) -> dict[int, ConstraintYear]:
    """{financial year start: ConstraintYear} summed from NESO's daily constraint files (£)."""
    out: dict[int, ConstraintYear] = {}
    files = sorted(directory.glob("constraint-breakdown-*.csv"))
    if not files:
        raise FileNotFoundError(f"no constraint-breakdown files in {directory}")
    for f in files:
        with open(f, encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.DictReader(fh))
        missing = [c for c in ("Date", *COST_COLUMNS, *VOLUME_COLUMNS) if c not in rows[0]]
        if missing:
            raise ValueError(f"{f.name}: layout changed, missing columns {missing}")
        fy = int(f.stem.split("-")[2])
        by_type = {c: 0.0 for c in COST_COLUMNS}
        vol = {c: 0.0 for c in CATEGORIES}
        days = set()
        for r in rows:
            d = _parse_day(r["Date"])
            if _fy(d) != fy:
                raise ValueError(f"{f.name}: date {d} outside FY {fy_label(fy)}")
            days.add(d)
            for c in COST_COLUMNS:
                by_type[c] += float(r[c] or 0)
            for c in CATEGORIES:
                vol[c] += abs(float(r[f"{c} volume"] or 0))
        start, n_days = date(fy, 4, 1), (date(fy + 1, 4, 1) - date(fy, 4, 1)).days
        last = max(days)
        span = n_days if len(days) >= n_days - 2 else (last - start).days + 1
        gaps = tuple(d for d in (start + timedelta(i) for i in range(span)) if d not in days)
        counted = sum(by_type[f"{c} cost"] for c in COUNTED)
        out[fy] = ConstraintYear(fy, counted, len(days), len(days) >= n_days - 2, by_type, vol, gaps)
    return out


def read_daily_balancing_constraints(directory: Path = DAILY_COST_DIR) -> dict[int, float]:
    """{FY: £} from NESO's half-hourly 'Daily Balancing Costs' files, 'Constraints' column."""
    out: dict[int, float] = {}
    for f in sorted(directory.glob("q1_cost_fy*.csv")):
        with open(f, encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.DictReader(fh))
        if "Constraints" not in rows[0] or "SETT_DATE" not in rows[0]:
            raise ValueError(f"{f.name}: layout changed")
        for r in rows:
            fy = _fy(_parse_day(r["SETT_DATE"]))
            out[fy] = out.get(fy, 0.0) + float(r["Constraints"] or 0)
    return out


def read_abc_constraints(path: Path = ABC_WORKBOOK) -> dict[int, float]:
    """{FY: £} from NESO 2025 Annual Balancing Costs Report data workbook, 'Total outturn'.

    NESO's report category 'Constraints' is wider than the portal's breakdown (it adds
    constraint-related trades and services); used only as a cross-check.
    """
    rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True)["Total outturn"]
                .iter_rows(values_only=True))
    hdr = next((i for i, r in enumerate(rows) if r[0] == "FY"), None)
    if hdr is None or rows[hdr][1] != "Constraints (£)":
        raise ValueError("NESO ABC workbook layout changed: 'FY' / 'Constraints (£)' header not found")
    out = {}
    for r in rows[hdr + 1:]:
        if not r[0]:
            break
        out[int(str(r[0])[:4])] = float(r[1])
    return out


def _real(value: float, from_year: int, to_year: int, idx: dict[int, float]) -> float:
    return value * idx[to_year] / idx[from_year]


def fy_index(fy: int, idx: dict[int, float]) -> float:
    """GDP deflator for financial year fy/fy+1: 9/12 of calendar year fy + 3/12 of fy+1."""
    return 0.75 * idx[fy] + 0.25 * idx[fy + 1]


def _fy_real(value: float, fy: int, price_year: int, idx: dict[int, float]) -> float:
    """£ of financial year ``fy`` restated in calendar-year ``price_year`` prices."""
    return value * idx[price_year] / fy_index(fy, idx)


def constraint_baseline(price_year: int, idx: dict[int, float] | None = None) -> Line:
    """ALTERNATIVE all-gas baseline (P6 bracket, high end): the 2017/18 counted constraint
    cost held constant in real terms (GDP deflator, financial-year weighted).

    This puts every rise in the constraint £/MWh, including the 2021-23 gas-price effect, into
    step 4(b), so it is *not* cautious. The central basis is ``constraint_all_gas_volume``.
    """
    idx = idx or read_deflator()[0]
    years = read_constraint_breakdown()
    earliest = min(fy for fy, y in years.items() if y.complete)
    if earliest != BASELINE_FY:
        raise ValueError(f"earliest complete NESO year moved to {fy_label(earliest)}")
    v = _fy_real(years[earliest].cost_gbp, earliest, price_year, idx)
    return Line(v, f"£ ({price_year} prices)", "estimated", SRC_CONSTRAINT,
                f"alternative basis: FY {fy_label(earliest)} thermal + voltage outturn held in real "
                "terms (GDP deflator); counts the gas-price effect as renewables (P6 high end)")


def constraint_all_gas_volume(fy: int) -> tuple[float, dict[str, dict]]:
    """CENTRAL all-gas constraint cost for FY ``fy``, nominal £ of that year.

    For each counted category: 2017/18 volume x this year's £/MWh for the category. If the
    category has no usable price this year (|volume| < 1% of its 2017/18 volume, or cost <= 0)
    its 2017/18 cost is held in real terms instead, and the detail says so.
    """
    years = read_constraint_breakdown()
    base, cy = years[BASELINE_FY], years[fy]
    idx = read_deflator()[0]
    total, detail = 0.0, {}
    for c in COUNTED:
        bv, v, cost = base.volume_mwh[c], cy.volume_mwh[c], cy.category_cost(c)
        if v >= MIN_PRICE_VOLUME * bv and cost > 0:
            price = cost / v
            ag, how = bv * price, "2017/18 volume x this year's £/MWh"
        else:
            price = float("nan")
            ag = base.category_cost(c) * fy_index(fy, idx) / fy_index(BASELINE_FY, idx)
            how = "no usable price this year: 2017/18 cost in real terms"
        detail[c] = {"actual_gbp": cost, "volume_mwh": v, "price_gbp_mwh": price,
                     "baseline_volume_mwh": bv, "all_gas_gbp": ag, "method": how}
        total += ag
    return total, detail


# --------------------------------------------------------------------------- DUKES demand

@lru_cache(maxsize=2)
def read_dukes_final_and_domestic(path: Path = DUKES_51) -> dict[int, tuple[float, float]]:
    """{calendar year: (final consumption GWh, domestic GWh)}, UK, DUKES 2026 Table 5.1."""
    rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True)["5.1"].iter_rows(values_only=True))
    hdr = next((i for i, r in enumerate(rows) if r[0] == "Category"), None)
    if hdr is None or not str(rows[0][0]).startswith("Table 5.1 electricity commodity balance (GWh)"):
        raise ValueError("DUKES 5.1 layout changed")
    cols = {int(v): j for j, v in enumerate(rows[hdr]) if v and str(v).isdigit()}
    fc = next(r for r in rows if r[0] == "Final consumption")
    dom = next(r for r in rows if r[0] == "Domestic")
    return {y: (float(fc[j]), float(dom[j])) for y, j in cols.items()}


@lru_cache(maxsize=1)
def read_subnational_gb_share(path: Path = SUBNATIONAL) -> dict[int, float]:
    """{calendar year: GB / (GB + NI)} electricity consumption, DESNZ subnational statistics
    2024 (published 18 Dec 2025), Table 7.1 "Subnational: GB" and "Subnational: NI" (GWh)."""
    import pymupdf
    t = re.sub(r"\s+", " ", " ".join(p.get_text() for p in pymupdf.open(path)))
    i = t.find("Table 7.1: Total electricity consumption (GWh)")
    j = t.find("Table 7.2", i)
    if i < 0 or j < 0 or "Subnational: GB Subnational: NI Subnational: GB + NI" not in t[i:j]:
        raise ValueError("DESNZ subnational Table 7.1 anchor missing; layout changed")
    out = {}
    for y, gb, ni, both in re.findall(r"(20\d\d) ([\d,]+) ([\d,]+) ([\d,]+) [\d,]+", t[i:j]):
        gb, ni, both = (float(x.replace(",", "")) for x in (gb, ni, both))
        if abs(gb + ni - both) > 1.5:
            raise ValueError(f"Table 7.1 {y}: GB + NI != GB + NI column")
        out[int(y)] = gb / both
    if not out:
        raise ValueError("DESNZ subnational Table 7.1: no rows read")
    return out


def gb_share_of_uk(year: int) -> tuple[float, str]:
    """GB / UK electricity consumption for a calendar year (DESNZ subnational Table 7.1), used to
    put UK-only cost lines (RO, CCL) on the GB basis of every other line (s30). Years outside the
    table take the nearest published year, tagged estimated. For the RO this slightly understates
    GB's part (the NIRO obligation per MWh is below GB's), the conservative side."""
    sh = read_subnational_gb_share()
    if year in sh:
        return sh[year], "modelled"
    return sh[min(sh, key=lambda k: (abs(k - year), k))], "estimated"


def gb_final_consumption_gwh(year: int) -> Line:
    """GB final electricity consumption, calendar year: DUKES 5.1 UK final consumption x the
    GB share of subnational consumption (Table 7.1). If the subnational table has no row for
    the year, the latest year's share is used and the line is tagged estimated."""
    dukes = read_dukes_final_and_domestic()
    if year not in dukes:
        raise KeyError(f"UNSOURCED: DUKES 5.1 has no {year}")
    sh = read_subnational_gb_share()
    if year in sh:
        return Line(dukes[year][0] * sh[year], "GWh", "modelled", f"{SRC_DUKES51}; {SRC_SUBNAT}",
                    f"UK final x GB share {sh[year]:.4f} ({year})")
    last = max(sh)
    return Line(dukes[year][0] * sh[last], "GWh", "estimated", f"{SRC_DUKES51}; {SRC_SUBNAT}",
                f"UK final x GB share {sh[last]:.4f} ({last}; {year} not yet published)")


def _dukes_row(path: Path, sheet: str, label: str, col_prefix: str) -> dict[int, float]:
    rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True)[sheet].iter_rows(values_only=True))
    hdr = next(r for r in rows if r[0] == "Category")
    row = next((r for r in rows if r[0] and str(r[0]).startswith(label)), None)
    if row is None:
        raise ValueError(f"{path.name}: row '{label}' missing")
    out = {}
    for h, v in zip(hdr, row):
        m = re.fullmatch(re.escape(col_prefix) + r"\(?(\d{4})\)?", str(h)) if h else None
        if m and isinstance(v, (int, float)):
            out[int(m.group(1))] = float(v)
    if not out:
        raise ValueError(f"{path.name}: no year columns for '{label}'")
    return out


@lru_cache(maxsize=1)
def domestic_share_cy() -> dict[int, float]:
    """The engine's one domestic share: calendar-year domestic share of UK electricity sales.

    Numerator: DUKES 2026 Table 5.2, row "Domestic [note 6]", column "Public distribution
    system (YYYY)". Denominator: DUKES 2026 Table 5.5.D, row "Gives UK electricity sales".
    A share of electricity *sold*, the volume levies and BSUoS are charged on (DUKES 5.1
    domestic / final consumption is 2 pp lower because final consumption includes on-site use).
    """
    dom = _dukes_row(DUKES_52, "5.2", "Domestic", "Public distribution system ")
    t55 = list(openpyxl.load_workbook(DUKES_55, read_only=True, data_only=True)["5.5"]
               .iter_rows(values_only=True))
    i = next((k for k, r in enumerate(t55) if r[0] and str(r[0]).startswith("Table 5.5.D")), None)
    if i is None:
        raise ValueError("DUKES 5.5: 'Table 5.5.D' anchor missing")
    hdr = t55[i + 1]
    row = next(r for r in t55[i:] if r[0] and str(r[0]).startswith("Gives UK electricity sales"))
    sales = {int(h): float(v) for h, v in zip(hdr[1:], row[1:]) if h and isinstance(v, (int, float))}
    return {y: dom[y] / sales[y] for y in dom if y in sales}


def domestic_share_records() -> dict[int, dict]:
    """0060: each year's own domestic share (DUKES 2026 5.2 domestic public supply / 5.5.D UK sales;
    supplier sales only, self-used generation excluded) with its own source id. No carry-forward."""
    return {y: {"value": v, "source_id": f"dukes-2026-5.2:{y};dukes-2026-5.5D:{y}"}
            for y, v in domestic_share_cy().items()}


def fy_weight(series: dict[int, float], fy: int) -> tuple[float, str]:
    """9/12 of calendar year fy + 3/12 of fy+1, `modelled`. If fy+1 is not yet published,
    calendar year fy alone, `estimated`. Raises KeyError if fy is missing."""
    if fy + 1 in series:
        return 0.75 * series[fy] + 0.25 * series[fy + 1], "modelled"
    if fy in series:
        return series[fy], "estimated"
    raise KeyError(fy)


# --------------------------------------------------------------------------- households

def gb_households(year: int) -> tuple[float, dict[str, Line]]:
    """GB household count for a calendar year: the engine's one count, `households.gb_parts`
    (ONS LFS April-June, England + Wales + Scotland; s30). Raises KeyError if unsourced."""
    from gascf import households as _hh
    parts = {n: Line(v, "households", "measured", src, "LFS April-June, survey estimate")
             for n, (v, src) in _hh.gb_parts(year).items()}
    return sum(p.value for p in parts.values()), parts


def gb_households_series() -> dict[int, tuple[float, str]]:
    """{calendar year: (GB households, basis)} for every year the count covers."""
    from gascf import households as _hh
    return {y: (_hh.gb_households_count(y), "measured") for y in _hh.read_households()}


def gb_households_fy(fy: int) -> tuple[float, str]:
    """GB households for financial year fy/fy+1 (9/12 + 3/12 weighting; see fy_weight)."""
    ser = gb_households_series()
    v, basis = fy_weight({y: hh for y, (hh, _) in ser.items()}, fy)
    used = (fy, fy + 1) if basis == "modelled" else (fy,)
    if any(ser[y][1] == "estimated" for y in used):
        basis = "estimated"
    return v, basis


# --------------------------------------------------------------------------- constraint per household

def bsuos_demand_share(fy: int) -> Line:
    """Share of BSUoS (and so of constraint cost) charged to demand and so to bills."""
    if fy >= CMP308_FROM_FY:
        return Line(1.0, "ratio", "measured", "ofgem_cmp308_decision",
                    "BSUoS on final demand only from 1 April 2023 (CMP308)")
    return Line(BSUOS_DEMAND_SHARE_PRE_CMP308, "ratio", "measured", SRC_CMP201,
                "before 1 April 2023 BSUoS was 'recovered equally (50:50) from demand ... and "
                "generation' (CMP201 report para 2.1); the generator half sat in wholesale prices")


def constraint_per_household(fy: int, price_year: int | None = None) -> dict[str, Line]:
    """Actual and all-gas constraint cost (thermal + voltage) for one financial year, per GB
    household, for the part charged to demand (``bsuos_demand_share``).

    Central all-gas basis: 2017/18 volume at this year's £/MWh (``all_gas_*``,
    ``difference_*``). Alternative: 2017/18 cost in real terms (``all_gas_real_terms_*``,
    ``difference_real_terms_*``), the P6 high end. Nominal unless ``price_year`` is given.
    Demand, households and the deflator are financial-year weighted (9/12 + 3/12).
    Raises for an incomplete year.
    """
    years = read_constraint_breakdown()
    if fy not in years:
        raise KeyError(f"UNSOURCED: no NESO constraint data for FY {fy_label(fy)}")
    cy = years[fy]
    if not cy.complete:
        raise ValueError(f"FY {fy_label(fy)} incomplete ({cy.days} days); not annualised")
    idx = read_deflator()[0]
    # nominal = £ of the financial year itself; else calendar-year price_year £
    to_py = 1.0 if price_year is None else idx[price_year] / fy_index(fy, idx)
    py = f"{price_year}" if price_year is not None else f"FY {fy_label(fy)} nominal"
    actual = cy.cost_gbp * to_py
    ag_nom, detail = constraint_all_gas_volume(fy)
    base_vol = ag_nom * to_py
    base_real = (years[BASELINE_FY].cost_gbp * fy_index(fy, idx) / fy_index(BASELINE_FY, idx)) * to_py
    share = bsuos_demand_share(fy)
    fc = {y: gb_final_consumption_gwh(y) for y in (fy, fy + 1)
          if y in read_dukes_final_and_domestic()}
    if fy not in fc:
        raise KeyError(f"UNSOURCED: DUKES 5.1 has no {fy}")
    final_gwh, fc_basis = fy_weight({y: ln.value for y, ln in fc.items()}, fy)
    if any(ln.basis == "estimated" for ln in fc.values()):
        fc_basis = "estimated"
    dom, dom_basis = fy_weight(domestic_share_cy(), fy)
    hh, hh_basis = gb_households_fy(fy)
    unit = f"£/household/yr ({py} prices)" if price_year is not None else f"£/household/yr ({py})"
    tag = "estimated" if "estimated" in (fc_basis, dom_basis, hh_basis) else "modelled"
    miss = (f"; NESO file lacks {len(cy.missing_dates)} day(s): "
            f"{', '.join(d.isoformat() for d in cy.missing_dates)} (counted as zero)"
            if cy.missing_dates else "")
    out = {
        "demand_share": share,
        "inertia_not_counted_total": Line(
            sum(cy.category_cost(c) for c in NOT_COUNTED) * to_py, f"£ ({py})", "measured",
            SRC_CONSTRAINT, "NESO 'Inertia' group (reducing largest loss + increasing system "
            "inertia): stability, not counted (spec §6)"),
        "gb_final_consumption_gwh": Line(final_gwh, "GWh", fc_basis, f"{SRC_DUKES51}; {SRC_SUBNAT}",
                                         f"FY {fy_label(fy)}, 9/12 + 3/12"),
        "domestic_share": Line(dom, "ratio", dom_basis, "DUKES26-5.2; DUKES26-5.5",
                               f"FY {fy_label(fy)}, domestic / UK sales"),
        "gb_households": Line(hh, "households", hh_basis, "gb_households", f"FY {fy_label(fy)}"),
    }
    for key, v, basis, note in (
            ("actual", actual, "measured", f"FY {fy_label(fy)} NESO thermal + voltage outturn{miss}"),
            ("all_gas", base_vol, "estimated", "2017/18 volume x this year's £/MWh, by category "
             "(central; " + "; ".join(f"{c.split()[0].lower()}: {d['method']}" for c, d in detail.items()) + ")"),
            ("all_gas_real_terms", base_real, "estimated", "alternative: 2017/18 cost in real terms "
             "(P6 high end; counts the gas-price effect as renewables)")):
        out[f"{key}_total"] = Line(v, f"£ ({py})", basis, SRC_CONSTRAINT, note)
        out[f"{key}_per_2500kwh"] = Line(
            v * share.value / (final_gwh * 1000) * 2.5, unit, tag, f"{SRC_CONSTRAINT}; {SRC_SUBNAT}",
            f"x demand share {share.value:g} / GB final consumption (FY) x 2.5 MWh")
        out[f"{key}_per_avg_household"] = Line(
            v * share.value * dom / hh, unit, tag, SRC_CONSTRAINT,
            f"x demand share {share.value:g} x domestic share {dom:.4f} / GB households (FY)")
    for k in ("per_2500kwh", "per_avg_household"):
        out[f"difference_{k}"] = Line(out[f"actual_{k}"].value - out[f"all_gas_{k}"].value, unit,
                                      tag, SRC_CONSTRAINT, "step 4(b), central (volume basis)")
        out[f"difference_real_terms_{k}"] = Line(
            out[f"actual_{k}"].value - out[f"all_gas_real_terms_{k}"].value, unit, tag, SRC_CONSTRAINT,
            "step 4(b), alternative real-terms basis (P6 high end)")
    return out


# --------------------------------------------------------------------------- any 12-month window
#
# The financial-year functions above refuse an incomplete year. For the current year the
# staircase instead uses the latest 12 months NESO has published ("trailing 12 months"): the
# same method, with every calendar-year weight (demand, households, deflator) taken from the
# months the window covers. A window of exactly one financial year gives 9/12 + 3/12, so the
# two routes agree on any complete year (tested).

@lru_cache(maxsize=4)
def read_constraint_days(directory: Path = CONSTRAINT_DIR) -> dict[date, dict[str, float]]:
    """{day: {cost column: £, category: |net volume| MWh}} from every NESO constraint file."""
    out: dict[date, dict[str, float]] = {}
    for f in sorted(directory.glob("constraint-breakdown-*.csv")):
        with open(f, encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.DictReader(fh))
        missing = [c for c in ("Date", *COST_COLUMNS, *VOLUME_COLUMNS) if c not in rows[0]]
        if missing:
            raise ValueError(f"{f.name}: layout changed, missing columns {missing}")
        for r in rows:
            d = _parse_day(r["Date"])
            day = out.setdefault(d, {**{c: 0.0 for c in COST_COLUMNS}, **{c: 0.0 for c in CATEGORIES}})
            for c in COST_COLUMNS:
                day[c] += float(r[c] or 0)
            for c in CATEGORIES:
                day[c] += abs(float(r[f"{c} volume"] or 0))
    return out


def latest_constraint_day() -> date:
    return max(read_constraint_days())


def trailing_year(end: date | None = None) -> tuple[date, date]:
    """(start, end) of the 12 months ending on ``end`` (default: latest stored NESO day)."""
    end = end or latest_constraint_day()
    try:
        start = end.replace(year=end.year - 1) + timedelta(1)
    except ValueError:            # 29 February
        start = date(end.year - 1, 3, 1)
    return start, end


def month_weights(start: date, end: date) -> dict[int, float]:
    """{calendar year: months of the window in that year / 12}; part months count by days.

    For 1 April t to 31 March t+1 this is exactly {t: 0.75, t+1: 0.25} (the FY rule)."""
    w: dict[int, float] = {}
    d = start
    while d <= end:
        m_end = (date(d.year + (d.month == 12), d.month % 12 + 1, 1) - timedelta(1))
        last = min(m_end, end)
        w[d.year] = w.get(d.year, 0.0) + ((last - d).days + 1) / m_end.day
        d = last + timedelta(1)
    tot = sum(w.values())
    return {y: v / tot for y, v in w.items()}


def window_weight(series: dict[int, float], w: dict[int, float]) -> tuple[float, str]:
    """Weighted value over the window's calendar years, `modelled`; a year not yet published
    takes the latest earlier year (`estimated`, as ``fy_weight`` does for FY t+1)."""
    v, basis = 0.0, "modelled"
    for y, wy in w.items():
        if y in series:
            v += wy * series[y]
        else:
            prior = [x for x in series if x < y]
            if not prior:
                raise KeyError(y)
            v += wy * series[max(prior)]
            basis = "estimated"
    return v, basis


def _window_label(start: date, end: date) -> str:
    return f"{start.isoformat()} to {end.isoformat()}"


def constraint_window(start: date, end: date) -> ConstraintYear:
    """Sum the NESO daily files over one 12-month window (as ``read_constraint_breakdown`` does
    for a financial year). ``fy`` is the FY the window ends in; ``complete`` means at most two
    days are absent (the FY rule)."""
    days = read_constraint_days()
    want = [start + timedelta(i) for i in range((end - start).days + 1)]
    have = [d for d in want if d in days]
    if not have:
        raise KeyError(f"UNSOURCED: no NESO constraint data for {_window_label(start, end)}")
    by_type = {c: sum(days[d][c] for d in have) for c in COST_COLUMNS}
    vol = {c: sum(days[d][c] for d in have) for c in CATEGORIES}
    gaps = tuple(d for d in want if d not in days)
    counted = sum(by_type[f"{c} cost"] for c in COUNTED)
    return ConstraintYear(_fy(end), counted, len(have), len(have) >= len(want) - 2, by_type, vol, gaps)


def constraint_per_household_window(start: date, end: date, price_year: int | None = None,
                                    ) -> dict[str, Line]:
    """``constraint_per_household`` for any one-year window (central volume basis and the
    real-terms alternative). Used for the current, incomplete financial year with the
    trailing 12 months to the latest stored NESO day; tagged `modelled` (or `estimated` where
    an input year is not yet published). Raises if the window is not one year or has gaps."""
    if trailing_year(end)[0] != start:
        raise ValueError(f"window {_window_label(start, end)} is not exactly 12 months")
    cy = constraint_window(start, end)
    if not cy.complete:
        raise ValueError(f"window {_window_label(start, end)} incomplete ({cy.days} days)")
    years = read_constraint_breakdown()
    base = years[BASELINE_FY]
    idx = read_deflator()[0]
    w = month_weights(start, end)
    w_idx, _ = window_weight(idx, w)
    to_py = 1.0 if price_year is None else idx[price_year] / w_idx
    lab = _window_label(start, end)
    py = f"{price_year}" if price_year is not None else f"{lab} nominal"
    # central all-gas: 2017/18 volume at the window's £/MWh, by category (same rule as FY)
    ag_nom, detail = 0.0, {}
    for c in COUNTED:
        bv, v, cost = base.volume_mwh[c], cy.volume_mwh[c], cy.category_cost(c)
        if v >= MIN_PRICE_VOLUME * bv and cost > 0:
            ag, how = bv * cost / v, "2017/18 volume x this year's £/MWh"
        else:
            ag = base.category_cost(c) * w_idx / fy_index(BASELINE_FY, idx)
            how = "no usable price this year: 2017/18 cost in real terms"
        detail[c] = how
        ag_nom += ag
    actual = cy.cost_gbp * to_py
    base_vol = ag_nom * to_py
    base_real = base.cost_gbp * w_idx / fy_index(BASELINE_FY, idx) * to_py
    share = bsuos_demand_share(_fy(start))
    if _fy(start) != _fy(end) and _fy(start) < CMP308_FROM_FY <= _fy(end):
        raise ValueError("window straddles the CMP308 change (1 April 2023); use financial years")
    dukes = read_dukes_final_and_domestic()
    fc = {y: gb_final_consumption_gwh(y) for y in dukes}
    final_gwh, fc_basis = window_weight({y: ln.value for y, ln in fc.items()}, w)
    if any(fc[y].basis == "estimated" for y in w if y in fc):
        fc_basis = "estimated"
    dom, dom_basis = window_weight(domestic_share_cy(), w)
    ser = gb_households_series()
    hh, hh_basis = window_weight({y: v for y, (v, _) in ser.items()}, w)
    if any(ser[y][1] == "estimated" for y in w if y in ser):
        hh_basis = "estimated"
    unit = f"£/household/yr ({py} prices)" if price_year is not None else f"£/household/yr ({py})"
    tag = "estimated" if "estimated" in (fc_basis, dom_basis, hh_basis) else "modelled"
    miss = (f"; NESO file lacks {len(cy.missing_dates)} day(s): "
            f"{', '.join(d.isoformat() for d in cy.missing_dates)} (counted as zero)"
            if cy.missing_dates else "")
    wtxt = ", ".join(f"{y}: {v:.3f}" for y, v in w.items())
    out = {
        "window": Line(cy.days, "days", "measured", SRC_CONSTRAINT, f"{lab}{miss}"),
        "demand_share": share,
        "gb_final_consumption_gwh": Line(final_gwh, "GWh", fc_basis, f"{SRC_DUKES51}; {SRC_SUBNAT}",
                                         f"calendar-year weights {wtxt}"),
        "domestic_share": Line(dom, "ratio", dom_basis, "DUKES26-5.2; DUKES26-5.5", f"weights {wtxt}"),
        "gb_households": Line(hh, "households", hh_basis, "gb_households", f"weights {wtxt}"),
    }
    for key, v, basis, note in (
            ("actual", actual, "measured", f"{lab} NESO thermal + voltage outturn{miss}"),
            ("all_gas", base_vol, "estimated", "2017/18 volume x this window's £/MWh, by category "
             "(central; " + "; ".join(f"{c.split()[0].lower()}: {h}" for c, h in detail.items()) + ")"),
            ("all_gas_real_terms", base_real, "estimated", "alternative: 2017/18 cost in real terms "
             "(P6 high end)")):
        out[f"{key}_total"] = Line(v, f"£ ({py})", basis, SRC_CONSTRAINT, note)
        out[f"{key}_per_2500kwh"] = Line(
            v * share.value / (final_gwh * 1000) * 2.5, unit, tag, f"{SRC_CONSTRAINT}; {SRC_SUBNAT}",
            f"x demand share {share.value:g} / GB final consumption (window) x 2.5 MWh")
        out[f"{key}_per_avg_household"] = Line(
            v * share.value * dom / hh, unit, tag, SRC_CONSTRAINT,
            f"x demand share {share.value:g} x domestic share {dom:.4f} / GB households (window)")
    for k in ("per_2500kwh", "per_avg_household"):
        out[f"difference_{k}"] = Line(out[f"actual_{k}"].value - out[f"all_gas_{k}"].value, unit,
                                      tag, SRC_CONSTRAINT, f"step 4(b), central, trailing 12 months {lab}")
        out[f"difference_real_terms_{k}"] = Line(
            out[f"actual_{k}"].value - out[f"all_gas_real_terms_{k}"].value, unit, tag, SRC_CONSTRAINT,
            "step 4(b), alternative real-terms basis (P6 high end)")
    return out


def constraint_per_household_current(fy: int, price_year: int | None = None) -> dict[str, Line]:
    """Financial year ``fy`` if NESO's file for it is complete; otherwise the trailing 12 months
    to the latest stored day, provided that day falls in ``fy`` (an incomplete current year).
    Raises for a year with no data or one already past without a complete file."""
    years = read_constraint_breakdown()
    if fy in years and years[fy].complete:
        return constraint_per_household(fy, price_year)
    last = latest_constraint_day()
    if _fy(last) != fy:
        raise KeyError(f"UNSOURCED: no NESO constraint data for FY {fy_label(fy)} "
                       f"(latest stored day {last.isoformat()})")
    s, e = trailing_year(last)
    return constraint_per_household_window(s, e, price_year)


# --------------------------------------------------------------------------- OFTO

@dataclass(frozen=True)
class TnuosYear:
    fy: int
    ofto_revenue: float          # £m, Table 33 'Offshore Transmission Pass-Through', this FY
    offshore_local: float        # £m, revenue from offshore local tariffs (O), January Final
    onshore_local: float         # £m, onshore local substation + circuit (LG + SG), January Final
    total_revenue: float         # £m, MAR, January Final
    demand_revenue: float        # £m, D.R, January Final
    domestic_proportion: float   # Table B domestic consumption proportion of the demand residual
    domestic_sites: float        # Table B domestic site count
    domestic_tdr_per_day: float  # Table B final domestic residual charge, £/site/day


def _tnuos_path(fy: int) -> Path:
    return TNUOS_DIR / f"Final_TNUoS_Tariffs_{fy}-{(fy + 1) % 100:02d}_Tables.xlsx"


def read_tnuos_year(fy: int) -> TnuosYear:
    """Read the figures we need from NESO's Final TNUoS tariff tables for one charging year."""
    path = _tnuos_path(fy)
    if not path.exists():
        raise KeyError(f"UNSOURCED: no Final TNUoS tables stored for {fy_label(fy)}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    def find_row(label: str, sheets=None):
        for n in sheets or wb.sheetnames:
            for r in wb[n].iter_rows(values_only=True):
                for j, v in enumerate(r):
                    if isinstance(v, str) and v.strip() == label:
                        return n, r, j
        raise ValueError(f"TNUoS {fy_label(fy)} layout changed: {label!r} not found")

    def last_num(r, j):  # January Final is the last of the four forecast columns
        nums = [v for v in r[j + 1:] if isinstance(v, (int, float))]
        if len(nums) < 4:
            raise ValueError(f"TNUoS {fy_label(fy)}: expected four forecast columns")
        return float(nums[3])

    _, r, j = find_row("Revenue recovered from offshore local tariffs (£m)")
    o = last_num(r, j)
    lg = last_num(*find_row("Revenue recovered from onshore local substation tariffs (£m)")[1:])
    sg = last_num(*find_row("Revenue recovered from onshore local circuit tariffs (£m)")[1:])
    mar = last_num(*find_row("Total Revenue (£m)")[1:])
    dr = last_num(*find_row("Revenue recovered from demand (£m)")[1:])

    # Table 33: header row 'Regulatory Year' gives the FY columns.
    t33 = next(n for n in wb.sheetnames if n.strip() in ("T33", "T32", "T29"))
    rows = list(wb[t33].iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r[0] == "Regulatory Year")
    col = next((k for k, v in enumerate(rows[hi]) if v == fy_label(fy)), None)
    pt = next((r for r in rows[hi:] if r[0] and str(r[0]).startswith("Offshore Transmission Pass-Through")), None)
    if col is None or pt is None:
        raise ValueError(f"TNUoS {fy_label(fy)}: Table 33 layout changed")
    ofto = float(pt[col])

    if "TB" not in wb.sheetnames:
        raise KeyError(f"UNSOURCED: TNUoS {fy_label(fy)} has no Table B (domestic band starts 2023/24)")
    tb = list(wb["TB"].iter_rows(values_only=True))
    th = next(i for i, r in enumerate(tb) if "Band" in r)
    hdr = [str(v).strip() if v else "" for v in tb[th]]
    dom = next(r for r in tb[th:] if "Domestic" in r)
    c_prop, c_site = hdr.index("Consumption Proportion %"), hdr.index("Site Count")
    tdr_cols = [k for k, h in enumerate(hdr) if "Final" in h and "TDR" in h or h.endswith("Final Tariffs")]
    tdr = float(dom[tdr_cols[-1]]) if tdr_cols else float("nan")
    return TnuosYear(fy, ofto, o, lg + sg, mar, dr, float(dom[c_prop]), float(dom[c_site]), tdr)


def ofto_demand_share(fy: int) -> dict[str, Line]:
    """OFTO revenue paid by GB domestic customers, per domestic site (household), nominal £."""
    t = read_tnuos_year(fy)
    src = SRC_TNUOS.format(fy=fy_label(fy))
    pooled = t.ofto_revenue - t.offshore_local
    demand_frac = t.demand_revenue / (t.total_revenue - t.offshore_local - t.onshore_local)
    demand_m = pooled * demand_frac
    dom_m = demand_m * t.domestic_proportion
    upper_m = (demand_m + t.offshore_local) * t.domestic_proportion
    unit = "£/household/yr (nominal)"
    return {
        "ofto_revenue_m": Line(t.ofto_revenue, "£m", "measured", src, "Table 33 pass-through"),
        "offshore_local_m": Line(t.offshore_local, "£m", "measured", src, "Table 23 'O'"),
        "demand_fraction_of_pool": Line(demand_frac, "ratio", "modelled", src, "D.R / (MAR - local charges)"),
        "demand_share_m": Line(demand_m, "£m", "estimated", src, "(OFTO - offshore local) x demand fraction"),
        "domestic_share_m": Line(dom_m, "£m", "estimated", src, "x Table B domestic consumption proportion"),
        "per_household": Line(dom_m * 1e6 / t.domestic_sites, unit, "estimated", src,
                              "/ Table B domestic site count"),
        "per_household_upper": Line(upper_m * 1e6 / t.domestic_sites, unit, "estimated", src,
                                    "upper bound: offshore local charges assumed passed through at the "
                                    "domestic residual proportion"),
    }
