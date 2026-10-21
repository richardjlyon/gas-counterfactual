"""How much the with-versus-without-net-zero answer depends on the gas price, and why our
all-gas cost of power is not DESNZ's ~£70/MWh (brief s12).

Two things live here.

**1. Gas-price sensitivity of the headline.** The headline is *what you pay* minus *step 1,
the pure gas price* (spec §2; carbon is a net-zero cost, never an offset). Here the
staircase for one cap period (default Q4 2026, 2,500 kWh, carbon first, fleets A and B) is
re-run with Ofgem's gas index replaced by a list of trial gas prices. Both columns move:

- *All-gas Britain* follows the fleet cost: the staircase is handed the trial gas price as
  if Ofgem had printed it, so fleet A or B, its carbon step and everything downstream are
  recomputed by the staircase's own code.
- *What you pay* follows the break-even regression already in ``breakeven.py``: Ofgem's
  electricity index E against its gas index G over the 32 cap periods since 2015, used
  *anchored* (the line through this period's observed point with the fitted slope), so at
  today's gas price nothing changes. Ofgem's wholesale line moves by
  uplift-and-losses x slope x (G − G_today) x consumption. The CfD line moves the other way,
  by the CfD volume share x the change in E (LCCC data, exactly as ``breakeven.GapModel``).
  Supplier percentage lines (headroom, EBIT, bad debt, payment-method percentage) are
  recomputed with Ofgem's own formulas (``staircase.apply_pct``) and VAT at the period rate.

*Slope* means how many £/MWh the electricity index moves for each £/MWh of gas. Two slopes
are reported: all 32 periods (central) and the 29 periods without the Energy Price
Guarantee, Oct 2022 – Jun 2023.

**2. Reconciliation to DESNZ.** DESNZ *Electricity Generation Costs 2025*, Annex A,
"Additional Estimates 2030", column "Gas CCGT 93% Load Factor": total £111/MWh of which
carbon £41, so £70 without carbon (2024 money). That is one new CCGT running 86% of the year
(Annex A's net load factor for that column), burning gas at DESNZ's own central price. Our
fleet costs £107/MWh (A) or £114/MWh (B) in Q4 2026. ``reconcile()`` walks from one to the
other in steps, each computed with the engine's own plant and fleet code from stored
sources, and the steps add up exactly (they telescope: each step is the next state minus
this one).

DESNZ's gas price is read from DESNZ *Fossil Fuel Price Assumptions 2025* (January 2026),
Table 1 (publication workbook, sheet "Gas"), p/therm, real 2024 prices, Assumption B (the
central case; the EGC 2025 report, p.20, says it uses "the 2025 fuel price projections").
p/therm is turned into £/MWh with Ofgem's own factor from Annex 2 tab 1b cell C7
("Converting pence per therm into £/MWh", 2.9307111).
"""
from __future__ import annotations

from contextlib import contextmanager
from dataclasses import dataclass, replace
from functools import lru_cache
from pathlib import Path

import openpyxl

from . import receipt_actual as ra
from . import staircase as sc

ROOT = Path(__file__).parents[2]
RAW = ROOT / "data/raw"
FFPA_XLSX = RAW / "desnz_ffpa/Fossil_Fuel_Price_Assumptions_2025_publication.xlsx"
ANNEX_A = RAW / "DESNZ_EGC2025_AnnexA.xlsx"
ANNEX2 = sc.ANNEX2

SRC_FFPA = "desnz-ffpa-2025"
SRC_EGC = "desnz-egc2025-annexA"
SRC_THERM = "ofgem-annex2-v1.30 (1b!C7)"

Q4_2026 = (2026, 10, 2026, 12)
TRIAL_PRICES = (20.0, 30.0, 60.0, 80.0, 100.0)     # £/MWh gas, the period's money (brief)
DESNZ_COLUMN = "Gas CCGT 93% Load Factor"
DESNZ_YEAR = 2030


# --------------------------------------------------------------------------- DESNZ inputs

@lru_cache(maxsize=1)
def therm_factor() -> float:
    """Ofgem's pence-per-therm to £/MWh divisor (Annex 2 tab 1b, C7)."""
    ws = openpyxl.load_workbook(ANNEX2, read_only=True, data_only=True)["1b Direct Fuel Cost Component"]
    rows = list(ws.iter_rows(min_row=7, max_row=7, min_col=2, max_col=3, values_only=True))
    label, val = rows[0]
    if str(label).strip() != "Converting pence per therm into £/MWh" or not isinstance(val, (int, float)):
        raise ValueError(f"Annex 2 1b C7 moved: {label!r} {val!r}")
    return float(val)


@lru_cache(maxsize=1)
def ffpa_gas() -> dict[str, dict[int, float]]:
    """DESNZ FFPA 2025 gas assumptions, p/therm real 2024: {'A'|'B'|'C': {year: value}}."""
    rows = list(openpyxl.load_workbook(FFPA_XLSX, read_only=True, data_only=True)["Gas"]
                .iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows) if r and len(r) > 1 and r[1] == "p/Therm")
    if [str(c).strip() for c in rows[hdr][2:5]] != ["A", "B", "C"] or \
            str(rows[hdr + 1][1]).strip() != "Real 2024 Prices":
        raise ValueError("FFPA 2025 Gas sheet layout moved")
    out: dict[str, dict[int, float]] = {"A": {}, "B": {}, "C": {}}
    for r in rows[hdr + 2:]:
        if not r or len(r) < 5 or not isinstance(r[1], int):
            continue
        for j, s in enumerate("ABC", start=2):
            out[s][r[1]] = float(r[j])
    if not {2021, 2030, 2050} <= set(out["B"]):
        raise ValueError("FFPA 2025 Gas sheet: years missing")
    return out


@dataclass(frozen=True)
class Quoted:
    value: float
    unit: str
    basis: str
    source_id: str
    note: str = ""


def desnz_gas(year: int = DESNZ_YEAR, case: str = "B", money_year: int = 2026) -> dict[str, Quoted]:
    """DESNZ's gas price assumption, as printed (p/therm, 2024 money), in £/MWh (2024 money)
    and in `money_year` money (HMT GDP deflator; 2026 is an OBR forecast year, flagged)."""
    from .deflator import read_deflator, rebase
    p = ffpa_gas()[case][year]
    g24 = p / therm_factor()
    _, fc = read_deflator()
    flag = f"; {money_year} deflator is an OBR forecast" if money_year in fc else ""
    return {
        "p_therm_2024": Quoted(p, "p/therm, real 2024", "measured", SRC_FFPA,
                               f"FFPA 2025 Table 1, Assumption {case}, {year}"),
        "gbp_mwh_2024": Quoted(g24, "£/MWh gas (gross CV), 2024 money", "modelled",
                               f"{SRC_FFPA}; {SRC_THERM}", "p/therm / 2.9307111"),
        "gbp_mwh": Quoted(rebase(g24, 2024, money_year), f"£/MWh gas, {money_year} money", "modelled",
                          f"{SRC_FFPA}; {SRC_THERM}; hmt-gdp-deflator-sep-2026", "rebased" + flag),
    }


def desnz_printed(column: str = DESNZ_COLUMN, year: int = DESNZ_YEAR) -> dict[str, float]:
    """DESNZ Annex A 'Additional Estimates <year>' rows for one column, £/MWh 2024 money
    ('<1' read as 0.5, the midpoint of what DESNZ's rounding allows)."""
    from .official import read_desnz_levelised
    rows = read_desnz_levelised(ANNEX_A, year)[column]
    return {k: (0.5 if v is None else v) for k, v in rows.items()}


# --------------------------------------------------------------------------- the actual side

def _slope(which: str) -> float:
    from .breakeven import fit_cap_periods
    return fit_cap_periods(exclude_epg={"all": False, "ex_epg": True}[which]).b


def actual_receipt_at_gas(k: tuple, gas: float, slope: str = "all", cfd_moves: bool = True) -> ra.Receipt:
    """Ofgem's receipt for period k with its wholesale line moved to gas price `gas` along the
    anchored regression; the CfD line moved the other way; percentage lines and VAT
    recomputed. Lines are £/household/yr at Ofgem's own consumption Q."""
    from .breakeven import cfd_volume
    r = ra.receipt(k)
    p = sc.annex2_periods()[k]
    Q = r.consumption_mwh
    ul = sum(p.supplied.values()) / len(p.supplied) / p.elec_index
    dE = _slope(slope) * (gas - p.gas_index)
    if p.elec_index + dE <= 0:
        raise ValueError(f"electricity index would be {p.elec_index + dE:.2f} at gas £{gas}: outside the fit")
    s = cfd_volume().share if cfd_moves else 0.0
    new = {"direct_fuel": ul * dE * Q, "cfd": -s * dE * Q}
    old = {ln.key: (ln.typical, ln.nil) for ln in r.lines if ln.key != "vat"}
    moved = {kk: ((v[0] + new.get(kk, 0.0)), v[1]) for kk, v in old.items()}
    rules = sc._rules_for(k)
    f0, f1 = sc.apply_pct(old, rules), sc.apply_pct(moved, rules)
    for code, key in sc.PCT.items():
        moved[key] = tuple(old[key][i] + f1[code][i] - f0[code][i] for i in (0, 1))
    lines = []
    for ln in r.lines:
        if ln.key == "vat":
            continue
        t, n = moved[ln.key]
        lines.append(replace(ln, typical=t, nil=n,
                             basis=ln.basis if (t, n) == (ln.typical, ln.nil) else "modelled"))
    ex_t = sum(ln.typical for ln in lines)
    ex_n = sum(ln.nil for ln in lines)
    vat = r.line("vat")
    lines.append(replace(vat, typical=ex_t * r.vat_rate, nil=ex_n * r.vat_rate,
                         basis=vat.basis if gas == p.gas_index else "modelled"))
    return replace(r, lines=tuple(lines), notes=r.notes + (
        f"sensitivity: wholesale moved to gas £{gas:.2f}/MWh along the anchored Ofgem regression "
        f"({slope} periods); CfD moved by LCCC volume share {s:.4f}; percentage lines and VAT recomputed",))


class _RaShim:
    """receipt_actual, with one period's receipt swapped (used only inside `_gas_override`)."""

    def __init__(self, k, receipt):
        self._k, self._r = k, receipt

    def receipt(self, period):
        key = period if isinstance(period, tuple) else ra.period_key(period)
        return self._r if key == self._k else ra.receipt(period)

    def __getattr__(self, name):
        return getattr(ra, name)


@contextmanager
def _gas_override(k: tuple, gas: float, slope: str, cfd_moves: bool):
    """Inside the block, the staircase sees Ofgem's gas index for period k as `gas` and the
    actual receipt moved to match. Everything is put back on exit, success or failure."""
    p0 = sc.annex2_periods()[k]
    periods = dict(sc.annex2_periods())
    periods[k] = replace(p0, gas_index=float(gas))
    shim = _RaShim(k, actual_receipt_at_gas(k, gas, slope, cfd_moves))
    orig_periods, orig_ra = sc.annex2_periods, sc.ra
    sc._core.cache_clear()
    sc.annex2_periods = lambda: periods
    sc.ra = shim
    try:
        yield
    finally:
        sc.annex2_periods, sc.ra = orig_periods, orig_ra
        sc._core.cache_clear()


def staircase_at_gas(k: tuple, gas: float, fleet: str, kwh: float = 2500.0, slope: str = "all",
                     cfd_moves: bool = True, order: str = "carbon_first") -> sc.Staircase:
    """The staircase for period k with gas at `gas` £/MWh (the period's money)."""
    with _gas_override(k, gas, slope, cfd_moves):
        return sc.staircase(k, fleet, kwh=kwh, order=order)


# --------------------------------------------------------------------------- the sweep

@dataclass(frozen=True)
class Row:
    label: str
    gas_gbp_mwh: float
    fleet: str
    slope: str
    elec_index: float              # Ofgem electricity index implied by the regression
    fleet_gbp_mwh: float           # all-gas cost of power at transmission, no carbon
    actual_total: float            # what you pay, £/household/yr
    step1: float                   # pure gas price, £/household/yr
    headline: float                # actual − step 1
    steps: dict
    identity_residual: float
    basis: str = "modelled"
    source_id: str = ""


def sweep(k: tuple = Q4_2026, kwh: float = 2500.0, fleets=("A", "B"), slopes=("all", "ex_epg"),
          prices: dict[str, float] | None = None, cfd_moves: bool = True) -> list[Row]:
    """Headline at each trial gas price, fleet and slope."""
    p = sc.annex2_periods()[k]
    if prices is None:
        d = desnz_gas(money_year=sc.price_year(k))["gbp_mwh"].value
        prices = {"DESNZ central (FFPA 2025 B, 2030)": d, "£20": 20.0, "£30": 30.0,
                  "current (Ofgem Q4 2026 index)": p.gas_index, "£60": 60.0, "£80": 80.0, "£100": 100.0}
    out = []
    for label, g in sorted(prices.items(), key=lambda kv: kv[1]):
        for sl in slopes:
            e = p.elec_index + _slope(sl) * (g - p.gas_index)
            for fl in fleets:
                s = staircase_at_gas(k, g, fl, kwh, sl, cfd_moves)
                out.append(Row(label, g, fl, sl, e, s.inputs["fleet_gbp_mwh"], s.actual_total,
                               s.steps["1"], s.actual_total - s.steps["1"], dict(s.steps),
                               s.identity_residual,
                               source_id=f"{sc.SRC_A2}; {sc.SRC_FLEET[fl]}; ofgem cap model; "
                                         f"lccc-actual-cfd-generation; {SRC_FFPA}"))
    return out


# --------------------------------------------------------------------------- reconciliation

def rebuild_rows(gas24: float | None = None) -> dict[str, tuple[float, float]]:
    """DESNZ's printed rows for the CCGT 93% column (2030) beside our rebuild from Annex A
    inputs: {row: (ours, printed)}, £/MWh 2024 money. Fuel uses DESNZ's central 2030 gas
    price unless `gas24` (£/MWh gas, 2024 money) is given."""
    from .plant import levelised_capital_gbp_mwh
    c = sc._fleet_kit()["ccgt"]
    g = desnz_gas()["gbp_mwh_2024"].value if gas24 is None else gas24
    pre, con = levelised_capital_gbp_mwh(c)
    fom = (c.fixed_om_gbp_mw_yr + c.insurance_gbp_mw_yr + c.connection_uos_gbp_mw_yr) / 1000 / (8.76 * c.net_load_factor)
    pr = desnz_printed()
    return {"Pre-development costs": (pre, pr["Pre-development costs"]),
            "Construction costs": (con, pr["Construction costs"]),
            "Fixed O&M costs": (fom, pr["Fixed O&M costs"]),
            "Variable O&M costs": (c.variable_om_gbp_mwh, pr["Variable O&M costs"]),
            "Fuel costs": (g / c.efficiency_gcv, pr["Fuel costs"])}


@dataclass(frozen=True)
class Step:
    name: str          # plain words
    delta: float       # £/MWh change this step makes
    level: float       # £/MWh after the step
    money: str         # price base of `level`
    basis: str
    source_id: str
    note: str = ""


def reconcile(k: tuple = Q4_2026) -> dict[str, list[Step]]:
    """Walk from DESNZ's £70/MWh (CCGT, 2030, 86% load factor, no carbon, 2024 money) to our
    fleet A and fleet B cost of power for period k (no carbon, the period's money).

    Everything is computed in 2024 money (Annex A's) and turned into the period's money in
    the last step. Each step = next state − this state, so the steps add up exactly."""
    from .deflator import rebase
    from .fleeta import cost_fleet_a
    from .fleetb import cost_fleet_b
    from .plant import annualised_fixed_gbp_kw_yr, levelised_capital_gbp_mwh, variable_gbp_mwh

    kit = sc._fleet_kit()
    ccgt, ocgt, tr = kit["ccgt"], kit["ocgt"], kit["transport"]
    if ccgt.name != DESNZ_COLUMN or ccgt.commissioning != DESNZ_YEAR:
        raise ValueError(f"staircase kit CCGT is {ccgt.name} {ccgt.commissioning}, not DESNZ's column")
    py = sc.price_year(k)
    to24, back = rebase(1.0, py, 2024), rebase(1.0, 2024, py)
    p = sc.annex2_periods()[k]
    g_now24 = p.gas_index * to24
    g_desnz24 = desnz_gas()["gbp_mwh_2024"].value
    lf = ccgt.net_load_factor
    hrs = 8.76 * lf                                     # MWh per kW per year

    printed = desnz_printed()
    desnz_ex_carbon = printed["Total"] - printed["Carbon costs"]

    def one_plant(gas: float, transport: bool) -> tuple[float, float]:
        """(fixed £/MWh, variable £/MWh) for one CCGT at Annex A's own load factor."""
        fx = annualised_fixed_gbp_kw_yr(ccgt) + (tr.fixed_gbp_kw_e_yr(ccgt.efficiency_gcv) if transport else 0)
        var = variable_gbp_mwh(ccgt, gas, tr.commodity_gbp_mwh_gas if transport else 0.0)
        return fx / hrs, var

    ldc = list(sc._ldc(sc.nearest(py, sc.RESIDUAL_YEARS)[0]))
    s0 = sum(one_plant(g_desnz24, False))
    s1 = sum(one_plant(g_now24, False))
    f2, v2 = one_plant(g_now24, True)
    s2 = f2 + v2
    nospare = cost_fleet_b(ldc, g_now24, ccgt, ocgt, -1.0, derating=kit["derating"], transport=tr)
    fixed_ns, var_ns = nospare.fixed_gbp_m / nospare.energy_twh, nospare.variable_gbp_m / nospare.energy_twh
    fb = cost_fleet_b(ldc, g_now24, ccgt, ocgt, kit["margin"], derating=kit["derating"], transport=tr)
    s_b = fb.gbp_per_mwh

    # fleet A: same inputs as the staircase (DUKES eta for the year, §8 margin, £5 VOM)
    s_a_full = sc.staircase(k, "A")
    eta, margin = s_a_full.inputs["eta_existing"], s_a_full.inputs["gas_market_margin_gbp_mwh"]
    m24 = rebase(margin, py, 2024)

    def fleet_a(eta_: float, m24_: float) -> float:
        return cost_fleet_a(ldc, g_now24, eta_, kit["vom"], m24_, kit["existing"], ccgt, ocgt,
                            kit["margin"], kit["derating"], transport=tr).gbp_per_mwh
    a_new_eta = fleet_a(ccgt.efficiency_gcv, 0.0)
    a_dukes_eta = fleet_a(eta, 0.0)
    a_full = fleet_a(eta, m24)

    pre, con = levelised_capital_gbp_mwh(ccgt)
    rows_sum = sum(v for kk, v in printed.items() if kk not in ("Total", "Carbon costs"))
    src_a = "desnz-egc2025-annexA (Additional Estimates 2030; Technical and Cost Assumptions)"
    common = [
        ("DESNZ, as printed: CCGT '93%' column, 2030, total less carbon", desnz_ex_carbon, "2024",
         "measured", src_a, f"total {printed['Total']:.0f} − carbon {printed['Carbon costs']:.0f}; "
                            f"runs {lf:.0%} of the year"),
        ("DESNZ's own printed rows without carbon, added up (DESNZ rounds each row to £1)", rows_sum,
         "2024", "measured", src_a, "pre-development '<1' read as 0.5"),
        ("Our rebuild of the same plant from Annex A inputs and DESNZ's central gas price",
         s0, "2024", "modelled", f"{src_a}; {SRC_FFPA}",
         f"capital {pre + con:.2f} + fixed running {s0 - pre - con - one_plant(g_desnz24, False)[1]:.2f} "
         f"+ variable {one_plant(g_desnz24, False)[1]:.2f} (gas £{g_desnz24:.2f}/MWh at η {ccgt.efficiency_gcv})"),
        ("Gas price: DESNZ's central assumption → Ofgem's Q4 2026 gas index", s1, "2024",
         "modelled", f"{SRC_FFPA}; {sc.SRC_A2}",
         f"£{g_desnz24:.2f} → £{g_now24:.2f}/MWh gas (2024 money), ÷ η {ccgt.efficiency_gcv}"),
        ("Gas pipeline charges (National Gas exit capacity + commodity), not in DESNZ's figure", s2,
         "2024", "modelled", "nationalgas-charges-oct-2026", "DESNZ's LCOE has no gas transport line"),
        ("Running part-time: a whole fleet follows demand, so capital is spread over fewer hours",
         s2 + (fixed_ns - f2), "2024", "modelled", "elexon-fuelhh+neso (residual 2025); desnz-egc2025-annexA",
         f"fleet fixed cost £{fixed_ns:.2f}/MWh vs £{f2:.2f} for one plant at {lf:.0%}"),
        ("Peaking plant (OCGT, η 0.36) burns more gas per MWh in the top slices", s2 + (fixed_ns - f2) + (var_ns - v2),
         "2024", "modelled", "desnz-egc2025-annexA",
         f"fleet running cost £{var_ns:.2f}/MWh vs £{v2:.2f}; {nospare.ccgt_gw:.1f} GW CCGT + {nospare.ocgt_gw:.1f} GW OCGT"),
    ]
    b_tail = [("Spare capacity: de-rated margin over peak (NESO ECR 2026), all paid for", s_b, "2024",
               "modelled", "neso-ecr-2026", f"{fb.reserve_gw:.1f} GW spare OCGT, fixed cost only")]
    a_tail = [
        (f"Today's {kit['existing']['ccgt'] / 1000:.1f} GW of CCGTs already exist: no build cost, but every MWh is paid the running "
         "cost of the dearest plant running (market price), plus spare capacity", a_new_eta, "2024",
         "modelled", "dukes-5.11; neso-ecr-2026", "at new-plant efficiency 0.54, no margin"),
        ("Today's plants are less efficient (DUKES 2025 fleet η)", a_dukes_eta, "2024", "modelled",
         "dukes-5.10", f"η 0.54 → {eta:.4f}"),
        ("Gas-market margin seen in the market when gas sets the price (spec §8)", a_full, "2024",
         "modelled", "elexon-mid-apx+nationalgas-sap (§8 validation)", f"£{margin:.2f}/MWh ({py} money)"),
    ]

    def build(states, final_staircase_value) -> list[Step]:
        out, prev = [], None
        for name, level, money, basis, src, note in states:
            out.append(Step(name, 0.0 if prev is None else level - prev, level, money, basis, src, note))
            prev = level
        last = prev * back
        out.append(Step(f"Price base: 2024 → {py} money (HMT GDP deflator)", last - prev, last, str(py),
                        "modelled", "hmt-gdp-deflator-sep-2026", f"× {back:.4f}"))
        if abs(last - final_staircase_value) > 1e-6:
            raise ValueError(f"reconciliation ends at {last}, staircase says {final_staircase_value}")
        return out

    s_b_full = sc.staircase(k, "B")
    return {"B": build(common + b_tail, s_b_full.inputs["fleet_gbp_mwh"]),
            "A": build(common + a_tail, s_a_full.inputs["fleet_gbp_mwh"])}


# --------------------------------------------------------------------------- printing

def table(rows: list[Row]) -> str:
    out = [f"{'gas price':38s} {'£/MWh':>7s} {'slope':>6s} {'fleet':>5s} {'E idx':>7s} {'fleet£/MWh':>10s} "
           f"{'you pay':>8s} {'pure gas':>8s} {'headline':>8s}"]
    for r in rows:
        out.append(f"{r.label[:38]:38s} {r.gas_gbp_mwh:7.2f} {r.slope:>6s} {r.fleet:>5s} {r.elec_index:7.2f} "
                   f"{r.fleet_gbp_mwh:10.2f} {r.actual_total:8.2f} {r.step1:8.2f} {r.headline:8.2f}")
    return "\n".join(out)


if __name__ == "__main__":
    print(table(sweep()))
    for fl, steps in reconcile().items():
        print(f"\nfleet {fl}")
        for s in steps:
            print(f"  {s.delta:+8.2f}  → {s.level:7.2f} ({s.money})  {s.name}")
