"""Fossil fleet of fossil Britain: coal and gas, 2002-2025, no carbon (spec v4 §4.2-§4.4, §15.1 decision 1).

Plain words. The fleet must make the residual demand (what nuclear, pumped storage and
interconnectors did not). It is built of three plant types:

- coal: supercritical coal with flue-gas desulphurisation (FGD) and NOx control, the cost of
  meeting air-pollution law. Capped each year at the coal capacity that actually existed in
  Great Britain (DUKES 5.7 less Northern Ireland's Kilroot); no new coal is invented.
- CCGT (combined-cycle gas turbine) and OCGT (open-cycle gas turbine, a peaker).

Every plant is priced as if built new from the official cost set in force that year (§4.3),
restated from the set's own price base to year-y money with CPIH. Sizing is the screening
curve of `fleetb.py`, extended to three plant types with coal capped. Fuel is then charged by
half-hourly dispatch in merit order at that month's fuel price (§4.2). No carbon price.

"Gross CV" (higher heating value, HHV) is the basis of UK fuel prices. Efficiencies published
on a net-CV (lower heating value, LHV) basis are converted with DUKES Annex A's own ratio of
net to gross calorific value for that fuel and year. Mixing bases raises an error.

Outputs are in money of year y (nominal). Every number carries `basis` and `source_id`.
"""
from __future__ import annotations

import bisect
import csv
import json
import math
import re
from dataclasses import dataclass, field, replace
from datetime import date
from functools import lru_cache
from pathlib import Path

import openpyxl

from .deflator_cpih import cpih
from .official import Plant, read_desnz_plants
from .plant import annualised_fixed_gbp_kw_yr

ROOT = Path(__file__).parents[2]
RAW = ROOT / "data/raw"
EGC = RAW / "desnz_egc"
EDITIONS_FILE = ROOT / "reference/cost_editions.yaml"
CALORIFIC = RAW / "dukes/s20b-DUKES_A.1-A.3.xlsx"
QEP = RAW / "desnz_qep/s18a-QEP_table_321.xlsx"
WMI = RAW / "ofgem/s18a-WMI_gas_dayahead_monthly_YFsS4g98P.js"
DUKES_57 = RAW / "DUKES_5.7.xlsx"
DUKES_511 = RAW / "DUKES_5.11.xlsx"
KWH_PER_THERM = 29.3071        # 1 therm = 29.3071 kWh; spec §4.4 writes "÷ 2.9307" for p/therm -> £/MWh

# De-rating (share of capacity counted at peak). NESO ECR 2026 Figure A10 (capacity.py) has
# no coal row: coal is given the CCGT factor, the lower of the fossil factors, which needs
# more spare plant and so makes the fossil fleet dearer (conservative rule). Labelled estimated.
COAL_DERATING_BASIS = "estimated"


# ----------------------------------------------------------------- calorific values

@lru_cache(maxsize=None)
def ncv_gcv_ratio() -> dict[str, dict[int, float]]:
    """{'gas': {year: NCV/GCV}, 'coal': {...}} from DUKES Annex A, Tables A.1 (gross) and A.2 (net).

    Gas: 'Natural Gas consumed' (MJ/m3). Coal: 'Power stations [note 2]' (GJ/t)."""
    wb = openpyxl.load_workbook(CALORIFIC, read_only=True, data_only=True)

    def series(sheet: str, title: str, label: str) -> dict[int, float]:
        rows = list(wb[sheet].iter_rows(values_only=True))
        t = next((i for i, r in enumerate(rows) if r[0] and str(r[0]).startswith(title)), None)
        if t is None:
            raise ValueError(f"DUKES Annex A layout changed: {title!r}")
        hdr = next(r for r in rows[t + 1:] if r[0] and str(r[1]).strip().isdigit())
        r = next((r for r in rows[t + 1:t + 40] if r[0] and str(r[0]).strip() == label), None)
        if r is None:
            raise ValueError(f"DUKES Annex A: row {label!r} missing under {title!r}")
        return {int(h): float(v) for h, v in zip(hdr[1:], r[1:])
                if h and str(h).strip().isdigit() and isinstance(v, (int, float))}

    out = {}
    for fuel, (tg, tn, lab) in {
        "gas": ("Table A.1.3", "Table A.2.3", "Natural Gas consumed"),
        "coal": ("Table A.1.1", "Table A.2.1", "Power stations [note 2]"),
    }.items():
        g, n = series("A.1", tg, lab), series("A.2", tn, lab)
        out[fuel] = {y: n[y] / g[y] for y in g if y in n}
    return out


def to_gross(eff: float, basis: str, fuel: str, year: int, aux: float = 0.0) -> float:
    """Efficiency on gross CV, sent-out (net of the plant's own use).

    basis: 'gross_cv' (already HHV, sent out), 'net_lhv' (LHV, sent out) or
    'gross_output_lhv' (LHV on generator output; Mott MacDonald 2010, `aux` its auxiliary power)."""
    r = ncv_gcv_ratio()[fuel]
    ratio = r[year] if year in r else r[min(r, key=lambda y: abs(y - year))]  # gas table starts 2007
    if basis == "gross_cv":
        return eff
    if basis == "net_lhv":
        return eff * ratio
    if basis == "gross_output_lhv":
        return eff * (1 - aux) * ratio
    raise ValueError(f"unknown efficiency basis {basis!r}")


# ----------------------------------------------------------------- cost editions

@lru_cache(maxsize=None)
def cost_editions() -> dict:
    """The hand-entered PDF tables (JSON, valid YAML). Tests check each row against its PDF."""
    return json.loads(EDITIONS_FILE.read_text())


def _even(n_years: float) -> tuple[float, ...]:
    """Equal spending over the whole years spanned; used where an edition gives only a duration."""
    n = max(1, round(n_years))
    return tuple([1.0 / n] * n)


def plant_from_table(key: str, hurdle: float, eff_year: int) -> Plant:
    """A `Plant` from one PDF table (medium column). Phasing is even over the stated periods
    (the tables give durations, not phasing): labelled in `spec-fossil.md`."""
    t = cost_editions()["tables"][key]
    r = {k: v["medium"] for k, v in t["rows"].items()}
    fuel = "coal" if t["tech"] == "coal" else "gas"
    size = r.get("size_mw", 1000.0)
    infra_k = r["infrastructure_gbp_k"] if "infrastructure_gbp_k" in r else r.get("infrastructure_gbp_kw", 0.0) * size
    eff = to_gross(r["efficiency"], t["eff_basis"], fuel, eff_year, r.get("auxiliary_power", 0.0)) \
        if "efficiency" in r else 1.0
    one = lambda v: {"low": v, "medium": v, "high": v}  # noqa: E731
    return Plant(
        name=key, commissioning=t["price_base"], size_mw=size, net_load_factor=0.93,
        efficiency_gcv=eff, life_years=int(r["life_years"]),
        predev_gbp_kw=one(r["prelicensing_gbp_kw"] + r["regulatory_gbp_kw"]),
        predev_phasing=_even(r["predev_years"]),
        construction_gbp_kw=one(r["epc_gbp_kw"]), construction_phasing=_even(r["construction_years"]),
        infrastructure_gbp_k=one(infra_k),
        fixed_om_gbp_mw_yr=r.get("fixed_om_gbp_mw_yr", 0.0), variable_om_gbp_mwh=r["variable_om_gbp_mwh"],
        insurance_gbp_mw_yr=r.get("insurance_gbp_mw_yr", 0.0),
        connection_uos_gbp_mw_yr=r.get("connection_uos_gbp_mw_yr", 0.0),
        hurdle=one(hurdle))


def _pct(v) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v or "").strip().rstrip("%")
    return float(s) / 100 if s and re.fullmatch(r"[\d.]+", s) else None


def read_egc2016(tech: str) -> Plant:
    """BEIS EGC 2016 Annexes, Table 19, commissioning 2018 column, medium case (2014 prices).

    Block layout (checked): name row; size/load factor/efficiency in the column right of the
    name; durations and spend shares in the name column and the two to its right; costs in
    the 'Commissioning year' row's 2018 column. Hurdle 7.8% ('CCGT & OCGT', report Table 17)."""
    wb = openpyxl.load_workbook(EGC / "s18d-BEIS_EGC2016_Annexes.xlsx", read_only=True, data_only=True)
    t = list(wb["Table 19"].iter_rows(values_only=True))
    if not str(t[0][0]).startswith("Table 19"):
        raise ValueError("EGC 2016 Table 19 layout changed")
    h = c = None
    for i, r in enumerate(t):
        for j, x in enumerate(r):
            if x and str(x).strip() == tech:
                h, c = i, j
                break
        if h is not None:
            break
    if h is None:
        raise ValueError(f"EGC 2016 Table 19: {tech!r} not found")
    block = t[h + 1:h + 31]

    def find(lbl: str, sub: str | None = None) -> tuple:
        cur = ""
        for r in block:
            cur = str(r[0]).strip() if r[0] else cur
            if cur.startswith(lbl) and (sub is None or str(r[1] or "").strip() == sub):
                return r
        raise ValueError(f"EGC 2016 Table 19 {tech}: {lbl!r}/{sub!r} missing")

    yr = next(r for r in block if r[c] == "Commissioning year")
    years = block[block.index(yr) + 1]
    if years[c] != 2018:
        raise ValueError("EGC 2016 Table 19: 2018 column moved")

    def phasing(lbl: str) -> tuple[float, ...]:
        i = block.index(find(lbl))
        v = [_pct(block[i][c + 1]), _pct(block[i][c + 2]), _pct(block[i + 1][c]),
             _pct(block[i + 1][c + 1]), _pct(block[i + 1][c + 2])]
        v = [x for x in v if x]
        return tuple(x / sum(v) for x in v)       # printed shares are rounded

    def cost(lbl: str) -> float:
        cur = ""
        for r in block:
            cur = str(r[0]).strip() if r[0] else cur
            if (cur.startswith(lbl) or (r[0] is None and cur.startswith(lbl))) and str(r[1] or "").strip() == "Medium":
                return float(r[c])
        raise ValueError(f"EGC 2016 Table 19 {tech}: {lbl} medium missing")

    def medium_after(lbl: str) -> float:
        i = block.index(find(lbl, "High"))
        if str(block[i + 1][1]).strip() != "Medium":
            raise ValueError(f"EGC 2016 Table 19 {tech}: {lbl} rows out of order")
        return float(block[i + 1][c])

    one = lambda v: {"low": v, "medium": v, "high": v}  # noqa: E731
    return Plant(
        name=f"EGC2016 {tech}", commissioning=2018, size_mw=float(find("Reference plant size")[c + 1]),
        net_load_factor=float(find("Average load factor")[c + 1]),
        efficiency_gcv=float(find("Efficiency (HHV)")[c + 1]),
        life_years=int(str(find("Operating period")[c]).split()[0]),
        predev_gbp_kw=one(medium_after("Pre-development")), predev_phasing=phasing("Pre-development period"),
        construction_gbp_kw=one(medium_after("Construction")), construction_phasing=phasing("Construction period"),
        infrastructure_gbp_k=one(medium_after("Infrastructure")),
        fixed_om_gbp_mw_yr=cost("Fixed O&M"), variable_om_gbp_mwh=cost("Variable O&M"),
        insurance_gbp_mw_yr=cost("Insurance"), connection_uos_gbp_mw_yr=cost("Connection and Use of System"),
        hurdle=one(0.078))


def read_annex_generic(path: Path, tech: str, year: int = 2025) -> Plant:
    """EGC 2020 / 2023 'Technical and cost assumptions' column (same layout as EGC 2025, but
    lower-case labels and single-valued periods). Medium case."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = next(s for s in wb.sheetnames if s.lower() == "technical and cost assumptions")
    t = list(wb[sheet].iter_rows(values_only=True))
    j = next((k for k, v in enumerate(t[0]) if v and str(v).strip() == tech and t[1][k] == year), None)
    if j is None:
        raise ValueError(f"{path.name}: {tech!r} {year} not found")
    labels = [(str(r[0]).strip().lower() if r[0] else "", str(r[1]).strip().lower() if r[1] else "") for r in t]

    def row(name: str, sub: str = "") -> int:
        cur = ""
        for i, (a, b) in enumerate(labels):
            cur = a or cur
            if cur.startswith(name) and (not sub or b == sub):
                return i
        raise ValueError(f"{path.name}: {name!r}/{sub!r} missing")

    def val(name: str, sub: str = "") -> float:
        try:
            return float(t[row(name, sub or "medium")][j])
        except ValueError:
            return float(t[row(name)][j])

    def phasing(name: str) -> tuple[float, ...]:
        i = row(name)
        v = [t[k][j] or 0 for k in range(i, i + 8) if (t[k][1] or "").lower().startswith("year")]
        while v and not v[-1]:
            v.pop()
        if abs(sum(v) - 1) > 0.02:
            raise ValueError(f"{path.name}: {name} phasing sums to {sum(v)}")
        return tuple(x / sum(v) for x in v)

    one = lambda v: {"low": v, "medium": v, "high": v}  # noqa: E731
    return Plant(
        name=f"{path.name} {tech}", commissioning=year, size_mw=val("reference plant size"),
        net_load_factor=val("average load factor"), efficiency_gcv=val("average fuel efficiency"),
        life_years=int(val("operating lifetime")),
        predev_gbp_kw=one(val("pre-development cost")), predev_phasing=phasing("pre-development phasing"),
        construction_gbp_kw=one(val("construction cost")), construction_phasing=phasing("construction phasing"),
        infrastructure_gbp_k=one(val("infrastructure")), fixed_om_gbp_mw_yr=val("fixed o&m"),
        variable_om_gbp_mwh=val("variable o&m"), insurance_gbp_mw_yr=val("insurance"),
        connection_uos_gbp_mw_yr=val("connection and use of system"),
        hurdle=one(val("hurdle rate")))


@dataclass(frozen=True)
class Edition:
    """The plant set used to cost year y, with each plant's price base and source."""
    year: int
    ccgt: Plant
    ocgt: Plant
    coal: Plant
    coal_nox: Plant | None          # NOx retrofit added to a coal set that lacks it
    price_base: dict[str, int]
    source_id: dict[str, str]
    basis: dict[str, str] = field(default_factory=dict)


def _hurdle_for(year: int) -> float:
    """Gas-plant hurdle of the edition in force (spec §4.3 table)."""
    if year <= 2015:
        return 0.075                 # DECC EGC Dec 2013 Table 16 (and spec: no step for 2002-11)
    if year <= 2019:
        return 0.078                 # BEIS EGC 2016 Table 17 'CCGT & OCGT'
    return 0.075                     # EGC 2020/2023 CCGT; EGC 2025 read from the Annex


def edition(year: int) -> Edition:
    """The latest official cost set published by the end of `year` (earliest set before 2010).

    Gas: MM2010 (to 2011), PB2012 (2012), PB2013 (2013-15), EGC2016 (2016-19), EGC2020 (2020-22),
    EGC2023 (2023-24), EGC2025 (2025). OCGT before 2012: Mott MacDonald 2010 has no OCGT table,
    so the earliest set with one (PB2012) is used, labelled.
    Coal: no official unabated new-coal set after PB2012. MM2010 (FGD + SCR in the reference
    plant) to 2011; PB2012 'ASC with FGD' from 2012 plus the NOx retrofit of the set in force
    (PB2012 SCR/SNCR to 2015, Leigh Fisher 2016 coal NOx abatement from 2016)."""
    if not 2002 <= year <= 2025:
        raise ValueError(f"no edition rule for {year}")
    h = _hurdle_for(year)
    coal_h = 0.082 if year >= 2016 else h            # BEIS 2016 Table 17 'Coal plants All retrofits 8.20%'
    src, base, basis = {}, {}, {}
    if year <= 2011:
        ccgt = plant_from_table("mm2010_ccgt", h, 2009); src["ccgt"], base["ccgt"] = "mm2010-tableA1", 2009
        ocgt = plant_from_table("pb2012_ocgt", h, 2012); src["ocgt"], base["ocgt"] = "pb2012-ocgt", 2012
        basis["ocgt"] = "estimated"
    elif year == 2012:
        ccgt = plant_from_table("pb2012_ccgt", h, 2012); src["ccgt"], base["ccgt"] = "pb2012-ccgt", 2012
        ocgt = plant_from_table("pb2012_ocgt", h, 2012); src["ocgt"], base["ocgt"] = "pb2012-ocgt", 2012
    elif year <= 2015:
        ccgt = plant_from_table("pb2013_ccgt", h, 2012); src["ccgt"], base["ccgt"] = "pb2013-ccgt", 2012
        ocgt = plant_from_table("pb2013_ocgt", h, 2012); src["ocgt"], base["ocgt"] = "pb2013-ocgt", 2012
    elif year <= 2019:
        ccgt = read_egc2016("CCGT H Class"); src["ccgt"], base["ccgt"] = "beis-egc2016-table19", 2014
        ocgt = read_egc2016("OCGT 600MW 500 hr"); src["ocgt"], base["ocgt"] = "beis-egc2016-table19", 2014
    elif year <= 2022:
        f = EGC / "s18d-BEIS_EGC2020_KeyData.xlsx"
        ccgt = read_annex_generic(f, "CCGT H Class"); ocgt = read_annex_generic(f, "OCGT 600MW 500 hr")
        src["ccgt"] = src["ocgt"] = "beis-egc2020-keydata"; base["ccgt"] = base["ocgt"] = 2018
    elif year <= 2024:
        f = EGC / "s18d-EGC2023_AnnexA.xlsx"
        ccgt = read_annex_generic(f, "CCGT H Class"); ocgt = read_annex_generic(f, "OCGT 600MW 500 hr")
        src["ccgt"] = src["ocgt"] = "desnz-egc2023-annexA"; base["ccgt"] = base["ocgt"] = 2021
    else:
        p = read_desnz_plants(EGC / "s18d-EGC2025_AnnexA_current.xlsx")
        ccgt, ocgt = p[("Gas CCGT 93% Load Factor", 2030)], p[("Gas OCGT 760MW 30% Load Factor", 2030)]
        # 2025 Annex starts at 2030 commissioning for these columns; 2030 is the nearest printed
        src["ccgt"] = src["ocgt"] = "desnz-egc2025-annexA"; base["ccgt"] = base["ocgt"] = 2024
        coal_h = ccgt.hurdle["medium"]
    if year <= 2011:
        coal = plant_from_table("mm2010_coal", coal_h, 2009); nox = None
        src["coal"], base["coal"] = "mm2010-tableA3", 2009
    else:
        coal = plant_from_table("pb2012_coal", coal_h, 2012); src["coal"], base["coal"] = "pb2012-coal", 2012
        if year <= 2015:
            nox = plant_from_table("pb2012_scr", coal_h, 2012); src["coal_nox"], base["coal_nox"] = "pb2012-scr", 2012
        else:
            nox = plant_from_table("lf2016_coal_nox", coal_h, 2014); src["coal_nox"], base["coal_nox"] = "lf2016-coal-nox", 2014
        basis["coal"] = "estimated"          # carried forward: no later official new-coal set
    for k in ("ccgt", "ocgt", "coal"):
        basis.setdefault(k, "measured")
    return Edition(year, ccgt, ocgt, coal, nox, base, src, basis)


def restate(value: float, price_base: int, year: int) -> float:
    """Price-base money to year-y money with CPIH (ONS L522), spec §4.3."""
    idx = cpih()
    return value * idx[year] / idx[price_base]


def fixed_gbp_kw_yr(ed: Edition, tech: str) -> float:
    """Annualised fixed cost (capital at the edition's hurdle over its life + fixed O&M,
    insurance, connection), year-y money. Coal includes its NOx retrofit where added."""
    p = getattr(ed, tech)
    v = restate(annualised_fixed_gbp_kw_yr(p), ed.price_base[tech], ed.year)
    if tech == "coal" and ed.coal_nox is not None:
        v += restate(annualised_fixed_gbp_kw_yr(ed.coal_nox), ed.price_base["coal_nox"], ed.year)
    return v


def vom_gbp_mwh(ed: Edition, tech: str) -> float:
    p = getattr(ed, tech)
    v = restate(p.variable_om_gbp_mwh, ed.price_base[tech], ed.year)
    if tech == "coal" and ed.coal_nox is not None:
        v += restate(ed.coal_nox.variable_om_gbp_mwh, ed.price_base["coal_nox"], ed.year)
    return v


# ----------------------------------------------------------------- fuel prices

def _qep_quarterly(col_title: str) -> dict[tuple[int, int], float]:
    wb = openpyxl.load_workbook(QEP, read_only=True, data_only=True)
    rows = list(wb["3.2.1"].iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if r[0] == "Year" and r[1] == "Quarter")
    j = next((k for k, v in enumerate(rows[h]) if v and str(v).startswith(col_title)), None)
    if j is None:
        raise ValueError(f"QEP 3.2.1 column {col_title!r} missing")
    qn = {"Jan to Mar": 1, "Apr to Jun": 2, "Jul to Sep": 3, "Oct to Dec": 4}
    return {(int(r[0]), qn[str(r[1]).strip()]): float(r[j]) * 10       # p/kWh -> £/MWh
            for r in rows[h + 1:] if isinstance(r[0], int) and str(r[1]).strip() in qn
            and isinstance(r[j], (int, float))}


@lru_cache(maxsize=None)
def coal_price_quarterly() -> dict[tuple[int, int], float]:
    """£/MWh, DESNZ QEP 3.2.1 'Major power producers: Coal (pence per kWh)', cash terms."""
    return _qep_quarterly("Major power producers: Coal (pence per kWh)")


@lru_cache(maxsize=None)
def qep_gas_quarterly() -> dict[tuple[int, int], float]:
    """£/MWh, QEP 3.2.1 'Major power producers: Natural gas (pence per kWh)'."""
    return _qep_quarterly("Major power producers: Natural gas (pence per kWh)")


@lru_cache(maxsize=None)
def nbp_monthly() -> dict[tuple[int, int], float]:
    """£/MWh gross CV, Ofgem WMI 'Gas Prices: Day Ahead Contracts - Monthly Average (Great Britain)'."""
    t = WMI.read_text()
    if "Gas Prices: Day Ahead Contracts - Monthly Average (Great Britain)" not in t:
        raise ValueError("Ofgem WMI chart title changed")
    i = t.find("var options = ")
    o = json.loads(t[i + 14:t.find("\n", i)].rstrip(";"))
    out = {}
    for r in list(csv.reader(o["data"]["csv"].splitlines()))[1:]:
        m, _, y = r[0].split("/")
        out[(int(y), int(m))] = float(r[1]) * 1000 / KWH_PER_THERM / 100   # p/therm -> £/MWh
    return out


# Which gas price fossil Britain pays before the WMI series starts (Oct 2009). Spec §15.2
# item 3: NBP throughout, so "nbp" (central) reads Ofgem, *Wholesale Energy Markets in 2015*,
# Figure 5 (NBP day-ahead, vector read in prices.fig5_gas_ptherm, Jan 2001 - May 2015).
# "qep" is the old route, kept as a switch: DESNZ QEP 3.2.1, the price major power producers
# actually paid (a contract-lagged price, not the hub price).
GAS_PRE_WMI_ROUTES = ("nbp", "qep")
GAS_PRE_WMI = "nbp"


@lru_cache(maxsize=None)
def fig5_nbp_monthly() -> dict[tuple[int, int], float]:
    """£/MWh gross CV, Ofgem WEM 2015 Fig 5 NBP day-ahead (p/therm read by prices.py)."""
    from . import prices                      # local import: prices is heavy (PDF reader)
    return {k: v * 1000 / KWH_PER_THERM / 100 for k, v in prices.fig5_gas_ptherm().items()}


def gas_price_month(year: int, month: int, pre_wmi: str | None = None) -> tuple[float, str, str]:
    """(£/MWh, basis, source_id): NBP day-ahead (Ofgem WMI) from Oct 2009. Before that,
    `pre_wmi` (default GAS_PRE_WMI): "nbp" = Ofgem WEM 2015 Fig 5 NBP day-ahead,
    "qep" = DESNZ QEP 3.2.1 MPP gas (spec §15.2 item 3 chose NBP throughout)."""
    route = GAS_PRE_WMI if pre_wmi is None else pre_wmi
    if route not in GAS_PRE_WMI_ROUTES:
        raise ValueError(f"unknown pre-WMI gas route {route!r}; one of {GAS_PRE_WMI_ROUTES}")
    n = nbp_monthly()
    if (year, month) in n:
        return n[(year, month)], "measured", "ofgem-wmi-nbp-dayahead"
    if route == "nbp":
        f = fig5_nbp_monthly()
        if (year, month) not in f:
            raise ValueError(f"Ofgem WEM 2015 Fig 5 has no NBP price for {year}-{month:02d}")
        return f[(year, month)], "measured", "ofgem-wholesale-energy-markets-2015-fig5-vector"
    q = qep_gas_quarterly()[(year, (month - 1) // 3 + 1)]
    return q, "measured", "desnz-qep-321-mpp-gas"


def coal_price_month(year: int, month: int) -> tuple[float, str, str]:
    """(£/MWh, basis, source_id). QEP suppresses some late quarters ('..', almost no coal
    burnt): the nearest earlier published quarter is carried forward, labelled estimated."""
    q = coal_price_quarterly()
    k = (year, (month - 1) // 3 + 1)
    if k in q:
        return q[k], "measured", "desnz-qep-321-mpp-coal"
    prior = max(x for x in q if x < k)
    return q[prior], "estimated", "desnz-qep-321-mpp-coal-carried"


# ----------------------------------------------------------------- coal capacity

@lru_cache(maxsize=None)
def coal_capacity_uk() -> dict[int, float]:
    """MW, DUKES 5.7.A 'All generating companies' / 'Coal fired' (UK)."""
    rows = list(openpyxl.load_workbook(DUKES_57, read_only=True, data_only=True)["5.7"].iter_rows(values_only=True))
    t = next(i for i, r in enumerate(rows) if r[0] and str(r[0]).startswith("Table 5.7.A Capacity by fuel"))
    hdr = rows[t + 1]
    r = next(r for r in rows[t + 2:t + 40] if r[0] == "All generating companies" and r[1] == "Coal fired")
    return {int(h): float(v) for h, v in zip(hdr[2:], r[2:]) if h and str(h).isdigit() and isinstance(v, (int, float))}


@lru_cache(maxsize=None)
def kilroot_mw() -> dict[int, float]:
    """Northern Ireland coal (Kilroot), MW, from the DUKES 5.11 annual station lists.

    Sheet 'DUKES y' is read for year y; years before the first sheet (2004) take the 2004
    value, labelled. Kilroot closed its coal units in September 2023."""
    wb = openpyxl.load_workbook(DUKES_511, read_only=True, data_only=True)
    out = {}
    for sn in wb.sheetnames:
        m = re.fullmatch(r"DUKES (\d{4})", sn)
        if not m:
            continue
        mw = 0.0
        for r in wb[sn].iter_rows(values_only=True):
            cells = [str(x) for x in r if isinstance(x, str)]
            if any("Kilroot" in c for c in cells) and any(re.search(r"(?i)\bcoal", c) for c in cells) \
                    and not any("OCGT" in c for c in cells):
                nums = [x for x in r if isinstance(x, (int, float)) and 100 < x < 1000]
                mw = max(mw, nums[0]) if nums else mw
        out[int(m.group(1))] = mw
    return out


def coal_capacity_gb(year: int) -> tuple[float, str]:
    """(MW, basis): UK coal capacity less Kilroot. The cap on coal in fossil Britain."""
    k = kilroot_mw()
    ni = k.get(year, k[min(k)] if year < min(k) else 0.0)
    basis = "measured" if year in k else "estimated"
    return max(0.0, coal_capacity_uk()[year] - ni), basis


# ----------------------------------------------------------------- the fleet

@dataclass(frozen=True)
class FossilYear:
    year: int
    coal_gw: float
    ccgt_gw: float
    ocgt_gw: float
    spare_gw: float
    energy_twh: float
    gen_twh: dict[str, float]
    fuel_gbp_m: dict[str, float]
    vom_gbp_m: float
    fixed_gbp_m: float
    coal_cap_gw: float
    basis: str
    source_id: str

    @property
    def total_gbp_m(self) -> float:
        return self.fixed_gbp_m + self.vom_gbp_m + sum(self.fuel_gbp_m.values())


COAL_RULES = ("forced", "allowed")
# Spec §15.2 decision 1 (Richard, 11 Oct 2026, binding): "Coal is forced up to the coal
# capacity that actually existed each year (DUKES 5.7, GB), not merely allowed. Fossil
# Britain runs the coal stations Britain had; gas fills the rest." 'allowed' (coal built
# only where it is the cheapest strip, the wave-A rule) is kept as a switch for the report.
DEFAULT_COAL_RULE = "forced"


def cost_fossil_year(year: int, rows, reserve_margin: float, derating: dict[str, float],
                     step_mw: float = 50.0, coal_cap_mw: float | None = None,
                     transport=None, coal_rule: str = DEFAULT_COAL_RULE,
                     firm_requirement_mw: float | None = None) -> FossilYear:
    """Size and cost the fossil fleet for one calendar year.

    rows: half-hours with `.day` (date) and `.residual` (MW) — `residual.residual_year` rows
    for 2016-2025, or `residual_neso.residual_neso_year` rows for any year (parent joins).
    derating: {'ccgt', 'ocgt'} (capacity.read_derating); coal takes the CCGT factor.
    transport: optional `GasTransport` already in year-y money (gas plant only).
    coal_rule: 'forced' (default, spec §15.2 decision 1): the bottom strips of the
    load-duration curve, up to the year's actual GB coal capacity, are coal whatever they
    cost; 'allowed': coal competes strip by strip with CCGT and OCGT (wave-A rule).

    1. Sizing (screening curve): strips of `step_mw` from the bottom of the load-duration
       curve; with coal forced, the first strips up to the coal cap are coal (the stations
       Britain had, run as baseload); every other strip is built as the cheaper of CCGT and
       OCGT at its running share, using demand-weighted mean fuel prices.
    2. Spare: de-rated capacity must reach peak x (1 + margin); shortfall bought as OCGT,
       fixed cost only.
    3. Fuel: each half-hour dispatched in merit order of that month's running cost
       (coal vs CCGT; OCGT last), fuel charged at the month's price."""
    if coal_rule not in COAL_RULES:
        raise ValueError(f"unknown coal rule {coal_rule!r}; one of {COAL_RULES}")
    ed = edition(year)
    cap = coal_capacity_gb(year)[0] if coal_cap_mw is None else coal_cap_mw
    load = [max(0.0, r.residual) for r in rows]
    months = [r.day.month for r in rows]
    n = len(load)
    hours = n / 2
    gas = {m: gas_price_month(year, m)[0] for m in range(1, 13)}
    coal_p = {m: coal_price_month(year, m)[0] for m in range(1, 13)}
    tv = transport.commodity_gbp_mwh_gas if transport else 0.0
    eta = {k: getattr(ed, k).efficiency_gcv for k in ("coal", "ccgt", "ocgt")}
    for k, v in eta.items():
        if not 0.2 < v < 0.65:
            raise ValueError(f"{year} {k}: efficiency {v} is not a gross-CV sent-out figure")
    mwh_m = {m: 0.0 for m in range(1, 13)}
    for l, m in zip(load, months):
        mwh_m[m] += l / 2
    tot = sum(mwh_m.values())
    gas_mean = sum(gas[m] * w for m, w in mwh_m.items()) / tot
    coal_mean = sum(coal_p[m] * w for m, w in mwh_m.items()) / tot
    fuel_mean = {"coal": coal_mean, "ccgt": gas_mean + tv, "ocgt": gas_mean + tv}
    fc = {k: fixed_gbp_kw_yr(ed, k) for k in eta}
    for k in ("ccgt", "ocgt"):
        if transport:
            fc[k] += transport.fixed_gbp_kw_e_yr(eta[k])
    vom = {k: vom_gbp_mwh(ed, k) for k in eta}
    vc = {k: fuel_mean[k] / eta[k] + vom[k] for k in eta}

    asc = sorted(load)
    peak = asc[-1]
    mw = {"coal": 0.0, "ccgt": 0.0, "ocgt": 0.0}
    lo = 0.0
    fixed = 0.0
    while lo < peak:
        step = step_mw
        if coal_rule == "forced" and 1e-9 < cap - mw["coal"] < step_mw:
            step = cap - mw["coal"]         # last coal strip: exactly up to the cap
        hi = min(lo + step, peak)
        share = (n - bisect.bisect_left(asc, (lo + hi) / 2)) / n
        w = hi - lo
        coal_left = mw["coal"] + w <= cap + 1e-9
        if coal_rule == "forced" and coal_left:
            p = "coal"                      # the coal stations Britain had (§15.2 dec. 1)
        else:
            opts = ["ccgt", "ocgt"] + (["coal"] if coal_rule == "allowed" and coal_left else [])
            p = min(opts, key=lambda q: fc[q] * 1000 * w + vc[q] * w * share * hours)
        mw[p] += w
        fixed += fc[p] * 1000 * w
        lo = hi
    drf = {"coal": derating["ccgt"], "ccgt": derating["ccgt"], "ocgt": derating["ocgt"]}
    derated = sum(mw[k] * drf[k] for k in mw)
    # firm_requirement_mw (0057/0065, contracted basis): gross peak × (1 + m) − contracted imports
    need = peak * (1 + reserve_margin) if firm_requirement_mw is None else firm_requirement_mw
    spare = max(0.0, need - derated) / drf["ocgt"]
    fixed += fc["ocgt"] * 1000 * spare

    # dispatch, month by month merit order (no carbon)
    gen = {k: 0.0 for k in mw}
    fuel = {k: 0.0 for k in mw}
    for l, m in zip(load, months):
        price = {"coal": coal_p[m], "ccgt": gas[m] + tv, "ocgt": gas[m] + tv}
        order = sorted(("coal", "ccgt"), key=lambda k: price[k] / eta[k] + vom[k]) + ["ocgt"]
        left = l
        for k in order:
            g = min(left, mw[k] if k != "ocgt" else mw[k] + spare)
            if g <= 0:
                continue
            gen[k] += g / 2
            fuel[k] += g / 2 / eta[k] * price[k]
            left -= g
        if left > 1e-6:
            raise ValueError(f"{year}: dispatch short by {left:.0f} MW")
    vom_total = sum(gen[k] * vom[k] for k in gen)
    basis = "estimated" if "estimated" in ed.basis.values() else "modelled"
    return FossilYear(
        year=year, coal_gw=mw["coal"] / 1000, ccgt_gw=mw["ccgt"] / 1000, ocgt_gw=mw["ocgt"] / 1000,
        spare_gw=spare / 1000, energy_twh=sum(gen.values()) / 1e6,
        gen_twh={k: v / 1e6 for k, v in gen.items()}, fuel_gbp_m={k: v / 1e6 for k, v in fuel.items()},
        vom_gbp_m=vom_total / 1e6, fixed_gbp_m=fixed / 1e6, coal_cap_gw=cap / 1000,
        basis=basis, source_id=";".join(sorted(set(ed.source_id.values()))))


# ----------------------------------------------------------------- s23b: no net-zero policy (spec §15.3)

# Flue-gas desulphurisation (FGD) retrofit. Neither PB 2012 nor Leigh Fisher 2016 prints a
# stand-alone FGD retrofit table (they print the NOx/SCR retrofit, used above). The only
# sourced figure found in the stored cost editions is Mott MacDonald (2010), *UK Electricity
# Generation Costs Update*, p.45 (PDF page 45): "FGD, which once accounted for some 15-20% of
# capex has fallen over time such that FGD and SCR together typically account for some 10-15%
# of capex." FGD retrofit capex is taken as 20% (the top of the FGD-alone range: dearer
# fossil Britain, the conservative side) of MM2010's new ASC coal EPC cost (Table A.3, Nth of a
# kind, medium, 2009 money). Life: the retrofit life of the NOx retrofit in force (LF2016
# 15 years; PB2012 5 years), so FGD is paid off as fast as SCR. FGD running costs (limestone,
# extra auxiliary power) are not printed in any stored edition: GAP, left out (it would make
# fossil Britain dearer, so leaving it out makes the cost of net zero larger; flagged).
FGD_SHARE_OF_COAL_EPC = 0.20
FGD_QUOTE = ("FGD, which once accounted for some 15-20% of capex has fallen over time such that "
             "FGD and SCR together typically account for some 10-15% of capex.")
FGD_SOURCE = "mm2010-p45-fgd-share"
ABATEMENT = ("FGD", "SCR", "FGD+SCR")


def fgd_retrofit(ed: Edition) -> tuple[Plant, int, str]:
    """(FGD retrofit as a `Plant`, price base, source_id) for the edition in force."""
    mm = cost_editions()["tables"]["mm2010_coal"]["rows"]
    epc = mm["epc_gbp_kw"]["medium"] * FGD_SHARE_OF_COAL_EPC
    life_src = ed.coal_nox if ed.coal_nox is not None else plant_from_table("pb2012_scr", ed.coal.hurdle["medium"], 2012)
    one = lambda v: {"low": v, "medium": v, "high": v}  # noqa: E731
    p = Plant(name="fgd_retrofit", commissioning=2009, size_mw=1000.0, net_load_factor=0.93,
              efficiency_gcv=1.0, life_years=life_src.life_years, predev_gbp_kw=one(0.0), predev_phasing=(1.0,),
              construction_gbp_kw=one(epc), construction_phasing=(1.0,), infrastructure_gbp_k=one(0.0),
              fixed_om_gbp_mw_yr=0.0, variable_om_gbp_mwh=0.0, insurance_gbp_mw_yr=0.0,
              connection_uos_gbp_mw_yr=0.0, hurdle=dict(ed.coal.hurdle))
    return p, 2009, FGD_SOURCE


def scr_retrofit(ed: Edition) -> tuple[Plant, int, str]:
    """(NOx / SCR retrofit `Plant`, price base, source_id): the edition's own NOx retrofit; before
    2012 (MM2010, whose reference plant already has SCR) the first edition that prints a
    retrofit, PB2012 'Coal - retrofit SCR or NSCR' (labelled estimated by the caller)."""
    if ed.coal_nox is not None:
        return ed.coal_nox, ed.price_base["coal_nox"], ed.source_id["coal_nox"]
    return plant_from_table("pb2012_scr", ed.coal.hurdle["medium"], 2012), 2012, "pb2012-scr"


def abatement_gbp_kw_yr(ed: Edition, need: str | None) -> tuple[float, float, list[str]]:
    """Annualised cost of the pollution filters a kept coal station must fit: (£/kW/yr fixed,
    £/MWh variable, source_ids), year-y money. `need` is 'FGD', 'SCR', 'FGD+SCR' or None."""
    if need is None:
        return 0.0, 0.0, []
    if need not in ABATEMENT:
        raise ValueError(f"unknown abatement {need!r}; one of {ABATEMENT}")
    fixed = vom = 0.0
    src = []
    if "FGD" in need:
        p, base, s = fgd_retrofit(ed)
        fixed += restate(annualised_fixed_gbp_kw_yr(p), base, ed.year)
        src.append(s)
    if "SCR" in need:
        p, base, s = scr_retrofit(ed)
        fixed += restate(annualised_fixed_gbp_kw_yr(p), base, ed.year)
        vom += restate(p.variable_om_gbp_mwh, base, ed.year)
        src.append(s)
    return fixed, vom, src


def keep_open_gbp_kw_yr(ed: Edition, tech: str) -> float:
    """What it costs a year to keep an existing station open (no new capital): fixed O&M,
    insurance and connection/use-of-system charges of the edition's plant of that type,
    year-y money. tech: 'coal', 'ccgt' or 'ocgt'. Coal excludes the NOx retrofit (added only
    where the station needs it, `abatement_gbp_kw_yr`)."""
    p = getattr(ed, tech)
    v = (p.fixed_om_gbp_mw_yr + p.insurance_gbp_mw_yr + p.connection_uos_gbp_mw_yr) / 1000
    return restate(v, ed.price_base[tech], ed.year)


# New nuclear in fossil Britain's choice of new plant (§15.3 item 4): "nuclear only where the
# edition prints it; never before its stated earliest commissioning year". Earliest year a plant
# ordered under each edition could run, from the edition's own words and timings:
NUCLEAR_EARLIEST = {
    # MM2010 p.81: "the first dates at which new projects, not currently under development,
    # could be ordered would be 2012-13"; Table A.10 PWR Nth of a kind, medium: 4.0 years
    # pre-development + 5.0 construction -> 2012 + 9 = 2021.
    "mm2010": (2021, "mm2010-p81-tableA10"),
    # PB2012 p.53 'Nuclear' and PB2013 p.50 'Nuclear - PWR': pre-development 5.0 years and
    # construction 6.0 (FOAK medium) from the edition year -> 2012 + 11 = 2023, 2013 + 11 = 2024.
    "pb2012": (2023, "pb2012-p53-nuclear"),
    "pb2013": (2024, "pb2013-p50-nuclear-pwr"),
    # BEIS EGC 2016 Table 19 'Nuclear - PWR FOAK': no costs in the 2018 and 2020 commissioning
    # columns; first costed commissioning year 2025.
    "egc2016": (2025, "beis-egc2016-table19-nuclear"),
    # EGC 2020, 2023 and 2025 annexes print no nuclear column: never available.
}


def nuclear_available(year: int) -> tuple[bool, str]:
    """(can fossil Britain build new nuclear that runs in `year`?, reason)."""
    key = ("mm2010" if year <= 2011 else "pb2012" if year == 2012 else "pb2013" if year <= 2015
           else "egc2016" if year <= 2019 else None)
    if key is None:
        return False, "edition in force (EGC 2020/2023/2025) prints no nuclear"
    earliest, src = NUCLEAR_EARLIEST[key]
    if year < earliest:
        return False, f"edition {key} ({src}): earliest commissioning {earliest}"
    return True, f"edition {key} ({src})"
